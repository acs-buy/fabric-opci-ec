"""fn_documents : les livrables de l'etape 6 de l'ecran Revision, une fonction par document.

Element User data functions `fn_documents` de l'espace OPCI-DEV, identifiant 01b1f037-0147-4046-9106-bd80d60c1310.
Publie par l'API : python 61_DEPLOIEMENT/poser_fonction_ecran.py --nom fn_documents --ecrire, puis relu par
getDefinition. Ecrit le 04/10/2026 ; remplace essai_modele, qui ecrivait au coffre (regle du 28/09/2026 : tout
fichier se depose dans SharePoint par la base et se lit par raccourci).

LES 4 FONCTIONS, chacune appelee par un bouton (elle rend une phrase) :
  produire_livrable(entite, arrete, livrable, format)  le livrable et le format passent en parametre : DOCX pour
                                               « Produire en Word », XLSX pour « Produire en Excel » ; le bouton ne
                                               produit que le fichier de son format
  produire_comptes_annuels(entite, arrete)     Word puis Excel, dans une meme version (livrable COMPTES_ANNUELS)
  produire_rapport(entite, arrete)             rapport de l'expert-comptable, forme arretee (livrable ATTESTATION)
  produire_dip(entite, arrete)                 document d'information periodique (livrable DIP)
Chacune pose son message dans la zone de message de l'ecran Revision (dbo.pr_poser_message_ecran), reussite ou refus.

UNE VERSION, PLUSIEURS FORMATS. ref_livrable.formats dit les formats d'une version : DOCX et XLSX pour les comptes
annuels, DOCX pour le rapport et le DIP. pr_enregistrer_document rattache un fichier a la demande en cours
(DEMANDEE) et ne la passe a PRODUITE que lorsque chaque format attendu est enregistre. Un fichier se rattache donc a
la version ouverte qui attend encore son format, et a une version neuve sinon (pr_demander_document). La version
ouverte n'est reprise que si ses fichiers sont posterieurs a la derniere decision de revue (le dossier n'a pas ete
rouvert depuis) et qu'aucun n'est perime : ses fichiers decrivent alors les memes comptes.

L'ENCHAINEMENT :
  1. interface de depot presente en base, sinon refus qui le dit ;
  2. modele Word du cabinet lu au site du cabinet par le raccourci du Coffre (Files/fec/Modeles/...) ;
  3. la version : celle qui attend encore ce format, ou pr_demander_document, la garde de la base ; une demande
     REFUSEE rend son motif tel quel ;
  4. lecture de la base, fabrication en memoire (docxtpl pour le Word, openpyxl pour l'Excel) ; pour les comptes
     annuels, la version de l'annexe est fixee a la 1re lecture (v_annexe_version), toutes les cellules se lisent a
     cette version (fn_annexe_cellules), et la production est refusee avant tout depot quand l'annexe a change de
     version depuis, que l'autre format du meme document porte une autre version de l'annexe, ou que ses cellules
     sont en retard sur une ecriture (les refus de pr_enregistrer_document, avec leurs textes) ; le rapport l'est
     aussi sur le retard ;
  5. pour chaque fichier : empreinte SHA-256 des octets, pr_deposer_classeur au sous-dossier Livrables avec cette
     empreinte (la base la recalcule et refuse l'ecrasement), puis l'empreinte de la copie stockee, relue par le
     raccourci fec ;
  6. l'enregistrement : pr_enregistrer_document pour les comptes annuels, avec la version de l'annexe que le
     fichier imprime (@annexe_version), et pour le DIP ; pour le rapport, pr_produire_attestation, qui le rattache a
     la version des comptes annuels qu'il couvre et l'enregistre.
Le rapport a une source unique, la base : la forme arretee, son intitule, son texte et la donnee de chaque
emplacement (v_rapport_expert_comptable, ref_forme_rapport). Le DIP lit v_dip_arrete et v_dip_rubrique_1 a 6.
Rien ne s'ecrit au coffre. La fabrication est le module pur 61_DEPLOIEMENT/fabrication_livrables.py, recopie
ci-dessous par assembler_fn_documents.py et essaye seul par produire_livrables_local.py.

Connexions : DossierOPCI (base) et Coffre (lakehouse, lecture du modele par le raccourci fec). Bibliotheques :
docxtpl (et python-docx qu'il embarque), openpyxl. Limites du service : 240 s par invocation.
"""

# >>> fabrication_livrables.py : copie integrale, ne pas modifier ici (61_DEPLOIEMENT/assembler_fn_documents.py)
# -*- coding: utf-8 -*-
u"""Fabrication des livrables de l'etape 6 : comptes annuels (Word et Excel), rapport de l'expert-comptable (Word),
DIP (Word). MODULE PUR : aucune dependance a Fabric, aucune ecriture.

UN SEUL CODE POUR LA FONCTION ET POUR LE BANC D'ESSAI. Ce fichier est la source ; assembler_fn_documents.py le
recopie tel quel dans 60_FABRIC/fn_documents.UserDataFunction/function_app.py, entre 2 marques, et le banc
produire_livrables_local.py refuse de tourner si la copie differe. Ce que le banc prouve vaut donc pour le code
publie.

CE QUE LE MODULE FAIT : il lit la base par un lecteur qu'on lui passe, lire(sql, params) -> liste de dict, en
SELECT seulement ; il fabrique les octets d'un fichier en memoire ; il calcule leur empreinte SHA-256. Il ne
demande, ne depose ni n'enregistre rien : la fonction Fabric le fait, par la base (pr_demander_document,
pr_deposer_classeur, pr_enregistrer_document).

REGLES DE FABRICATION, appliquees sans exception :
  - tout montant est lu dans une vue ou une table de la base ; le module ne calcule, ne somme ni n'arrondit aucun
    montant (le total de l'onglet Balance est une formule Excel, calculee par le tableur a l'ouverture) ;
  - LES CELLULES DE L'ANNEXE SONT CELLES DE LA TABLE DE LA BASE (dbo.cellule_annexe), lues a une version : celle que
    la production fixe a sa 1re lecture (dbo.fn_annexe_cellules), ou la derniere (dbo.v_annexe_cellule). Le module
    ne calcule, ne deduit ni ne complete aucune cellule : il imprime le texte imprime de la table (texte_imprime),
    les lignes de la table dans l'ordre de la table (ordre), et met en page ;
  - une cellule a remplir, en attente ou sans objet s'imprime VIDE, sans surlignage ni fond ; une cellule lue,
    calculee ou saisie imprime son texte ; le Word ne dit jamais « à remplir » dans l'annexe. La devise de
    comptabilite (311-1) s'imprime dans les titres des etats, du 335-8 et sur la page de garde, vide de meme ;
  - une cellule que le modele imprime et que la table ne porte pas s'imprime vide et se note dans d['sans_cellule'],
    liste de (article, ligne, colonne) ;
  - la ligne de chaque cellule imprimee porte sa cle (cles : (article, ligne, colonne) par cellule, None pour un
    texte du modele), que la mise en page ignore et que la preuve relit ;
  - dans le bilan et le compte de resultat, une valeur NULL hors rubrique s'imprime « à remplir », mise en evidence ;
  - l'appel de note de l'article 333-3 s'imprime en exposant, comme au recueil ;
  - les libelles et l'ordre des lignes des modeles 321-2, 321-6 et 322-2 sont ceux du recueil, embarques ici
    (MODELES_ETATS) ; la base doit rendre les memes codes dans le meme ordre, sinon la production est refusee ;
  - le rapport de l'expert-comptable a une source unique, la base : la forme arretee, son intitule, son texte (les
    exemples E1 a E4 de la NP 2300, ref_forme_rapport) et la donnee de chaque emplacement (v_rapport_expert_comptable) ;
    le module n'embarque aucun texte de la norme ;
  - le DIP lit v_dip_arrete et les vues v_dip_rubrique_1 a v_dip_rubrique_6, une par rubrique ; la 7e est le texte de
    la page « Le DIP » du 06/09/2026 ; une rubrique dont la vue rend la ligne « à remplir » s'imprime « à remplir » ;
  - le reglement s'ecrit « règlement ANC n° 2021-09, modifié par le règlement ANC n° 2024-01 » ;
  - le Word des comptes annuels est un livrable, depuis le 05/10/2026 : page de garde sans en-tete ni numero,
    sommaire automatique (champ TOC mis a jour par Word a l'ouverture), titres « Bilan », « Compte de résultat »,
    « Annexe », et chaque tableau a la forme de son modele au recueil (CARTES_ETATS, CARTES_ANNEXE) ; la forme ne
    change aucune valeur : les cellules arrivent de la table que l'ecran lit aussi.

FORMAT DES NOMBRES : 1 234 567,89 et -1 234,50 (espace insecable), parts sans decimale quand elles sont
entieres, pourcentage a 2 decimales suivi de « % », surface entiere suivie de « m² », date JJ/MM/AAAA.
"""
import datetime
import decimal
import hashlib
import io
import re

D = decimal.Decimal
A_REMPLIR = u'à remplir'
A_REMPLIR_INSECABLE = u'à' + u'\u00a0' + u'remplir'   # le Word l'imprime sans coupure de ligne
AR = u'\x00'                      # marque d'un « à remplir » dans un texte, rendu mis en evidence
EXPOSANT = u'\x01'                # ce qui suit dans le morceau s'imprime en exposant : l'appel de note de 333-3
NBSP = u' '
REGLEMENT = u'règlement ANC n° 2021-09, modifié par le règlement ANC n° 2024-01'
MODELE_WORD = u'fec/Modeles/modele_livrables_opci.docx'   # relatif a Files du Coffre : raccourci du site du cabinet


class Refus(Exception):
    u"""Production refusee : le message se rend tel quel a l'utilisateur."""


# ============================================================================ formats

def _dec(v):
    if v is None:
        return None
    return v if isinstance(v, D) else D(str(v))


def montant(v):
    u"""1 234 567,89 ; -1 234,50 ; jamais de parentheses. Le format a 2 decimales est celui de l'affichage."""
    d = _dec(v)
    if d is None:
        return None
    d = d.quantize(D(u'0.01'), rounding=decimal.ROUND_HALF_UP)
    t = u'{:,.2f}'.format(abs(d)).replace(u',', NBSP).replace(u'.', u',')
    return (u'-' if d < 0 else u'') + t


def nombre(v):
    u"""Un nombre de parts ou un denombrement : sans decimale s'il est entier, ses decimales significatives sinon."""
    d = _dec(v)
    if d is None:
        return None
    if d == d.to_integral_value():
        return (u'-' if d < 0 else u'') + u'{:,}'.format(abs(int(d))).replace(u',', NBSP)
    ent, _, dec = format(abs(d).normalize(), u'f').partition(u'.')
    return (u'-' if d < 0 else u'') + u'{:,}'.format(int(ent)).replace(u',', NBSP) + u',' + dec


def pourcentage(v):
    u"""Une fraction de la base (0,081441) en pourcentage a 2 decimales : 8,14 %."""
    d = _dec(v)
    if d is None:
        return None
    p = (d * 100).quantize(D(u'0.01'), rounding=decimal.ROUND_HALF_UP)
    return (u'-' if p < 0 else u'') + u'{:,.2f}'.format(abs(p)).replace(u',', NBSP).replace(u'.', u',') + NBSP + u'%'


def surface(v):
    d = _dec(v)
    if d is None:
        return None
    return u'{:,}'.format(int(d.to_integral_value(rounding=decimal.ROUND_HALF_UP))).replace(u',', NBSP) + NBSP + u'm²'


def date_fr(v):
    if v is None:
        return None
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime(u'%d/%m/%Y')
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', str(v))
    return u'%s/%s/%s' % (m.group(3), m.group(2), m.group(1)) if m else str(v)


def date_iso(v):
    if v is None:
        return None
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime(u'%Y-%m-%d')
    return str(v)[:10]


def ou_ar(t):
    u"""Le texte, ou la regle commune d'une cellule sans valeur (« à remplir ») quand il manque."""
    return a_remplir() if t is None or (isinstance(t, str) and not t.strip()) else t


def empreinte(octets):
    return hashlib.sha256(octets).hexdigest()


# Le prefixe du nom de fichier d'un livrable. Le livrable ATTESTATION de la base est le « Rapport de l'expert-
# comptable » (regle du 02/10/2026, NP 2300, paragraphe A10) : compte rendu de travaux (E4) pour une entite soumise au
# commissariat aux comptes, attestation (E1 a E3) sur motif. Le fichier porte ce nom, quelle que soit la forme ;
# pr_enregistrer_document ne lit que la version (« _v<n>. ») et l'extension du nom (definition relue le 04/10/2026).
PREFIXES = {u'ATTESTATION': u'RAPPORT_EXPERT_COMPTABLE'}

# LES NOMS QU'UN MESSAGE DONNE au livrable et a l'etat, depuis le 05/10/2026 : un message de la zone de l'ecran Revision
# ne montre aucun code de la base (ATTESTATION, BILAN_ACTIF), et tient en 2 lignes de 870 px (poser_ecran_revision,
# lignes_msg_py ; preuve : prouver_messages_fn_documents.py). L'etat se nomme avec l'article de son modele au recueil.
NOMS_LIVRABLES = {u'COMPTES_ANNUELS': u'Comptes annuels', u'ATTESTATION': u'Rapport de l’expert-comptable',
                  u'DIP': u'Document d’information périodique'}
NOMS_ETATS = {u'BILAN_ACTIF': (u'de l’actif du bilan', u'321-2'), u'BILAN_PASSIF': (u'du passif du bilan', u'321-6'),
              u'RESULTAT': (u'du compte de résultat', u'322-2')}


def guillemets(t):
    u"""« t », les guillemets tenus au texte par une espace insecable : la zone de message ne coupe pas la ligne entre
    un guillemet et son texte (elle ne coupe qu'a l'espace ordinaire)."""
    return u'«' + NBSP + t + NBSP + u'»'


def insecables(t):
    u"""L'espace avant « ; » et « : » rendue insecable, comme le veut la typographie : une ligne de la zone de message
    ne commence jamais par un signe double."""
    return t.replace(u' ;', NBSP + u';').replace(u' :', NBSP + u':')


def nom_fichier(livrable, entite, arrete, version, extension):
    u"""Nom versionne d'un livrable : <livrable>_<entite>_<arrete>_v<version>.<ext>, le livrable ATTESTATION sous le
    nom du rapport de l'expert-comptable, RAPPORT_EXPERT_COMPTABLE_<entite>_<arrete>_v<version>.docx."""
    return u'%s_%s_%s_v%d.%s' % (PREFIXES.get(livrable, livrable), entite, arrete, int(version), extension)


def de_elide(nom):
    u"""« de » s'elide devant une voyelle ou un h muet."""
    return (u'd’' if nom[:1].lower() in u'aeiouyhéèêàâîôû' else u'de ') + nom


# ============================================================================ modeles du recueil

# Modeles 321-2, 321-6 et 322-2 du reglement ANC n° 2021-09, modifie par le reglement ANC n° 2024-01 (recueil,
# pages 90 a 94 du fichier) : code de ref_ligne_etat, libelle imprime au recueil, graisse au recueil. Le banc
# produire_livrables_local.py retrouve chaque libelle dans le texte du PDF.
MODELES_ETATS = {
    u'BILAN_ACTIF': [
        (u'A_IMMO', u'Actifs à caractère immobilier', 1),
        (u'A_IMMO_1', u'Immeubles en cours, construits ou acquis et droits réels', 0),
        (u'A_IMMO_2', u'Contrats de crédit-bail', 0),
        (u'A_IMMO_3', u'Parts des sociétés de personnes article L214-36 2ème alinéa', 0),
        (u'A_IMMO_4', u'Parts et actions des sociétés article L214-36 3ème alinéa', 0),
        (u'A_IMMO_5', u'Actions négociées sur un marché réglementé article L214-36 4ème alinéa', 0),
        (u'A_IMMO_6', u'OPCI et organismes étrangers équivalents article L214-36 5ème alinéa', 0),
        (u'A_IMMO_7', u'Avances preneurs sur crédit-bail', 0),
        (u'A_IMMO_8', u'Autres actifs à caractère immobilier (1)', 0),
        (u'A_FIN', u'Dépôts et instruments financiers non immobiliers', 1),
        (u'A_FIN_1', u'Actions et valeurs assimilées', 0),
        (u'A_FIN_2', u'Obligations et valeurs assimilées', 0),
        (u'A_FIN_3', u'Titres de créances', 0),
        (u'A_FIN_4', u'Organismes de placement collectif à capital variable (OPCVM et FIA)', 0),
        (u'A_FIN_5', u'Dépôts', 0),
        (u'A_FIN_6', u'Opérations temporaires sur titres', 0),
        (u'A_FIN_7', u'Instruments financiers à  terme', 0),     # 2 espaces, lettre du recueil (page 90)
        (u'A_AUTRE', u'Autres actifs', 1),
        (u'A_AUTRE_1', u'Créances locataires', 0),
        (u'A_AUTRE_2', u'Autres créances', 0),
        (u'A_AUTRE_3', u'Dépôts à vue', 0),
        (u'A_TOTAL', u"Total de l'actif", 1),
    ],
    u'BILAN_PASSIF': [
        (u'P_CP', u'Capitaux propres (= actif net)', 1),
        (u'P_CP_1', u'Capital', 0),
        (u'P_CP_2', u'Report des plus-values nettes réalisées (1)', 0),
        (u'P_CP_3', u'Report des résultats nets antérieurs (1)', 0),
        (u'P_CP_4', u"Résultat de l'exercice (1)", 0),
        (u'P_CP_5', u"Acomptes versés au cours de l'exercice (1)", 0),
        (u'P_PROV', u'Provisions pour risques et charges', 1),
        (u'P_IMMO', u'Passifs liés aux actifs à caractère immobilier', 1),
        (u'P_IMMO_1', u'Dettes sur rentes viagères', 0),
        (u'P_IMMO_2', u'Provisions sur droits nés de contrats de crédit-bail', 0),
        (u'P_IMMO_3', u'Provisions sur filiales et participations', 0),
        (u'P_IMMO_4', u'Autres passifs liés aux actifs à caractère immobilier', 0),
        (u'P_IF', u'Instruments financiers', 1),
        (u'P_IF_1', u'Opérations de cessions sur instruments financiers', 0),
        (u'P_IF_2', u'Opérations temporaires sur titres', 0),
        (u'P_IF_3', u'Instruments financiers à terme', 0),
        (u'P_DET', u'Dettes', 1),
        (u'P_DET_1', u'Emprunts envers les établissements de crédit', 0),
        (u'P_DET_2', u'Autres emprunts et dettes financières', 0),
        (u'P_DET_3', u'Concours bancaires courants', 0),
        (u'P_DET_4', u'Dépôts de garantie reçus', 0),
        (u'P_DET_5', u"Autres dettes d'exploitation", 0),
        (u'P_TOTAL', u'Total du passif', 1),
    ],
    u'RESULTAT': [
        (u'R_PI', u"Produits de l'activité immobilière", 1),
        (u'R_PI_1', u'Produits immobiliers', 0),
        (u'R_PI_2', u'Produits sur parts et actions des entités à caractère immobilier', 0),
        (u'R_PI_3', u'Autres produits sur actifs à caractère immobilier', 0),
        (u'R_TI', u'Total I', 1),
        (u'R_CI', u"Charges de l'activité immobilière", 1),
        (u'R_CI_1', u'Charges immobilières', 0),
        (u'R_CI_2', u'Charges sur parts et actions des entités à caractère immobilier', 0),
        (u'R_CI_3', u'Autres charges sur actifs à caractère immobilier', 0),
        (u'R_CI_4', u"Charges d'emprunt sur les actifs à caractère immobilier", 0),
        (u'R_TII', u'Total II', 1),
        (u'R_RAI', u"Résultat de l'activité immobilière (I - II)", 1),
        (u'R_PF', u'Produits sur opérations financières', 1),
        (u'R_PF_1', u'Produits sur dépôts et instruments financiers non immobiliers', 0),
        (u'R_PF_2', u'Autres produits financiers', 0),
        (u'R_TIII', u'Total III', 1),
        (u'R_CF', u'Charges sur opérations financières', 1),
        (u'R_CF_1', u'Charges sur dépôts et instruments financiers non immobiliers', 0),
        (u'R_CF_2', u'Autres charges financières', 0),
        (u'R_TIV', u'Total IV', 1),
        (u'R_ROF', u'Résultat sur opérations financières (III - IV)', 1),
        (u'R_TV', u'Autres produits (V)', 0),
        (u'R_TVI', u'Frais de gestion et de fonctionnement externes (VI)', 0),
        (u'R_TVII', u'Autres charges (VII)', 0),
        (u'R_RN', u'Résultat net (A= I - II + III - IV + V - VI - VII)', 1),
        (u'R_TVIII', u'Régularisations sur résultat net (VIII)', 0),
        (u'R_RNAR', u'Résultat net après régularisations (A +/- VIII)', 1),
        (u'R_PV', u'Plus-values réalisées nettes', 1),
        (u'R_PV_1', u'Plus-values réalisées nettes de frais sur actifs à caractère immobilier', 0),
        (u'R_PV_2', u'Plus-values réalisées nettes de frais sur dépôts et instruments financiers non immobiliers', 0),
        (u'R_TIX', u'Total IX', 1),
        (u'R_MV', u'Moins-values réalisées nettes', 1),
        (u'R_MV_1', u'Moins-values réalisées nettes de frais sur actifs à caractère immobilier', 0),
        (u'R_MV_2', u'Moins-values réalisées nettes de frais sur dépôts et instruments financiers non immobiliers', 0),
        (u'R_TX', u'Total X', 1),
        (u'R_RPMV', u'Résultat des Plus et moins-values réalisées nettes (B = IX-X)', 1),
        (u'R_TXI', u'Régularisations sur plus et moins-values réalisées nettes (XI)', 0),
        (u'R_PMVAR', u'Plus et moins-values réalisées nettes après régularisations (B +/- XI)', 1),
        (u'R_RE', u"Résultat de l'exercice (C= A +/- VIII + B +/- XI)", 1),
    ],
}
TITRES_ETATS = {u'BILAN_ACTIF': (u'Bilan', u'Actif'), u'BILAN_PASSIF': (u'Bilan', u'Passif'),
                u'RESULTAT': (u'Compte de résultat', None)}
NOTES_ETATS = {
    u'BILAN_ACTIF': u'(1) Les autres actifs à caractère immobilier comprennent les avances en compte courant et les '
                    u'dépôts et cautionnements versés.',
    u'BILAN_PASSIF': u'(1) Y compris comptes de régularisation',
}

SECTIONS_ANNEXE = [(u'331', u'Section 1 - Contenu de l’annexe'),
                   (u'332', u'Section 2 - Informations générales'),
                   (u'333', u'Section 3 - Evolution des capitaux propres'),
                   (u'334', u'Section 4 - Informations relatives aux expositions d’un OPCI'),
                   (u'335', u'Section 5 - Autres informations relatives au bilan et au compte de résultat'),
                   (u'336', u'Section 6 - Inventaire du portefeuille d’un OPCI')]

# Lignes en gras au recueil, relevees sur les polices et sur le mode de rendu du PDF (le gras simule par contour
# compte) : CL_1 de 335-2 est maigre au recueil ; 335-8 et 335-10 portent leur gras italique (ITALIQUE_ANNEXE).
GRAS_ANNEXE = {
    u'333-1': [u'CAPITAL', u'SD', u'SD_TOTAL', u'CP_TOTAL'], u'333-2': [u'MONTANTS_NETS'],
    u'333-3': [u'AN_DEBUT', u'AN_FIN'], u'334-2': [u'TOTAL', u'TOTAL_HORS_FRAIS'], u'334-3': [u'TOTAL'],
    u'334-5': [u'TOTAL', u'TOTAL_HORS_FRAIS'], u'334-6': [u'TOTAL'],
    u'334-7': [u'TOTAL_IMMEUBLES', u'TOTAL_AUTRES', u'TOTAL_GENERAL'],
    u'335-2': [u'CL', u'CL_ST', u'AC', u'AC_ST', u'TOTAL', u'DEP_TOTAL'], u'335-3': [u'TOTAL'],
    u'335-4': [u'EC', u'AUTRES', u'TOTAL'], u'335-5': [u'TOTAL'], u'335-6': [u'TOTAL'],
    u'335-8': [u'P', u'P_ST', u'C', u'C_ST', u'TOTAL'], u'335-10': [u'TOTAL'], u'335-12': [u'I', u'II'],
}

# Le signe imprime au modele de l'article 333-3 (page 99 du fichier) est celui que la table de l'annexe porte, a la
# lettre du recueil, son espace comprise (« + /- » et « +/- ») : colonne signe de la table.

NN1 = [(u'Exercice N', 1), (u'Exercice N-1', 2)]
ACOMPTES = [(u'Date de paiement', None), (u'Montants sur résultat net, montant total', None),
            (u'Montant unitaire', None), (u'Montants sur plus et moins-values nettes réalisées, montant total', None),
            (u'Montant unitaire', None), (u'Distributions réalisées, montant total', None), (u'Montant unitaire', None)]

# LES BLOCS DE L'ANNEXE, article par article.
# genre : tableau (colonnes du recueil ; source : la colonne k de la table de l'annexe, 'sig' le signe du modele que
# la table porte, 'ech' l'echeance, colonne 6 du 334-6), texte (2 colonnes « Information », « Texte de l'annexe »),
# phrase, une_ligne (une colonne par ligne de la table), sans_ligne (tableau des acomptes : une ligne par paiement de
# la table, de cle prefixe:..., puis la ligne prefixe_TOTAL), inventaire (336-2). codes : les lignes du bloc ; None
# prend toutes les lignes de l'article qu'aucun autre bloc ne nomme. nombres et lignes_nombre : colonnes et lignes de
# nombres (le format est celui du texte imprime de la table).
def T(entete, colonnes=NN1, codes=None, titre=None, note=None, nombres=(), lignes_nombre=()):
    return dict(genre=u'tableau', colonnes=[(entete, u'lib')] + list(colonnes), codes=codes, titre=titre,
                note=note, nombres=set(nombres), lignes_nombre=set(lignes_nombre))


def X(codes=None, avant=(), exiges=None):
    u"""Bloc de texte ; avant : libelles des informations que le recueil exige, imprimes en tete du bloc, la valeur
    lue sur la ligne de la table que CODES_SANS_LIGNE designe. exiges : {code: format} des informations que le recueil
    exige sans modele ; chacune se lit sur sa ligne de la table, et s'imprime vide quand la table ne la porte pas
    (d['sans_cellule'])."""
    return dict(genre=u'texte', codes=codes, avant=list(avant), exiges=dict(exiges or {}))


# Article 332-2 (recueil, p. 96 du fichier) : « Outre les mentions relatives aux regles et methodes comptables
# prevues par le reglement ANC n° 2020-07 [...] ».
MENTIONS_2020_07 = (u'Mentions relatives aux règles et méthodes comptables prévues par le règlement ANC n° 2020-07 '
                    u'relatif aux comptes annuels des organismes de placement collectif à capital variable')
# Le code de la ligne de la table qui porte une information imprimee par son libelle (bloc « avant »).
CODES_SANS_LIGNE = {MENTIONS_2020_07: u'MENTIONS_2020_07'}


ANNEXE = {
    u'331-1': [dict(genre=u'phrase', codes=[u'PRINCIPE'])],
    u'332-1': [X([u'STRATEGIE']),
               T(u'', [(u'Exercice {a0}', 1), (u'Exercice {a1}', 2), (u'Exercice {a2}', 3), (u'Exercice {a3}', 4),
                       (u'Exercice {a4}', 5)],
                 codes=[u'ACTIF_NET', u'NB_PARTS', u'VL_UNIT', u'DIST_CESSION', u'DIST_RESULTAT', u'CAPITALISATION'],
                 titre=u'Un tableau des éléments caractéristiques de l’OPCI au cours des cinq derniers exercices '
                       u'fait apparaître distinctement :',
                 lignes_nombre=[u'NB_PARTS'])],
    u'332-2': [X(avant=[MENTIONS_2020_07])],
    u'333-1': [T(u'Décomposition du poste « capitaux propres » au bilan')],
    u'333-2': [T(u'Souscriptions et rachats de l’exercice',
                 [(u'Nombre de parts ou actions', 3), (u'Montants bruts (hors frais et commissions)', 1),
                  (u'Frais et commissions*', 4)],
                 note=u'*Frais et commissions perçus par l’OPCI et non rétrocédés', nombres=[3])],
    u'333-3': [T(u'Evolution de l’actif net (= capitaux propres) au cours de l’exercice', [(u'', u'sig')] + NN1,
                 note=u'1 Le contenu de cette ligne fait l’objet d’une explication précise de la part de l’OPCI '
                      u'(versements en garantie, provisions pour risques et charges et autres variations d’actif '
                      u'net non détaillées ci-dessus).'),
               # l'explication du renvoi 1 que l'OPCI donne, sous la note du modele ; rien quand elle est sans objet
               X([u'RENVOI_1'])],
    u'334-1': [X()],
    u'334-2': [T(u'Ventilation par nature',
                 [(u'Exercice N-1', 2), (u'Augmentations', 3), (u'Diminutions', 4),
                  (u'Variation des différences d’estimation', 5), (u'Exercice N', 1), (u'Frais', 6)])],
    # 334-3 : le Word imprime les informations (bloc 1) avant le tableau « Ventilation par contrat » (bloc 0), dans
    # l'ordre du recueil (page 100 du fichier), par avant_bloc de la carte (u'334-3', 1) ; le rang des blocs ne change
    # pas, l'ecran et les autres livrables les lisant par ce rang
    u'334-3': [T(u'Ventilation par contrat',
                 [(u'Assiette du contrat', 3), (u'Exercice N-1', 2), (u'Augmentations', 4), (u'Diminutions', 5),
                  (u'Variation des différences d’estimation', 6), (u'Exercice N', 1)], codes=[u'TOTAL']),
               T(u'Information', [(u'Montant', 1)],
                 codes=[u'ASSIETTE', u'REDEVANCES_N', u'REDEVANCES_CUM', u'RESTANT_1AN', u'RESTANT_1_5ANS',
                        u'RESTANT_5ANS', u'PRIX_RESIDUEL', u'PRIX_RESIDUEL_1AN', u'PRIX_RESIDUEL_1_5ANS',
                        u'PRIX_RESIDUEL_5ANS'])],
    u'334-5': [T(u'Évolution des parts et actions des sociétés visées à l’article L214-36 du code monétaire et '
                 u'financier (2° et 3° du I)',
                 [(u'Exercice N-1', 2), (u'Cessions', 3), (u'Acquisitions', 4),
                  (u'Variation des différences d’estimation', 5), (u'Exercice N', 1), (u'Frais', 6)],
                 codes=[u'PARTS_2', u'PARTS_3', u'TOTAL', u'TOTAL_HORS_FRAIS']),
               X([u'SECTEURS'])],
    # 334-6 : « Ventilation par contrat » (recueil, p. 101 du fichier) : une ligne par contrat, les lignes CC: de la
    # table, puis sa ligne TOTAL (genre 'contrats')
    u'334-6': [dict(genre=u'contrats',
                    colonnes=[(u'Ventilation par contrat', u'lib'), (u'Echéance', u'ech'), (u'Exercice N-1', 2),
                              (u'Augmentations', 3), (u'Diminutions', 4),
                              (u'Variation des différences d’estimation', 5), (u'Exercice N', 1)],
                    codes=[u'COMPTES_COURANTS', u'ECHEANCE', u'TOTAL'])],
    u'334-7': [T(u'Ventilation par nature',
                 [(u'Plus-values', 3), (u'Moins-values', 4), (u'Plus et moins-values N', 1),
                  (u'Plus et moins-values N-1', 2)])],
    u'334-8': [X()],
    u'335-1': [X()],
    u'335-2': [T(u'Décomposition des postes du bilan',
                 codes=[u'CL', u'CL_1', u'CL_2', u'CL_3', u'CL_ST', u'AC', u'AC_1', u'AC_2', u'AC_3', u'AC_4',
                        u'AC_5', u'AC_6', u'AC_ST', u'TOTAL']),
               T(u'Evolution des dépréciations',
                 [(u'Situation N-1', 2), (u'Dotations de l’exercice', 3), (u'Reprises de l’exercice', 4),
                  (u'Situation N', 1)], codes=[u'DEP_CL', u'DEP_AC', u'DEP_TOTAL'])],
    u'335-3': [T(u'Détail des provisions pour risques et charges',
                 [(u'Situation Exercice N-1', 2), (u'Dotations de l’exercice', 3), (u'Reprises de l’exercice', 4),
                  (u'Situation Exercice N', 1)])],
    u'335-4': [T(u'Détail des emprunts', codes=[u'EC', u'EC_IMMO', u'EC_NON_IMMO', u'AUTRES', u'AUTRES_IMMO',
                                                u'AUTRES_NON_IMMO', u'TOTAL']),
               X([u'TAUX_FIXE', u'RATIOS'])],
    # 335-5 : colonnes 1 a 4 selon la note de ref_tableau_annexe, seule convention ecrite en base
    u'335-5': [T(u'Ventilation par maturité résiduelle',
                 [(u'Jusqu’à 1 an', 1), (u'[1 - 5 ans]', 2), (u'> 5 ans', 3), (u'Total', 4)])],
    u'335-6': [T(u'Décomposition du poste « Autres dettes d’exploitation » au bilan')],
    u'335-7': [X()],
    u'335-8': [T(u'', titre=u'Revenus financiers au {date} à en {devise}'),
               X([u'DETAIL_NATURE'], exiges={u'DETAIL_NATURE': u'texte'})],
    u'335-9': [X()],
    u'335-10': [T(u'Frais de gestion et de fonctionnement externes'),
                X([u'PCT_ACTIF_NET_MOYEN'], exiges={u'PCT_ACTIF_NET_MOYEN': u'pct'}),
                X([u'FRAIS_EXTERNES_NATURE'], exiges={u'FRAIS_EXTERNES_NATURE': u'texte'}),
                X([u'SURPERFORMANCE'], exiges={u'SURPERFORMANCE': u'montant'})],
    u'335-11': [X()],
    u'335-12': [T(u'Tableau d’affectation des sommes distribuables'),
                dict(genre=u'sans_ligne', colonnes=[(u'Tableau des acomptes versés au titre de l’exercice', u'lib')]
                     + ACOMPTES, lignes=[(u'Total acomptes', True)], prefixe=u'ACPT_N'),
                dict(genre=u'sans_ligne',
                     colonnes=[(u'Tableau des acomptes versés au titre de l’exercice précédent', u'lib')] + ACOMPTES,
                     lignes=[(u'Total acomptes', True)], prefixe=u'ACPT_N1'),
                X([u'DETAIL_CATEGORIES'], exiges={u'DETAIL_CATEGORIES': u'texte'})],
    u'336-1': [X()],
    u'336-2': [dict(genre=u'inventaire')],
    u'336-3': [dict(genre=u'une_ligne', colonnes=[(u'Libellé', u'LIBELLE', u'texte'),
                                                  (u'Quantité / nominal', u'QUANTITE', u'nombre'),
                                                  (u'Valeur actuelle', u'VALEUR', u'montant'),
                                                  (u'Devise de cotation', u'DEVISE', u'texte'),
                                                  (u'% de l’actif net', u'PCT', u'pct')])],
}

# 336-2, textes du recueil (pages 106 et 107 du fichier), coquille « participation » du 1er groupe gardee.
# L'inventaire est etabli par poste du bilan (336-1) : le groupe des immeubles porte le libelle du poste de 321-2.
TITRE_IMMEUBLES = u'Immeubles en cours, construits ou acquis et droits réels'
ENTETES_IMMEUBLES = [u'Adresse', u'Surface', u'Secteur', u'Date d’acquisition', u'Prix de revient',
                     u'Valeur actuelle', u'% de l’actif net']
# Le texte de l'article 336-2 qui introduit chaque tableau, a la lettre (page 106 du fichier) : le chapeau, puis un
# tiret par nature d'actif, comme au recueil ; les tableaux tiennent la place des enumerations.
TEXTE_336_2 = (u'Un inventaire détaillé des actifs à caractère immobilier est établi, actif par actif, à la clôture de '
               u'l’exercice par poste du bilan et évalué à la valeur actuelle en détaillant :')
POUR_IMMEUBLES = u'- pour les immeubles et droits réels :'
POUR_PARTS = (u'- pour les parts des sociétés non cotées à prépondérance immobilière visées aux 2° et 3° du I de '
              u'l’article L214-36 du code monétaire et financier, le tableau suivant est présenté. Il tient lieu de '
              u'tableau des filiales et participations.')
POUR_AUTRES = u'- pour les autres actifs à caractère immobilier sont détaillés :'
# 336-3, le texte de l'article qui introduit l'inventaire, a la lettre (page 107 du fichier)
TEXTE_336_3 = (u'Un inventaire détaillé de la rubrique des dépôts et instruments financiers non immobiliers est établi '
               u'et évalué à la valeur actuelle en détaillant :')
ENTETES_FILIALES = [u'', u'Nature des immeubles détenus (1)', u'Secteur d’activité et localisation',
                    u'Date d’acquisition des parts ou des sous-jacents', u'Coût de revient des titres frais exclus (2)',
                    u'Valeur actuelle des titres comptabilisée (3)', u'Différence d’estimation (3) - (2)']
GROUPE_CONTROLEES = (u'Filiales ou participation contrôlées au sens de l’article R214-83 du code monétaire et '
                     u'financier')
GROUPE_NON_CONTROLEES = (u'Filiales ou participations non contrôlées au sens de l’article R214-83 du code '
                         u'monétaire et financier')
# LES FILIALES QUE LA BASE NE CLASSE DANS AUCUN DES 2 GROUPES : une 3e ligne de groupe du tableau des parts, sous les 2
# du modele, tant qu'une filiale reste a classer au sens de l'article R. 214-83 du CMF (ecart au modele de la
# p. 107 accepte le 06/10/2026 ; ligne de groupe FIL_A_CLASSER de la table). Son libelle est forme sur la lettre des 2
# groupes du recueil.
GROUPE_A_CLASSER = (u'Filiales ou participations à classer au sens de l’article R214-83 du code monétaire et '
                    u'financier')
# LES LIGNES DE GROUPE DE LA TABLE et l'intitule que le Word imprime pour chacune, a la lettre du recueil.
GROUPES_FILIALES = {u'FIL_CONTROLEES': GROUPE_CONTROLEES, u'FIL_NON_CONTROLEES': GROUPE_NON_CONTROLEES,
                    u'FIL_A_CLASSER': GROUPE_A_CLASSER}
NOTE_FILIALES = (u'(1) Préciser le nombre et la nature des immeubles détenus par la filiale / participation : type '
                 u'de bien (ex. terrain nu, terrain et construction, construction sur sol d’autrui, autre droit '
                 u'réel, immeuble en cours..), secteur d’activité, localisation')
COLONNES_AUTRES_IMMO = [(u'Libellé', u'AUTRES_LIBELLE', u'texte'), (u'Quantité / nominal', u'AUTRES_QUANTITE', u'nombre'),
                        (u'Valeur actuelle', u'AUTRES_VALEUR', u'montant'), (u'% de l’actif net', u'AUTRES_PCT', u'pct')]

# ============================================================================ le Word, mise en forme des modeles
#
# LE WORD DES COMPTES ANNUELS EST UN LIVRABLE, depuis le 05/10/2026 : aucun element de l'ecran, la navigation de Word
# (sommaire automatique, titres, numeros de page), et chaque tableau a la forme de son modele au recueil (reglement
# ANC n° 2021-09, modifie par le reglement ANC n° 2024-01, pages 90 a 107 du fichier) : titre du modele dans la 1re
# cellule d'en-tete, intitules centres sans fond, cadre et filets verticaux, filets horizontaux par bloc (plein ou
# pointille) la ou le modele les trace, retraits et graisses du modele. La lettre du recueil est gardee, coquilles
# comprises : mots, casse, apostrophe droite ou courbe, capitale accentuee ou non, telles que le recueil les porte
# (lettre()). Ce qui suit ne porte que la forme : les valeurs des cellules sont le texte imprime de la table de
# l'annexe (_imprime, _valeur_texte), celle que l'ecran lit aussi.
#
# MISE EN PAGE, RETENUE LE 06/10/2026 : marges de page de 1,5 cm ;
# corps de page en 12 pt ; contenu des tableaux en 10 pt (le corps de 9 pt du recueil devient 10 pt, ses exceptions
# de 10 pt restent en 10 pt) ; le titre du modele en 12 pt, gras, bleu fonce #1F3864, y compris dans la 1re cellule de
# son tableau ; les titres (Titre, Titre 1 a 3) gardent leur taille ; proportions des colonnes du modele gardees, ce
# qui deborde passe sur 2 lignes ; pages en paysage pour les tableaux larges (334-2, 334-3, 334-5, 334-6, acomptes du
# 335-12, 336-2), le reste en portrait, les 3 etats, 333-3 et 334-7 compris ; 3,5 pt a droite des cellules de montant
# des 12 tableaux a marge de cellule etroite (MARGE_MONTANTS).

ITAL = u'\x02'                    # bascule l'italique dans un libelle : le « id. » des libelles de 334-7
PUCE = u'\x03'                    # la puce du modele, en police Symbol comme au recueil (335-12)
CORPS_PAGE = 12                   # corps de page des comptes annuels, en points
CORPS_TABLEAU = 10                # contenu des tableaux : le corps de 9 pt du recueil, et ses exceptions de 10 pt
TITRE_MODELE = 12                 # titre du modele, gras, bleu fonce du modele Word du cabinet
BLEU_TITRE = u'1F3864'
MARGE_PAGE = 850                  # marges de page des comptes annuels : 1,5 cm, en twips
PAGE_A4 = (11906, 16838)          # largeur et hauteur d'une page A4 en portrait, en twips
LARGEUR_UTILE = PAGE_A4[0] - 2 * MARGE_PAGE      # justification en portrait : 18 cm (10 206 twips)
LARGEUR_PAYSAGE = PAGE_A4[1] - 2 * MARGE_PAGE    # justification en paysage : 26,7 cm (15 138 twips)
# LES TABLEAUX LARGES, en paysage (06/10/2026) : les blocs de ces articles ; (article, rang du bloc) pour
# les seuls acomptes de 335-12. Les pages qui les portent sont en paysage, le reste du document en portrait.
PAYSAGE = (u'334-2', u'334-3', u'334-5', u'334-6', (u'335-12', 1), (u'335-12', 2), u'336-2')
# LES AUTRES TABLEAUX RESTENT EN PORTRAIT, aux proportions du modele : une valeur plus large que sa colonne y passe
# sur 2 lignes (06/10/2026 : proportions du modele gardees, ce qui deborderait encore passe sur 2 lignes),
# un montant coupe entre 2 groupes de chiffres, quels que soient les montants du dossier. Le bilan actif, le bilan
# passif, le compte de resultat, 333-3 et 334-7 en portrait, comme le modele (06/10/2026).
# Un etat (BILAN_ACTIF, BILAN_PASSIF, RESULTAT) ou un article nomme ici passerait en
# paysage, quels que soient les montants ; vide : aucun.
PAYSAGE_SUR_DECISION = ()
# MARGES DE CELLULE DU MODELE, gauche et droite, en twips : ecart entre l'axe du filet vertical du cadre et l'origine
# d'un libelle cale a gauche sans retrait, releve au recueil (pages 90 a 107 du fichier, PyMuPDF, origine du 1er
# caractere). Les modeles traces en filet
# de 0,96 pt (etats, 334-2, 334-6, 335-2, 335-6, 335-8, 335-10, 335-12 affectation) : 3,48 pt ; les modeles traces en
# filet de 0,48 pt (333-1 a 333-3, 334-3, 334-5, 334-7, 335-3 a 335-5) : 0,24 pt, le texte part du filet ; les 2
# tableaux des acomptes et le tableau des parts de 336-2 : 0,72 pt. Les tableaux de presentation libre (332-1, 334-3
# informations, 336-2 immeubles et autres actifs, 336-3), traces en filet de 0,96 pt, prennent la marge de cette
# famille. Word compte la marge depuis l'axe du filet (rendu Word du 05/10/2026 : 45 twips donnent 2,28 pt).
MARGE_PLEIN, MARGE_FIN, MARGE_ACOMPTES = 70, 5, 14
# MARGE DES MONTANTS, retenue le 06/10/2026 : 3,5 pt (70 twips, la marge
# des modeles traces en filet de 0,96 pt) a droite des cellules de montant des 12 tableaux dont le modele a une marge de
# cellule de 0,24 a 0,72 pt (333-1, 333-2, 333-3, 334-3, 334-5, 334-7, 335-3, 335-4, 335-5, les 2 tableaux des acomptes
# de 335-12, le tableau des parts de 336-2) ; 12 ecarts au modele acceptes. Le recueil n'imprime aucun montant : sa
# marge, relevee sur les libelles, reste celle du modele a gauche de toute cellule et a droite des libelles. None
# rendrait la marge du modele des 2 cotes.
MARGE_MONTANTS = MARGE_PLEIN
# PROPORTIONS DES COLONNES, retenues le 06/10/2026 : celles du modele, fixes, dans la justification ; ce
# qui deborde passe sur 2 lignes (un montant se coupe a l'espace d'un groupe de chiffres, _coupable). Un tableau de
# presentation libre (aucun modele dessine au recueil : carte « libre ») regle ses colonnes de montants sur leur
# contenu (_allouer).
GARDE_ENTIER = 15                # un tableau court, d'au plus 15 lignes en-tete compris, ne se coupe pas entre 2 pages


def _depuis_cadre(points, marge):
    u"""Le retrait (w:ind, twips) d'un libelle que le recueil place a « points » de l'axe du filet du cadre."""
    return int(round(points * 20)) - marge

# LA LETTRE DU RECUEIL, coquilles comprises. Le recueil ecrit l'apostrophe courbe, sauf dans les formes qui suivent,
# ou elle est droite ; il n'accentue pas la capitale des formes qui suivent (« Evolution », « Ecart », « Echéance »,
# « Etat »), et l'accentue ailleurs (« Évolution des parts et actions » de 334-5, « À l’exception »). Formes relevees
# mot a mot sur le texte du recueil, pages 88 a 107 du fichier.
FORMES_DU_RECUEIL = (
    # apostrophe droite
    u"Charges d'emprunt", u"Résultat de l'exercice (C=", u"sol d'autrui", u"faisant l'objet de dépréciations",
    u"« Autres dettes d'exploitation »", u"Date d'acquisition des parts", u"Secteur d'activité et localisation",
    u"^Différence d'estimation$", u"au sens de l'article R214-83", u"secteur d'activité, localisation",
    u"afférentes à l'exercice", u"à plus d'un an", u"prix d'achat résiduel", u"nature de l'engagement",
    u"montant de l'engagement", u"des frais d'acquisition en frais exclus",
    u"la méthode d'évaluation et de comptabilisation des dépréciations",
    # capitale sans accent
    u'Section 3 - Evolution', u'Evolution du nombre', u'Evolution de l’actif net', u'Evolution des comptes courants',
    u'Evolution des dépréciations', u'Ecart résultant', u'Echéance', u'Etat et autres collectivités')


def _motif_forme(forme):
    u"""Le motif d'une forme du recueil : la meme suite de caracteres, l'apostrophe droite ou courbe, la capitale
    accentuee ou non ; une forme entre ^ et $ est le texte entier (l'intitule de colonne du tableau des parts de
    336-2, « Différence d'estimation », que 333-3 ecrit avec l'apostrophe courbe)."""
    ancre = forme.startswith(u'^') and forme.endswith(u'$')
    m = u''
    for c in (forme[1:-1] if ancre else forme):
        m += {u"'": u"['’]", u'’': u"['’]", u'E': u'[EÉ]', u'É': u'[EÉ]'}.get(c) or re.escape(c)
    return re.compile(u'^%s$' % m if ancre else m)


_FORMES = [(_motif_forme(f), f[1:-1] if f.startswith(u'^') else f) for f in FORMES_DU_RECUEIL]


def lettre(t):
    u"""Un titre, un intitule, un libelle ou un texte du recueil, a sa lettre : l'apostrophe courbe, sauf dans les
    formes du recueil qui la portent droite ; la capitale telle que le recueil l'ecrit. Jamais applique a une valeur
    saisie ou lue."""
    if not t:
        return t
    t = t.replace(u"'", u'’')
    for motif, forme in _FORMES:
        t = motif.sub(lambda m, f=forme: f, t)
    return t


def typo(t):
    u"""Un texte qui ne vient pas du recueil (nom du livrable sur la page de garde) : l'apostrophe
    typographique et la capitale accentuee. Jamais applique a un texte du recueil, qui garde sa lettre."""
    if not t:
        return t
    return re.sub(u'\\bE(?=(?:volution|cart|chéance|tat)\\b)', u'É', t.replace(u"'", u'’'))


# Intitules des articles de l'annexe, au recueil (pages 96 a 107 du fichier), coquilles comprises (« 3 du I ... °du »
# de 334-5, « Exposition liées » de 334-8).
TITRES_ARTICLES = {
    u'331-1': u'Contenu de l’annexe', u'332-1': u'Caractéristiques et activité de l’OPCI',
    u'332-2': u'Règles et méthodes comptables', u'333-1': u'Décomposition des capitaux propres',
    u'333-2': u'Evolution du nombre de parts ou actions au cours de l’exercice',
    u'333-3': u'Evolution de l’actif net (= capitaux propres)',
    u'334-1': u'Informations relatives aux expositions immobilières',
    u'334-2': u'Immeubles en cours, construits ou acquis et droits réels', u'334-3': u'Contrats de crédit-bail',
    u'334-5': u'Parts et actions des sociétés, visées aux 2° et 3 du I de l’article L214-36 °du code monétaire et '
              u'financier',
    u'334-6': u'Evolution des comptes courants', u'334-7': u'Résultat sur plus et moins-values',
    u'334-8': u'Exposition liées aux actifs financiers d’OPCI', u'335-1': u'Principe de présentation',
    u'335-2': u'Décomposition des créances', u'335-3': u'Détail des provisions pour risques et charges',
    u'335-4': u'Détail des emprunts', u'335-5': u'Ventilation des emprunts par maturité résiduelle',
    u'335-6': u'Décomposition des dettes', u'335-7': u'Produits et charges de l’activité immobilière',
    u'335-8': u'Produits et charges sur opérations financières', u'335-9': u'Autres produits et autres charges',
    u'335-10': u'Frais de gestion et frais de fonctionnement externes', u'335-11': u'Engagements reçus et donnés',
    u'335-12': u'Tableau d’affectation du résultat', u'336-1': u'Principe de présentation',
    u'336-2': u'Inventaire détaillé des actifs à caractère immobilier',
    u'336-3': u'Inventaire détaillé des dépôts et des instruments financiers non immobiliers'}

# ARTICLES DE PRINCIPE (331-1, 334-1, 335-1, 336-1) : le texte de l'article, au mot pres, LU EN BASE depuis le
# 10/10/2026, une seule source pour le Word et l'ecran : ref_tableau_annexe.principe le designe, texte_principe le
# porte, un alinea par ligne, separes par NCHAR(10) (63_SQL/288_articles_de_principe.sql, texte verifie sur le recueil,
# pages 96, 99, 101 et 106 du fichier) ; lire_comptes_annuels le coupe en d['principes']. Aucune information n'y est
# demandee ; une saisie que la base porterait sur leurs lignes s'imprime sous le texte.


# 336-2, « pour les autres actifs à caractère immobilier » : chaque poste restant de l'actif immobilier de 321-2 recoit
# son inventaire (lecture du texte retenue le 05/10/2026), dans l'ordre du bilan, sous son intitule. Une ligne AUT: de
# la table porte son poste dans sa colonne groupe ; un poste sans actif garde une ligne du modele, « Néant ».
POSTES_AUTRES_IMMO = [u'A_IMMO_2', u'A_IMMO_5', u'A_IMMO_6', u'A_IMMO_7', u'A_IMMO_8']
NEANT = u'Néant'                  # ligne du modele gardee sans occurrence (une ligne par actif, et aucun actif)
# 336-3 : la ligne de la base ne porte pas son poste ; elle s'inscrit sous la rubrique de 321-2.
RUBRIQUE_336_3 = u'Dépôts et instruments financiers non immobiliers'

R9, R18 = 198, 360               # retraits du modele, en twips : 9,9 pt (niveau d'une ligne), 18 pt (detail des modeles)
R_ETATS = 180                    # detail des 3 etats : 12,48 pt du cadre au recueil, soit 3,48 pt de marge et 9 pt
# FILETS DU MODELE, epaisseur relevee au recueil : PLEIN 0,96 pt (w:sz 8, 1 pt), FIN 0,48 pt (w:sz 4, 0,5 pt),
# POINTILLE 0,48 pt. Une carte « fin » trace en FIN tout son tableau, cadre et verticaux compris.
PLEIN, FIN, POINTILLE = u'plein', u'fin', u'pointille'


def _filets(codes, style=PLEIN):
    return dict((c, dict(filet=style)) for c in codes)


def _maj(*dicts):
    r = {}
    for x in dicts:
        for k, v in x.items():
            r[k] = dict(r.get(k) or {}, **v)
    return r


# LES ETATS (321-2, 321-6, 322-2) : proportions des colonnes au recueil, retrait du detail, filets par bloc, hauteur
# de l'en-tete relevee entre les filets du recueil (pages 90 a 94 du fichier), en points (h_entete), et le pas de
# chaque rangee du corps (BASES_MODELES).
_SANS_FILET_RESULTAT = (u'R_PI', u'R_PI_1', u'R_PI_2', u'R_CI', u'R_CI_1', u'R_CI_2', u'R_CI_3', u'R_PF', u'R_PF_1',
                        u'R_CF', u'R_CF_1', u'R_PV', u'R_PV_1', u'R_MV_1')
# LE PAS DES LIGNES DES MODELES, rangee par rangee. Au recueil, le pas des lignes varie d'un bloc a l'autre (rubriques
# plus espacees, lignes de detail serrees a 10,9 pt dans certains blocs) et le texte n'est pas toujours centre dans sa
# rangee. Releve au recueil (pages 90 a 106 du fichier, PyMuPDF) pour chaque modele de structure fixe : la ligne de
# base du libelle de chaque rangee, en points depuis le filet sous l'en-tete ; le filet qui la separe de la suivante,
# s'il y en a un (position, epaisseur) ; la position du filet qui ferme le corps ; l'epaisseur du filet sous l'en-tete.
# Un tableau coupe entre 2 pages du recueil se raccorde au bas de sa derniere rangee.
BASES_MODELES = {
    u'BILAN_ACTIF': ([(u'A_IMMO', 11.04), (u'A_IMMO_1', 23.76), (u'A_IMMO_2', 36.6), (u'A_IMMO_3', 49.44),
                      (u'A_IMMO_4', 62.16), (u'A_IMMO_5', 75.0), (u'A_IMMO_6', 86.88), (u'A_IMMO_7', 98.64),
                      (u'A_IMMO_8', 111.48), (u'A_FIN', 124.56), (u'A_FIN_1', 137.52), (u'A_FIN_2', 150.36),
                      (u'A_FIN_3', 163.2), (u'A_FIN_4', 175.92), (u'A_FIN_5', 188.76), (u'A_FIN_6', 201.6),
                      (u'A_FIN_7', 214.32), (u'A_AUTRE', 229.2), (u'A_AUTRE_1', 243.96), (u'A_AUTRE_2', 256.8),
                      (u'A_AUTRE_3', 269.64)], 272.40, 0.96),
    u'BILAN_PASSIF': ([(u'P_CP', 14.76), (u'P_CP_1', 31.08), (u'P_CP_2', 43.92), (u'P_CP_3', 56.64),
                       (u'P_CP_4', 69.36), (u'P_CP_5', 82.08), (u'P_PROV', 101.04), (u'P_IMMO', 123.48),
                       (u'P_IMMO_1', 138.96), (u'P_IMMO_2', 149.88), (u'P_IMMO_3', 160.68), (u'P_IMMO_4', 171.6),
                       (u'P_IF', 187.08), (u'P_IF_1', 203.52), (u'P_IF_2', 216.24), (u'P_IF_3', 228.96),
                       (u'P_DET', 245.4), (u'P_DET_1', 260.88), (u'P_DET_2', 271.8), (u'P_DET_3', 282.6),
                       (u'P_DET_4', 294.48), (u'P_DET_5', 307.2)], 309.96, 0.96),
    u'RESULTAT': ([
        (u'R_PI', 12.12), (u'R_PI_1', 27.12), (u'R_PI_2', 42.12), (u'R_PI_3', 57.48, (61.8, 0.96)),
        (u'R_TI', 74.28, (78.48, 0.96)), (u'R_CI', 90.6), (u'R_CI_1', 105.6), (u'R_CI_2', 120.6), (u'R_CI_3', 135.6),
        (u'R_CI_4', 151.08, (155.28, 0.96)), (u'R_TII', 167.76, (171.96, 0.96)), (u'R_RAI', 184.56, (188.76, 0.96)),
        (u'R_PF', 200.88), (u'R_PF_1', 215.88), (u'R_PF_2', 231.24, (235.56, 0.96)),
        (u'R_TIII', 248.04, (252.24, 0.96)), (u'R_CF', 264.36), (u'R_CF_1', 279.36),
        (u'R_CF_2', 294.72, (299.04, 0.96)), (u'R_TIV', 311.52, (315.72, 0.96)), (u'R_ROF', 328.32, (332.52, 0.96)),
        (u'R_TV', 345.0, (349.32, 0.96)), (u'R_TVI', 361.8, (366.0, 0.96)), (u'R_TVII', 378.48, (382.8, 0.96)),
        (u'R_RN', 395.28, (399.48, 0.96)), (u'R_TVIII', 412.08, (416.28, 0.96)), (u'R_RNAR', 432.96, (441.24, 0.96)),
        (u'R_PV', 453.36), (u'R_PV_1', 468.36), (u'R_PV_2', 483.72, (488.04, 0.96)),
        (u'R_TIX', 500.52, (504.72, 0.96)), (u'R_MV', 517.32, (521.52, 0.96)), (u'R_MV_1', 533.64),
        (u'R_MV_2', 549.0, (553.32, 0.96)), (u'R_TX', 565.8, (570.0, 0.96)), (u'R_RPMV', 582.48, (586.8, 0.96)),
        (u'R_TXI', 599.28, (603.48, 0.96)), (u'R_PMVAR', 620.16, (628.62, 0.96)), (u'R_RE', 642.24)],
        648.72, 0.96),
    (u'333-1', 0): ([
        (u'CAPITAL', 9.48, (14.28, 0.48)), (u'SD', 23.76, (28.44, 0.48)), (u'SD_REPORT_PMV', 38.88, (43.8, 0.48)),
        (u'SD_REG_PMV', 54.24, (59.16, 0.48)), (u'SD_REPORT_RES', 69.6, (74.52, 0.48)),
        (u'SD_REG_RES', 84.96, (90.0, 0.48)), (u'SD_RES', 100.44, (105.48, 0.48)),
        (u'SD_RES_NET', 115.92, (120.96, 0.48)), (u'SD_RES_PMV', 131.4, (136.32, 0.48)),
        (u'SD_AC', 146.88, (151.68, 0.48)), (u'SD_AC_NET', 162.24, (167.04, 0.48)),
        (u'SD_AC_PMV', 177.6, (182.52, 0.48)), (u'SD_REG_AC', 192.96, (197.88, 0.48)),
        (u'SD_REG_AC_NET', 208.32, (213.24, 0.48)), (u'SD_REG_AC_PMV', 223.68, (228.6, 0.48)),
        (u'SD_TOTAL', 238.08, (242.76, 0.48)), (u'CP_TOTAL', 253.2)],
        257.04, 0.48),
    (u'333-2', 0): ([
        (u'SOUSCRIPTIONS', 10.56, (14.4, 0.48)), (u'RACHATS', 24.84, (28.44, 0.48)), (u'MONTANTS_NETS', 39.6)],
        44.16, 0.48),
    (u'333-3', 0): ([
        (u'AN_DEBUT', 11.52), (u'SOUSCRIPTIONS', 27.48), (u'RACHATS', 52.8), (u'FRAIS_ACQ', 68.28),
        (u'FRAIS_ACQ_N', 81.0), (u'FRAIS_ACQ_RES', 93.72), (u'CHANGE', 107.4), (u'VAR_DE_IMMO', 121.2),
        (u'DE_IMMO_N', 134.04), (u'DE_IMMO_N1', 145.32), (u'VAR_DE_FIN', 157.44), (u'DE_FIN_N', 182.04),
        (u'DE_FIN_N1', 193.2), (u'DISTRIB_PREC', 205.32), (u'RESULTAT', 219.0), (u'RES_NET', 231.96),
        (u'RES_PMV', 243.84), (u'ACOMPTES', 256.44), (u'AC_RES_NET', 270.36), (u'AC_CESSIONS', 282.24),
        (u'FUSION', 296.76), (u'AUTRES', 308.64), (u'AN_FIN', 322.44)],
        328.32, 0.48),
    (u'334-2', 0): ([
        (u'TERRAINS_NUS', 11.76), (u'TERRAINS_CONSTR', 26.04), (u'SOL_AUTRUI', 40.32), (u'DROITS_REELS', 54.48),
        (u'EN_COURS', 68.76), (u'AUTRES', 83.28, (87.12, 0.96)), (u'TOTAL', 98.64, (101.16, 0.48)),
        (u'TOTAL_HORS_FRAIS', 111.96)],
        114.24, 0.96),
    (u'334-5', 0): ([
        (u'PARTS_2', 10.56), (u'PARTS_3', 33.24, (46.92, 0.48)), (u'TOTAL', 56.64, (60.6, 0.48)),
        (u'TOTAL_HORS_FRAIS', 70.2)],
        79.56, 0.48),
    (u'334-7', 0): ([
        (u'TERRAINS_NUS', 10.56), (u'TERRAINS_CONSTR', 22.44), (u'SOL_AUTRUI', 34.32), (u'DROITS_REELS', 46.2),
        (u'EN_COURS', 58.08), (u'AUTRES', 69.96, (72.72, 0.48)), (u'TOTAL_IMMEUBLES', 82.2, (92.88, 0.48)),
        (u'PARTS_2', 103.44), (u'PARTS_3', 115.44), (u'ACTIONS_4', 127.2), (u'OPCI_5', 139.08),
        (u'AUTRES_IMMO', 161.88, (164.64, 0.48)), (u'TOTAL_AUTRES', 174.12, (184.8, 0.48)),
        (u'TOTAL_GENERAL', 198.72)],
        205.68, 0.48),
    (u'335-2', 0): ([
        (u'CL', 9.96), (u'CL_1', 22.92), (u'CL_2', 35.76), (u'CL_3', 48.6), (u'CL_ST', 63.6, (68.64, 0.96)),
        (u'AC', 79.68), (u'AC_1', 92.52), (u'AC_2', 105.36), (u'AC_3', 118.2), (u'AC_4', 131.04), (u'AC_5', 143.88),
        (u'AC_6', 156.72), (u'AC_ST', 171.84, (176.76, 0.96)), (u'TOTAL', 190.08)],
        195.12, 0.96),
    (u'335-2', 1): ([
        (u'DEP_CL', 10.8), (u'DEP_AC', 23.16, (25.8, 0.96)), (u'DEP_TOTAL', 38.28)],
        42.6, 0.96),
    # 335-3 : la ligne « … » du modele porte la ligne des provisions (page 102 du fichier)
    (u'335-3', 0): ([
        (u'PROVISIONS', 9.98), (u'TOTAL', 51.5)],
        54.4, 0.48),
    (u'335-4', 0): ([
        (u'EC', 10.2), (u'EC_IMMO', 27.0), (u'EC_NON_IMMO', 39.0, (46.2, 0.48)), (u'AUTRES', 56.4),
        (u'AUTRES_IMMO', 73.2), (u'AUTRES_NON_IMMO', 90.0, (97.32, 0.48)), (u'TOTAL', 107.4)],
        114.72, 0.48),
    (u'335-5', 0): ([
        (u'TF', 10.56), (u'TF_AMORT', 23.4), (u'TF_IN_FINE', 35.28, (38.04, 0.48)), (u'TV', 48.0),
        (u'TV_AMORT', 60.0), (u'TV_IN_FINE', 71.88, (74.52, 0.48)), (u'TOTAL', 85.08)],
        88.8, 0.48),
    (u'335-6', 0): ([
        (u'D_1', 11.28, (14.52, 0.48)), (u'D_2', 25.08, (28.32, 0.48)), (u'D_3', 38.88, (42.12, 0.48)),
        (u'D_4', 52.68, (55.92, 0.48)), (u'D_5', 66.84, (70.2, 0.96)), (u'TOTAL', 81.72)],
        85.2, 0.96),
    (u'335-8', 0): ([
        (u'P', 12.12), (u'P_1', 27.0, (31.08, 0.48)), (u'P_2', 42.48, (46.56, 0.48)), (u'P_3', 57.96, (62.04, 0.48)),
        (u'P_4', 73.44, (77.52, 0.48)), (u'P_5', 88.8, (92.88, 0.48)), (u'P_6', 104.28, (108.36, 0.48)),
        (u'P_7', 119.76, (123.84, 0.48)), (u'P_8', 135.24, (139.32, 0.48)), (u'P_9', 150.6, (154.68, 0.48)),
        (u'P_ST', 166.44, (170.76, 0.96)), (u'C', 182.88), (u'C_1', 197.76, (201.84, 0.48)),
        (u'C_2', 213.24, (217.32, 0.48)), (u'C_3', 228.72, (232.8, 0.48)), (u'C_4', 244.08, (248.16, 0.48)),
        (u'C_5', 259.56, (263.64, 0.48)), (u'C_6', 275.04, (279.12, 0.48)), (u'C_7', 290.52, (294.6, 0.48)),
        (u'C_ST', 306.36, (310.56, 0.96)), (u'TOTAL', 325.32)],
        331.8, 0.96),
    (u'335-10', 0): ([
        (u'F_1', 13.92, (19.92, 0.48)), (u'F_2', 33.12, (39.12, 0.48)), (u'F_3', 52.32, (58.32, 0.48)),
        (u'F_4', 71.52, (77.52, 0.48)), (u'F_5', 90.72, (96.48, 0.48)), (u'F_6', 109.92, (115.68, 0.48)),
        (u'TOTAL', 132.96)],
        142.44, 0.96),
    (u'335-12', 0): ([
        (u'I', 10.08), (u'I_D', 28.92), (u'I_D_1', 47.88), (u'I_D_2', 58.8), (u'I_D_3', 69.71), (u'I_D_T', 80.51),
        (u'I_A', 102.23), (u'I_A_1', 113.15), (u'I_A_2', 124.06), (u'I_A_3', 134.98),
        (u'I_A_T', 145.8, (147.6, 0.48)), (u'II', 157.2), (u'II_D', 193.2), (u'II_D_1', 205.2), (u'II_D_2', 217.2),
        (u'II_D_3', 229.2), (u'II_D_T', 241.2), (u'II_A', 261.36), (u'II_A_1', 272.27), (u'II_A_2', 283.19),
        (u'II_A_3', 296.99), (u'II_A_T', 307.8)],
        310.92, 0.96),
}
# Du haut d'un paragraphe de 9 pt en Arial, interligne simple, a sa ligne de base, en points (rendu Word du 06/10/2026)
B0_ARIAL_9 = 8.47


def _b0(forme=None):
    u"""Du haut du paragraphe du libelle a sa ligne de base, en points, rendu Word du 06/10/2026 : 8,47 pt en 9 pt,
    au prorata du corps (contenu des tableaux en 10 pt) ; a interligne exact, Word pose la ligne de base aux 4/5 de
    l'interligne ; la puce en police Symbol abaisse la ligne de 0,65 pt."""
    forme = forme or {}
    taille = forme.get(u'taille_libelle') or forme.get(u'taille') or CORPS_TABLEAU
    if forme.get(u'interligne'):
        b = 0.8 * _interligne(forme[u'interligne'], taille)
    else:
        b = B0_ARIAL_9 * taille / 9.0
    return b + (0.65 if forme.get(u'puce') else 0.0)


# LES LIBELLES SUR 2 LIGNES AU RECUEIL, et l'interligne qui les separe (releve au recueil) : la rangee du modele
# porte ces 2 lignes ; quand le libelle tient sur une ligne au Word, la
# rangee perd cet interligne (aucune bande blanche que le modele n'a pas).
LIBELLES_2_LIGNES = {
    (u'333-3', 0): {u'SOUSCRIPTIONS': 11.6, u'VAR_DE_FIN': 11.7},
    (u'334-5', 0): {u'PARTS_2': 10.9, u'PARTS_3': 11.0},
    (u'334-7', 0): {u'OPCI_5': 10.9},
    (u'335-12', 0): {u'II': 12.0}}


def _hauteur_de_ligne(forme=None):
    u"""La hauteur d'une ligne de texte du libelle au Word, en points : l'interligne exact du modele, au moins
    l'interligne simple du corps (11,5 pt en 10 pt) ; 12,2 pt pour une ligne a puce en police Symbol (rendu Word du
    06/10/2026)."""
    forme = forme or {}
    taille = forme.get(u'taille_libelle') or forme.get(u'taille') or CORPS_TABLEAU
    h = _interligne(forme.get(u'interligne'), taille) or INTERLIGNE_ARIAL * taille
    return max(h, 1.22 * taille) if forme.get(u'puce') else h


def _au_pas_du_recueil(bases, fin, ep_entete, formes=None):
    u"""Les attributs de rangee qui posent chaque libelle sur la ligne de base relevee au recueil : texte cale en haut
    de sa rangee, un espace avant qui place sa ligne de base, une hauteur qui mene au haut de la rangee suivante (la
    ligne de base suivante, ou le filet qui les separe), la derniere jusqu'au filet qui ferme le corps. L'epaisseur du
    filet au-dessus d'une rangee se deduit de sa hauteur dans _modele : elle se compte ici dans la rangee. formes :
    corps, interligne et puce du libelle de chaque rangee (_b0) ; ligne_blanche : le modele met une ligne blanche
    avant la rangee (335-12) ; pas_2_lignes : la rangee porte un libelle de 2 lignes au recueil."""
    formes = formes or {}
    r = {}
    haut, ep = 0.0, (0.5 if ep_entete < 0.7 else 1.0)
    for k, x in enumerate(bases):
        code, b = x[0], x[1]
        dessous = x[2] if len(x) > 2 else None
        avant = max(0.0, b - haut - ep - _b0(formes.get(code)))
        if dessous is not None:
            suivant, ep_suivant = dessous[0], (0.5 if dessous[1] < 0.7 else 1.0)
        elif k + 1 < len(bases):
            # LA LIGNE BLANCHE AVANT LA RANGEE SUIVANTE est l'espace avant de celle-ci, et non le bas de celle-ci : un
            # libelle qui passe sur 2 lignes au Word ne la mange pas. Elle vaut l'ecart des 2 lignes de base du
            # recueil, moins la hauteur de ligne de la rangee au Word (et la 2e ligne d'un libelle de 2 lignes)
            f = formes.get(bases[k + 1][0]) or {}
            blanche = 0.0
            if f.get(u'ligne_blanche'):
                fc = formes.get(code) or {}
                blanche = max(0.0, bases[k + 1][1] - b - _hauteur_de_ligne(fc) - (fc.get(u'pas_2_lignes') or 0))
            suivant, ep_suivant = bases[k + 1][1] - _b0(f) - blanche, 0.0
        else:
            suivant, ep_suivant = fin, 0.0
        r[code] = dict(va=u'top', espace_avant=round(avant, 2), espace_apres=0, h=round(suivant - haut, 2))
        haut, ep = suivant, ep_suivant
    return r


CARTES_ETATS = {
    # « OPCI et organismes étrangers équivalents… » part a 13,68 pt du cadre au recueil, 1,2 pt plus loin que les autres
    # lignes de detail (12,48 pt, page 90 du fichier)
    u'BILAN_ACTIF': dict(largeurs=[69.68, 13.65, 16.67], exercices=(u'Exercice N', u'Exercice N-1'), retrait=R_ETATS,
                         h_entete=26.3, pas=12.8,
                         lignes=_maj(_filets([u'A_AUTRE_3']), {u'A_TOTAL': dict(h=21.2)},
                                     {u'A_IMMO_6': dict(retrait=_depuis_cadre(13.68, MARGE_PLEIN))},
                                     _au_pas_du_recueil(*BASES_MODELES[u'BILAN_ACTIF']))),
    u'BILAN_PASSIF': dict(largeurs=[71.2, 13.2, 15.6], exercices=(u'Exercice N', u'Exercice N-1'), retrait=R_ETATS,
                          h_entete=28.0, pas=12.7,
                          lignes=_maj(_filets([u'P_DET_5']), {u'P_TOTAL': dict(h=21.1)},
                                      _au_pas_du_recueil(*BASES_MODELES[u'BILAN_PASSIF']))),
    u'RESULTAT': dict(largeurs=[81.0, 9.5, 9.5], exercices=(u'Exercice\nN', u'Exercice\nN-1'), retrait=0,
                      h_entete=25.9, pas=15.4, interligne_entete=10.9,
                      lignes=_maj(_filets([c for c, _, _ in MODELES_ETATS[u'RESULTAT']
                                           if c not in _SANS_FILET_RESULTAT]),
                                  _au_pas_du_recueil(*BASES_MODELES[u'RESULTAT'])))}


_TOTAL_FRAIS_EXCLUS = dict(dispo=[(0, 5, u'right'), (5, 2, u'right')], retrait=0)
_EXERCICE_N, _EXERCICE_N1 = u'Exercice\nN', u'Exercice\nN-1'
_VARIATION_DE = u'Variation des différences d’estimation'
_ENTETE_ACOMPTES = [[None, dict(t=u'Montants sur résultat net', span=2),
                     dict(t=u'Montants sur plus et moins-values nettes réalisées', span=2),
                     dict(t=u'Distributions réalisées', span=2)],
                    [dict(t=u'Date de paiement', gras=False)] + [dict(t=t, gras=False) for t in
                                                                  (u'Montant\ntotal', u'Montant\nunitaire') * 3]]
# 2e ligne d'en-tete du 1er tableau des acomptes, coupures du recueil (page 106 du fichier) : « Montant total » sur une
# ligne dans la colonne des distributions (espace insecable : la colonne s'elargit a l'intitule)
_ENTETE_ACOMPTES_N = [_ENTETE_ACOMPTES[1][0]] + [dict(t=t, gras=False) for t in (
    u'Montant\ntotal', u'Montant\nunitaire', u'Montant\ntotal', u'Montant\nunitaire', u'Montant\u00a0total',
    u'Montant\nunitaire')]

# LES MODELES DE L'ANNEXE, par (article, rang du bloc dans ANNEXE). largeurs : proportions des colonnes au recueil,
# mesurees entre ses filets verticaux (PyMuPDF, get_drawings ; ecart de 0,1 point au plus, releve du 06/10/2026), fixes
# dans la justification ; libre : tableau de presentation libre, sans modele dessine (les colonnes de montants
# s'elargissent a leur contenu, la 1re prend le reste) ; titre_modele : la 1re cellule d'en-tete porte le titre du
# modele (12 pt, gras, bleu fonce), et non l'intitule de la colonne des libelles (« Ventilation par… ») ;
# marge : marge de cellule du modele en twips (MARGE_PLEIN par defaut, MARGE_FIN pour une carte « fin ») ; entete : lignes
# d'en-tete (texte, ou dict t, span, jc, va, gras, taille) ; jc0 : alignement de la 1re cellule d'en-tete ; fin : tout
# le tableau en filets de 0,48 pt ; cadre_bas : filet du bas ; h_entete, pas :
# hauteur de l'en-tete et des lignes, en points, relevee entre les filets du recueil (pages 98 a 107 du fichier) ;
# defaut et lignes : attributs de ligne (retrait en twips, filet sous la ligne et filet_cols : colonnes qu'il couvre,
# h et va : hauteur et alignement vertical, jc0 du libelle, italique, couleur, taille, espace_avant, espace_apres,
# puce, affiche : libelle du recueil, dispo : (cellule, colonnes couvertes, alignement), avant et apres : lignes du
# modele imprimees avant et apres celle-ci ; intitule : une rangee d'intitule que le modele ajoute, ses autres
# cellules par la regle commune) ; forme texte : les informations redigees, en paragraphes a accroche (lie_au_suivant :
# la valeur reste avec ce qui la suit) ; avant_bloc : le rang du bloc avant lequel celui-ci s'imprime
# (_ordre_impression).
CARTES_ANNEXE = {
    # LES INFORMATIONS REDIGEES (forme texte) : le texte du recueil qui exige l'information, a la lettre, puis sa
    # valeur. Une phrase qui exige une information : la phrase (intro), puis la valeur seule (sans_accroche). Une
    # enumeration : sa phrase d'introduction, puis une puce par element, l'element en gras, sa ponctuation finale
    # remplacee par « : » et la valeur (puce) ; un element sans valeur propre en puce maigre (puces). L'ordre est celui
    # du recueil.
    (u'332-1', 0): dict(forme=u'texte', lignes={u'STRATEGIE': dict(sans_accroche=True, avant=[dict(
        intro=u'Selon les informations du prospectus, la stratégie d’investissement de l’OPCI est rappelée.')])}),
    (u'332-1', 1): dict(largeurs=[34.0, 13.2, 13.2, 13.2, 13.2, 13.2], libre=True, titre_au_dessus=True, pas=12.8,
                        defaut=dict(retrait=0),
                        egales=[(1, 2, 3, 4, 5)],
                        lignes={
        u'ACTIF_NET': dict(affiche=u'actif net'), u'NB_PARTS': dict(affiche=u'nombre de parts'),
        u'VL_UNIT': dict(affiche=u'valeur liquidative unitaire'),
        u'DIST_CESSION': dict(affiche=u'distribution unitaire sur résultat de cession (y compris les acomptes)'),
        u'DIST_RESULTAT': dict(affiche=u'distribution unitaire sur résultat net (y compris les acomptes)'),
        u'CAPITALISATION': dict(affiche=u'capitalisation unitaire')}),
    (u'332-2', 0): dict(forme=u'texte', lignes={
        MENTIONS_2020_07: dict(sans_accroche=True, avant=[dict(
            intro=u'Outre les mentions relatives aux règles et méthodes comptables prévues par le règlement ANC n° '
                  u'2020-07 relatif aux comptes annuels des organismes de placement collectif à capital variable, '
                  u'l’annexe comporte les informations relatives aux règles et méthodes comptables spécifiques aux '
                  u'OPCI en raison de leurs actifs immobiliers, dès lors qu’elles sont significatives.')]),
        u'VALORISATION': dict(puce=True, avant=[dict(
            intro=u'Les modalités d’application des méthodes comptables sont détaillées dans l’annexe et, si '
                  u'nécessaire, justifiées. Sont notamment concernés :')],
            affiche=u'le mode de valorisation des actifs à caractère immobilier, qui distingue les immeubles en cours de '
                    u'construction, acquis en vue de leur restructuration ou faisant l’objet d’une restructuration, les '
                    u'immeubles construits et acquis, les droits réels et les parts et actions des sociétés ou '
                    u'organismes ainsi que les comptes courants. À l’exception des actifs à caractère immobilier dont '
                    u'la valeur actuelle est issue d’une cotation accessible, les modalités de détermination de la '
                    u'valeur actuelle doivent être précisées en distinguant notamment les actifs à caractère immobilier '
                    u'que l’OPCI s’est engagé à céder. Une attention particulière sera enfin apportée, le cas échéant, '
                    u'à la description des modes de valorisation faisant appel à des calculs ou à des modèles '
                    u'financiers utilisant des calculs ou des paramètres faisant place à l’anticipation'),
        u'FRAIS_ACQ': dict(puce=True, affiche=u'les modalités d’enregistrement des frais d’acquisition en frais exclus, '
                                              u'détaillées par nature d’actifs'),
        u'COUTS_ULTERIEURS': dict(puce=True, affiche=u'le mode de comptabilisation des coûts et dépenses ultérieurs'),
        u'COMMISSIONS': dict(puce=True, affiche=u'le mode de comptabilisation des commissions de souscription, en '
                                                u'distinguant la part fixe de la part variable'),
        u'DEPRECIATIONS': dict(puce=True, affiche=u'la méthode d’évaluation et de comptabilisation des dépréciations sur '
                                                  u'créances locatives'),
        u'PROVISIONS': dict(puce=True, affiche=u'la méthode d’évaluation et de comptabilisation des provisions pour '
                                               u'risques et charges'),
        u'CREDIT_BAIL': dict(puce=True, affiche=u'la méthode d’évaluation des contrats de crédit-bail'),
        u'RENTES_VIAGERES': dict(puce=True, affiche=u'la méthode d’évaluation des biens acquis moyennant paiement de '
                                                    u'rentes viagères et de la dette associée')}),
    (u'333-1', 0): dict(largeurs=[64.34, 17.05, 18.61], titre_modele=True, jc0=u'center', quadrillage=True, fin=True,
                        h_entete=20.1,
                        pas=15.4, defaut=dict(retrait=0), lignes=_maj(
        # detail du resultat a 11,40 pt du cadre, detail des acomptes et de leur regularisation a 14,16 pt (page 98)
        dict((c, dict(retrait=_depuis_cadre(11.40, MARGE_FIN))) for c in (u'SD_RES_NET', u'SD_RES_PMV')),
        dict((c, dict(retrait=_depuis_cadre(14.16, MARGE_FIN))) for c in (u'SD_AC_NET', u'SD_AC_PMV', u'SD_REG_AC_NET',
                                                                          u'SD_REG_AC_PMV')),
        dict((c, dict(h=14.2)) for c in (u'CAPITAL', u'SD', u'SD_TOTAL', u'CP_TOTAL')))),
    (u'333-2', 0): dict(largeurs=[31.6, 22.3, 30.1, 15.9], titre_modele=True, jc0=u'center', fin=True, h_entete=29.8,
                        pas=14.2,
                        interligne_entete=11.8, entete=[[
        u'Souscriptions et rachats de l’exercice', u'Nombre de parts ou actions',
        u'Montants bruts\n(hors frais et commissions)', u'Frais et commissions*']], defaut=dict(retrait=_depuis_cadre(18.24, MARGE_FIN)),
        lignes={u'SOUSCRIPTIONS': dict(filet=POINTILLE), u'RACHATS': dict(filet=PLEIN),
                u'MONTANTS_NETS': dict(h=15.8)}),
    # 333-3 : filets verticaux du recueil a x 48,2, 409,8, 452,3 et 501,8 pt (pages 98 et 99 du fichier, get_drawings) :
    # libelles et signes 79,7 %, Exercice N 9,4 %, Exercice N-1 10,9 % ; les signes, sans filet, centres a x 397,1 pt :
    # leur colonne 5,6 %, les libelles 74,1 %
    (u'333-3', 0): dict(largeurs=[74.1, 5.6, 9.4, 10.9], titre_modele=True, verticales=(0, 2, 3, 4), centrees=(1,),
                        fin=True,
                        h_entete=24.0, pas=12.7, va=u'top', interligne_entete=11.7, entete=[[
        dict(t=u'Evolution de l’actif net (= capitaux propres) au cours de l’exercice', span=2, va=u'top'),
        _EXERCICE_N, _EXERCICE_N1]], lignes=_maj(
        # frais d'acquisition a 7,32 pt du cadre, differences d'estimation a 8,76 pt (page 99)
        dict((c, dict(retrait=_depuis_cadre(7.32, MARGE_FIN))) for c in (u'FRAIS_ACQ_N', u'FRAIS_ACQ_RES')),
        dict((c, dict(retrait=_depuis_cadre(8.76, MARGE_FIN))) for c in (u'DE_IMMO_N', u'DE_IMMO_N1', u'DE_FIN_N',
                                                                         u'DE_FIN_N1')),
        dict((c, dict(retrait=R18)) for c in (u'RES_NET', u'RES_PMV', u'AC_RES_NET', u'AC_CESSIONS')),
        dict((c, dict(italique_libelle=True)) for c in (u'DE_IMMO_N', u'DE_IMMO_N1', u'DE_FIN_N', u'DE_FIN_N1')),
        {u'DE_IMMO_N1': dict(affiche=u'Différence d’estimation exercice N-1:'),
         u'DE_FIN_N1': dict(affiche=u'Différence d’estimation exercice N-1:'),
         u'ACOMPTES': dict(affiche=u'Acomptes versés au cours de l’exercice avant comptes de régularisations :')},
        # les 2 libelles sur 2 lignes, a l'interligne du recueil (page 98 et 99 du fichier : 11,6 pt)
        dict((c, dict(interligne=11.6)) for c in (u'SOUSCRIPTIONS', u'VAR_DE_FIN')))),
    # 333-3 : l'explication du renvoi 1, sous la note du modele qui l'appelle (page 99 du fichier)
    (u'333-3', 1): dict(forme=u'texte', lignes={u'RENVOI_1': dict(sans_accroche=True)}),
    (u'334-1', 0): dict(forme=u'texte'),
    # 334-2 : filet de 0,48 pt entre « Total » et « Total frais exclus », 0,96 pt ailleurs ; libelles des 2 totaux en
    # 10 pt au recueil
    (u'334-2', 0): dict(largeurs=[34.08, 9.11, 15.17, 12.49, 12.69, 8.95, 7.52],
                        h_entete=44.4, pas=14.2, interligne_entete=10.9, entete=[[
        dict(t=u'Ventilation par nature', jc=u'left'), u'Exercice\nN-1', u'Augmentations', u'Diminutions',
        _VARIATION_DE, _EXERCICE_N, u'Frais']], defaut=dict(retrait=R18), lignes={
        u'AUTRES': dict(filet=PLEIN), u'TOTAL': dict(retrait=0, jc0=u'center', taille_libelle=10, filet=FIN, h=14.1),
        u'TOTAL_HORS_FRAIS': dict(_TOTAL_FRAIS_EXCLUS, taille_libelle=10, h=13.1)}),
    # 334-3 : une ligne par contrat avant « Total », les lignes CB: de la table (un contrat, le constat « Néant », ou
    # le contrat a inventorier), chacune a la forme de la ligne du modele
    (u'334-3', 0): dict(largeurs=[16.9, 12.7, 8.4, 14.2, 12.7, 14.2, 11.3], jc0=u'center', fin=True, h_entete=35.6,
                        interligne_entete=11.8, entete=[[
        u'Ventilation par contrat', u'Assiette du contrat', u'Exercice\nN-1', u'Augmentations', u'Diminutions',
        _VARIATION_DE, _EXERCICE_N]], lignes={u'TOTAL': dict(jc0=u'center', h=20.3)},
                        occurrences=dict(avant=u'TOTAL', prefixe=u'CB:', forme=dict(retrait=R18, filet=PLEIN, h=16.3))),
    # 334-3, informations enumerees (presentation libre) : ventilees « selon les postes du bilan dont auraient releve
    # les biens concernes », soit le poste « Immeubles en cours, construits ou acquis et droits reels » de 321-2 pour
    # un immeuble pris en credit-bail ; celles du dernier alinea « selon les echeances a un an au plus, a plus d'un an
    # et cinq ans au plus et a plus de cinq ans » (recueil, page 100 du fichier). Libelles : les mots des puces du
    # recueil, a la lettre. Les 3 echeances du prix d'achat residuel se lisent sur leurs lignes de la table
    # (PRIX_RESIDUEL_1AN, PRIX_RESIDUEL_1_5ANS, PRIX_RESIDUEL_5ANS) ; une ligne absente s'imprime vide.
    # ORDRE DU RECUEIL (page 100 du fichier) : les informations enumerees, puis le tableau « Ventilation par contrat » ;
    # ce bloc s'imprime donc avant le bloc 0 (avant_bloc), son rang dans ANNEXE, qui designe sa carte, restant 1.
    (u'334-3', 1): dict(largeurs=[68.0, 32.0], libre=True, pas=12.4, avant_bloc=0,
                        entete=[[dict(t=u'Contrats de crédit-bail', jc=u'left'),
                                                        u'Immeubles en cours, construits ou acquis et droits réels']],
                        defaut=dict(retrait=0), lignes={
        u'ASSIETTE': dict(affiche=u'valeur des biens pris en crédit-bail au moment de la signature du contrat '
                                  u'(assiette d’origine du contrat)'),
        u'REDEVANCES_N': dict(affiche=u'montant des redevances afférentes à l’exercice'),
        u'REDEVANCES_CUM': dict(affiche=u'montant cumulé des redevances des exercices précédents'),
        u'RESTANT_1AN': dict(retrait=R18, affiche=u'à un an au plus', avant=[
            dict(intitule=u'évaluation des redevances restant à payer', garder=True)]),
        u'RESTANT_1_5ANS': dict(retrait=R18, affiche=u'à plus d’un an et cinq ans au plus'),
        u'RESTANT_5ANS': dict(retrait=R18, affiche=u'à plus de cinq ans'),
        u'PRIX_RESIDUEL': dict(affiche=u'prix d’achat résiduel de ces biens stipulé au contrat', garder=True,
                               suivantes=[u'PRIX_RESIDUEL_1AN', u'PRIX_RESIDUEL_1_5ANS', u'PRIX_RESIDUEL_5ANS']),
        u'PRIX_RESIDUEL_1AN': dict(retrait=R18, affiche=u'à un an au plus'),
        u'PRIX_RESIDUEL_1_5ANS': dict(retrait=R18, affiche=u'à plus d’un an et cinq ans au plus'),
        u'PRIX_RESIDUEL_5ANS': dict(retrait=R18, affiche=u'à plus de cinq ans')}),
    # 334-5 : le titre du modele tient sur 4 lignes dans la 1re cellule (page 100 du fichier)
    (u'334-5', 0): dict(largeurs=[31.22, 10.94, 10.94, 14.07, 15.62, 9.4, 7.81], titre_modele=True, jc0=u'center',
                        fin=True,
                        h_entete=47.4,
                        lignes_entete=4, interligne_entete=11.8,
                        entete=[[
        u'Évolution des parts et actions des sociétés visées à l’article L214-36 du code monétaire et financier '
        u'(2° et 3° du I)', u'Exercice\nN-1', u'Cessions', u'Acquisitions', _VARIATION_DE, _EXERCICE_N, u'Frais']],
        defaut=dict(retrait=0), lignes={u'PARTS_2': dict(h=23.4, interligne=10.9),
                                        u'PARTS_3': dict(filet=PLEIN, h=23.4, interligne=10.9),
                                        u'TOTAL': dict(jc0=u'center', filet=PLEIN, h=13.7),
                                        u'TOTAL_HORS_FRAIS': dict(_TOTAL_FRAIS_EXCLUS, h=19.0)}),
    (u'334-5', 1): dict(forme=u'texte', lignes={u'SECTEURS': dict(sans_accroche=True, avant=[dict(
        intro=u'La valeur actuelle des exercices N et N-1 des parts et actions relevant des 2° et 3° du I de l’article '
              u'L214-36 du code monétaire et financier est également communiquée par secteur d’activité et/ou par '
              u'secteur géographique.')])}),
    # 334-6 : une ligne par contrat, au pas du modele (12,8 pt), pointilles entre contrats (page 101 du fichier)
    (u'334-6', 0): dict(largeurs=[18.76, 12.5, 10.93, 17.2, 14.07, 15.63, 10.91], jc0=u'center', h_entete=33.6,
                        pas=12.8,
                        interligne_entete=10.8,
                        entete=[[
        u'Ventilation par contrat', u'Echéance', u'Exercice\nN-1', u'Augmentations', u'Diminutions', _VARIATION_DE,
        _EXERCICE_N]], lignes={u'TOTAL': dict(h=17.0)}),
    # 334-7 : les 2 lignes de sous-total portent la hauteur du bloc, texte en haut ; le bloc suivant part du filet
    (u'334-7', 0): dict(largeurs=[67.1, 7.8, 7.8, 7.8, 9.4], fin=True, h_entete=47.4, pas=11.9, lignes_entete=4,
                        interligne_entete=11.7,
                        entete=[[
        dict(t=u'Ventilation par nature', jc=u'left'), u'Plus-values', u'Moins-values', u'Plus et moins-values\nN',
        u'Plus et moins-values\nN-1']], defaut=dict(retrait=R18), lignes={
        u'AUTRES': dict(filet=PLEIN), u'TOTAL_IMMEUBLES': dict(retrait=0, filet=PLEIN, h=20.1, va=u'top'),
        u'PARTS_3': dict(affiche=u'Parts et actions des sociétés visées au 3° du I, art. L214-36 ' + ITAL + u'id.'),
        u'ACTIONS_4': dict(affiche=u'Actions négociées sur un marché réglementé (4° du I, art. L214-36 ' + ITAL
                                   + u'id.' + ITAL + u')'),
        u'OPCI_5': dict(affiche=u'Parts ou actions d’OPCI et organismes équivalents (5° du I, art. L214-36 ' + ITAL
                                + u'id.)'),
        u'AUTRES_IMMO': dict(filet=PLEIN), u'TOTAL_AUTRES': dict(retrait=0, filet=PLEIN, h=20.2, va=u'top'),
        u'TOTAL_GENERAL': dict(retrait=0, jc0=u'center', h=20.8)}),
    # 334-8 : les modeles d'exposition du reglement ANC n° 2020-07 (articles 434-2 a 434-6) quand l'OPCI gere des
    # actifs financiers ; sinon « Néant » (forme retenue le 05/10/2026 : ligne gardee sans occurrence). La cellule de
    # la table porte l'un ou l'autre ; aucune phrase de regle
    (u'334-8', 0): dict(forme=u'texte', lignes={u'RENVOI_2020_07': dict(sans_accroche=True)}),
    (u'335-1', 0): dict(forme=u'texte'),
    (u'335-2', 0): dict(largeurs=[62.5, 15.6, 21.9], titre_modele=True, h_entete=22.3, pas=12.9,
                        defaut=dict(retrait=R18), lignes={
        u'CL': dict(retrait=0), u'CL_ST': dict(filet=PLEIN, h=15.5), u'AC': dict(retrait=0),
        u'AC_ST': dict(filet=PLEIN, h=15.5), u'TOTAL': dict(retrait=0, h=18.4)}),
    (u'335-2', 1): dict(largeurs=[34.4, 15.6, 15.6, 15.5, 18.8], titre_modele=True, h_entete=24.0, pas=12.4,
                        defaut=dict(retrait=0),
                        interligne_entete=10.8,
                        entete=[[
        dict(t=u'Evolution des dépréciations', jc=u'left'), u'Situation N-1', u'Dotations de\nl’exercice',
        u'Reprises de\nl’exercice', u'Situation N']],
                        lignes={u'DEP_AC': dict(filet=PLEIN), u'DEP_TOTAL': dict(h=16.8)}),
    # 335-3 : corps d'un seul tenant de 54,4 pt, « Total » en pied (recueil, page 102 du fichier) ; la ligne des
    # provisions tient sur 2 lignes de texte (20,7 pt)
    (u'335-3', 0): dict(largeurs=[27.5, 18.9, 17.4, 18.9, 17.4], titre_modele=True, jc0=u'center', fin=True,
                        h_entete=23.3, pas=12.4,
                        interligne_entete=10.4,
                        defaut=dict(retrait=_depuis_cadre(18.24, MARGE_FIN)), entete=[[
        u'Détail des provisions\npour risques et charges', u'Situation\nExercice N-1', u'Dotations de\nl’exercice',
        u'Reprises de\nl’exercice', u'Situation\nExercice N']], lignes={u'TOTAL': dict(h=33.7, va=u'bottom')}),
    # 335-4 : texte en haut de sa ligne ; 17 pt sous chaque rubrique (recueil, page 102 du fichier)
    (u'335-4', 0): dict(largeurs=[70.2, 14.2, 15.7], titre_modele=True, jc0=u'center', fin=True, h_entete=20.2,
                        va=u'top',
                        defaut=dict(retrait=_depuis_cadre(34.92, MARGE_FIN)), lignes={
        u'EC': dict(retrait=_depuis_cadre(18.12, MARGE_FIN), h=17.1), u'EC_IMMO': dict(h=12.0),
        u'EC_NON_IMMO': dict(filet=PLEIN, h=16.5), u'AUTRES': dict(retrait=_depuis_cadre(18.12, MARGE_FIN), h=17.1),
        u'AUTRES_IMMO': dict(h=16.8), u'AUTRES_NON_IMMO': dict(filet=PLEIN, h=16.6),
        u'TOTAL': dict(retrait=_depuis_cadre(18.12, MARGE_FIN), h=17.4)}),
    (u'335-4', 1): dict(forme=u'texte', lignes={
        u'TAUX_FIXE': dict(sans_accroche=True, avant=[dict(
            intro=u'Pour les emprunts à taux fixe, les informations complémentaires suivantes sont présentées :')],
            puces=[
            u'les caractéristiques des emprunts (taux, échéance…) ;',
            u'l’écart entre la valeur contractuelle (ou de remboursement) de l’emprunt inscrite au bilan et sa valeur '
            u'actuelle, c’est-à-dire sa valeur actualisée au taux actuel du marché ;',
            u'le montant des pénalités de remboursement anticipé de l’emprunt.']),
        u'RATIOS': dict(sans_accroche=True, avant=[dict(
            intro=u'Pour l’ensemble des emprunts assortis de ratios financiers les principales caractéristiques sont '
                  u'précisées.')])}),
    (u'335-5', 0): dict(largeurs=[37.5, 17.2, 14.1, 14.1, 17.1], fin=True, h_entete=20.2, pas=12.4,
                        defaut=dict(retrait=R18), lignes={
        u'TF': dict(retrait=0), u'TF_IN_FINE': dict(filet=PLEIN), u'TV': dict(retrait=0),
        u'TV_IN_FINE': dict(filet=PLEIN), u'TOTAL': dict(retrait=0, h=14.3)}),
    (u'335-6', 0): dict(largeurs=[67.16, 15.63, 17.21], titre_modele=True, h_entete=25.7, pas=13.9,
                        defaut=dict(retrait=R18, filet=POINTILLE), lignes={
        u'D_5': dict(filet=PLEIN), u'TOTAL': dict(retrait=0, filet=None, h=15.0)}),
    # 335-7 : la phrase « Le cas échéant les postes et rubriques… » est une puce au recueil (page 103 du fichier)
    (u'335-7', 0): dict(forme=u'texte', lignes={
        u'PRODUITS_IMMO': dict(puce=True, avant=[dict(
            intro=u'Le détail des éléments compris dans chacune des lignes de produits et charges de l’activité '
                  u'immobilière fait apparaître les détails suivants :')],
            affiche=u'produits immobiliers en distinguant les loyers, les charges facturées et les autres revenus '
                    u'immobiliers en détaillant ce dernier poste'),
        u'CHARGES_IMMO': dict(puce=True, affiche=u'charges immobilières en distinguant les charges ayant leur '
                                                 u'contrepartie en produits, les charges d’entretien, le coût initial des '
                                                 u'éléments d’actifs renouvelés ou remplacés et les autres charges en '
                                                 u'détaillant ce dernier poste'),
        u'PRODUITS_PARTS': dict(puce=True, affiche=u'produits sur parts et actions', avant=[dict(
            puce=True, intro=u'Le cas échéant les postes et rubriques des résultats de l’activité immobilière sont '
                             u'détaillés et il est précisé la nature et le montant et traitements comptables des '
                             u'éléments suivants :')]),
        u'CHARGES_PARTS': dict(puce=True, affiche=u'charges sur parts et actions'),
        u'PRODUITS_AUTRES': dict(puce=True, affiche=u'produits issus des autres actifs à caractère immobilier'),
        u'CHARGES_AUTRES': dict(puce=True, affiche=u'autres charges issues des autres actifs à caractère immobilier'),
        u'ANTERIEURS': dict(puce=True, affiche=u'charges et produits imputables aux exercices antérieurs', avant=[dict(
            intro=u'Enfin, le cas échéant, il est notamment précisé la nature, le montant et le traitement comptable '
                  u'des éléments suivants :')]),
        u'EXCEPTIONNELS': dict(puce=True, affiche=u'produits exceptionnels et charges exceptionnelles'),
        u'REMBOURSEMENTS': dict(puce=True, affiche=u'remboursements de charges (ventilation par nature des charges)')}),
    (u'335-8', 0): dict(largeurs=[73.4, 12.3, 14.3], titre_modele=True, h_entete=21.5, pas=15.4,
                        defaut=dict(retrait=R18, filet=POINTILLE),
                        lignes={
        u'P': dict(retrait=0, italique=True, filet=None), u'P_ST': dict(italique=True, filet=PLEIN, h=16.0),
        u'P_5': dict(affiche=u'Produits sur Instruments financiers à terme'),
        u'C': dict(retrait=0, italique=True, filet=None), u'C_ST': dict(italique=True, filet=PLEIN, h=16.0),
        u'C_2': dict(affiche=u'Charges sur Instruments financiers à terme'),
        u'C_3': dict(affiche=u'Charges sur Opérations temporaires sur titres'),
        u'TOTAL': dict(retrait=0, italique=True, couleur=u'002060', filet=None, h=21.3)}),
    # 335-8, information exigee sans modele (page 104 du fichier)
    (u'335-8', 1): dict(forme=u'texte', lignes={u'DETAIL_NATURE': dict(sans_accroche=True, avant=[dict(
        intro=u'Le cas échéant les autres produits financiers et autres charges financières sont détaillés par '
              u'nature.')])}),
    # 335-9 : ni grille ni phrase de regle, le detail par nature seul (forme retenue le 05/10/2026)
    (u'335-9', 0): dict(forme=u'texte', lignes={u'DETAIL': dict(sans_accroche=True)}),
    # 335-10 : filets de 0,48 pt de part et d'autre de « Commissions de surperformance », 0,96 pt ailleurs
    (u'335-10', 0): dict(largeurs=[66.7, 15.8, 17.5], titre_modele=True, h_entete=26.4, pas=19.2,
                         defaut=dict(retrait=R18, filet=POINTILLE), lignes={
        u'F_5': dict(filet=FIN), u'F_6': dict(filet=FIN),
        u'TOTAL': dict(retrait=0, italique=True, filet=None, h=26.7)}),
    # 335-10, informations exigees sans modele (page 105 du fichier) : une ligne de la base chacune
    (u'335-10', 1): dict(forme=u'texte', lignes={u'PCT_ACTIF_NET_MOYEN': dict(puce=True, avant=[dict(
        intro=u'Les frais de gestion et frais de fonctionnement externes figurant au compte de résultat sont '
              u'explicités en détaillant les différents types de frais :')],
        affiche=u'les frais de fonctionnement et de gestion : indication du pourcentage de l’actif net moyen représenté '
                u'par ces frais au cours de l’exercice, le cas échéant, par catégorie de parts et actions')}),
    (u'335-10', 2): dict(forme=u'texte', lignes={u'FRAIS_EXTERNES_NATURE': dict(
        puce=True, affiche=u'les frais de fonctionnement externes détaillés par nature le cas échéant')}),
    (u'335-10', 3): dict(forme=u'texte', lignes={u'SURPERFORMANCE': dict(
        puce=True, affiche=u'les commissions de surperformance : indication du montant des frais de l’exercice, le cas '
                           u'échéant, par catégorie de parts et actions')}),
    # 335-11 : 2 enumerations du recueil (page 105 du fichier), coquille « instrument financiers » comprise
    (u'335-11', 0): dict(forme=u'texte', lignes={
        u'INVESTISSEMENTS': dict(sans_accroche=True, avant=[dict(
            intro=u'Les engagements donnés et/ou reçus affectant les investissements en actifs à caractère immobilier '
                  u'sont présentés comme suit :')],
            puces=[u'nature de l’engagement ;', u'montant de l’engagement.'],
            suite=u'Le cas échéant ces informations sont détaillées ou complétées, notamment avec l’échéance de '
                  u'l’engagement.'),
        u'IFT': dict(intro_rubrique=True, affiche=u'Les engagements sur instrument financiers à terme sont présentés de '
                                                  u'la manière suivante :'),
        u'IFT_REGLEMENTES': dict(puce=True, affiche=u'engagements sur marchés réglementés ou assimilés (par nature de '
                                                    u'produits)'),
        u'IFT_GRE_A_GRE': dict(puce=True, affiche=u'engagements de gré à gré (par nature de produit)'),
        u'IFT_AUTRES': dict(puce=True, affiche=u'autres engagements (par nature de produit)')}),
    # 335-12 : le filet de 0,48 pt entre les parties I et II ne couvre que la colonne des libelles ; filet du bas de
    # 0,48 pt ; une ligne blanche avant chaque « Détail… » et « Affectation… », et apres le « Détail… » de la partie I
    # (recueil, pages 105 et 106 du fichier) ; les 2 totaux de la partie II a leur place au modele, 14,88 pt et 28,80 pt
    # du cadre (page 106 du fichier), ceux de la partie I a 3,48 pt
    (u'335-12', 0): dict(largeurs=[72.6, 13.7, 13.7], titre_modele=True, h_entete=22.6, pas=10.9, cadre_bas=FIN,
                         interligne_entete=10.8,
                         entete=[[
        dict(t=u'Tableau d’affectation des sommes distribuables', jc=u'left'), _EXERCICE_N, _EXERCICE_N1]],
        defaut=dict(retrait=0), lignes=_maj(
        dict((c, dict(puce=True)) for c in (u'I_D_1', u'I_D_2', u'I_D_3', u'I_A_1', u'I_A_2', u'I_A_3', u'II_D_1',
                                            u'II_D_2', u'II_D_3', u'II_A_1', u'II_A_2', u'II_A_3')),
        # une ligne blanche avant chaque « Détail… » et chaque « Affectation… », et apres le « Détail… » de la partie I
        # (le 1er point de la partie II suit au pas de 12 pt, page 105 du fichier)
        dict((c, dict(ligne_blanche=True)) for c in (u'I_D', u'I_D_1', u'II_D', u'I_A', u'II_A')),
        {u'I_A_T': dict(filet=FIN, filet_cols=(0,)), u'II_D_T': dict(retrait=_depuis_cadre(14.88, MARGE_PLEIN)),
         u'II_A_T': dict(retrait=_depuis_cadre(28.80, MARGE_PLEIN))})),
    (u'335-12', 1): dict(largeurs=[34.9, 10.1, 11.5, 10.2, 11.9, 11.3, 10.1], titre_modele=True, h_entete=[36.8, 24.5],
                         pas=12.6,
                         marge=MARGE_ACOMPTES, interligne_entete=11.8,
                         entete=[
        [dict(t=u'Tableau des acomptes versés au titre de l’exercice', jc=u'left')] + _ENTETE_ACOMPTES[0][1:],
        _ENTETE_ACOMPTES_N]),
    (u'335-12', 2): dict(largeurs=[33.0, 9.8, 12.3, 13.8, 13.9, 9.1, 8.1], titre_modele=True, h_entete=[25.2, 24.5],
                         pas=12.6,
                         marge=MARGE_ACOMPTES, interligne_entete=11.7,
                         cadre_bas=FIN, entete=[
        [dict(t=u'Tableau des acomptes versés au titre de l’exercice précédent', jc=u'left')]
        + _ENTETE_ACOMPTES[0][1:], _ENTETE_ACOMPTES[1]]),
    # 335-12 : « Le cas échéant, il est à détailler pour chaque catégorie de parts », sous le titre de l'article et
    # avant le tableau d'affectation qu'il designe (page 105 du fichier, y 487,2, le tableau a y 513,8) : imprime avant le
    # bloc 0 (avant_bloc), sa valeur liee au tableau qui suit
    (u'335-12', 3): dict(forme=u'texte', avant_bloc=0, lignes={u'DETAIL_CATEGORIES': dict(
        sans_accroche=True, lie_au_suivant=True, avant=[dict(
            intro=u'Le cas échéant, il est à détailler pour chaque catégorie de parts.')])}),
    (u'336-1', 0): dict(forme=u'texte'),
    # 336-3 : intitules de l'enumeration du recueil, a la lettre (page 107 du fichier)
    (u'336-3', 0): dict(largeurs=[35.0, 15.0, 18.0, 15.0, 17.0], libre=True, pas=12.8, entete=[[
        u'libellé', u'quantité/ nominal', u'valeur actuelle', u'devise de cotation', u'% de l’actif net']]),
}
# CORPS DE LA NOTE d'un modele, celui du recueil : 9 pt (pages 91, 92, 98 et 107 du fichier), 6 pt pour la note de
# 333-3 (page 99). Une note n'est ni le contenu d'un tableau (10 pt) ni le corps de page (12 pt) : elle garde le
# corps du modele.
CORPS_NOTE = 9
CORPS_NOTES = {u'333-3': 6}
# hauteur des lignes « Total acomptes » des 2 tableaux des acomptes (recueil, page 106 du fichier)
H_TOTAL_ACOMPTES = {1: 16.0, 2: 16.6}
# 336-2 : immeubles (en-tete « Libellé » sur ses 4 colonnes), parts des societes (renvois en pied d'en-tete, 20,5 pt
# par societe, texte en haut, recueil, page 107 du fichier), autres actifs a caractere immobilier (un intitule par
# poste du bilan)
# Intitules des immeubles et des autres actifs : les mots de l'enumeration du recueil, a la lettre (page 106 du fichier :
# « libellé : adresse, surface, secteur, date d’acquisition ; prix de revient ; La valeur actuelle et le % de l’actif
# net » ; page 107 : « libellé, quantité/ nominal ; valeur actuelle ; % de l’actif net »).
CARTE_IMMEUBLES = dict(largeurs=[30.0, 11.0, 12.0, 12.0, 12.5, 12.5, 10.0], libre=True, souples=(1, 3, 4, 5, 6),
                       pas=12.8, entete=[
    [dict(t=u'libellé', span=4), dict(t=u'prix de revient', vm=u'restart'), dict(t=u'valeur actuelle', vm=u'restart'),
     dict(t=u'% de l’actif net', vm=u'restart')],
    [u'adresse', u'surface', u'secteur', u'date d’acquisition', dict(vm=u'continue'), dict(vm=u'continue'),
     dict(vm=u'continue')]])
# Tableau des parts : proportions des colonnes au recueil (page 106 du fichier) ; une denomination, une nature, un
# secteur passent a la ligne comme les intitules du recueil.
CARTE_FILIALES = dict(largeurs=[10.1, 13.1, 14.5, 20.3, 18.8, 14.5, 8.7], marge=MARGE_ACOMPTES,
                      lignes_entete=4, interligne_entete=11.7,
                      filets_entete=[None], h_entete=[53.0, 13.7], pas=20.5, va=u'top',
                      entete=[[u'', u'Nature des immeubles détenus', u'Secteur d’activité et localisation',
                               u'Date d’acquisition des parts ou des sous-jacents',
                               u'Coût de revient des titres frais exclus', u'Valeur actuelle des titres comptabilisée',
                               u'Différence d’estimation'],
                              [u'', u'(1)', u'', u'', u'(2)', u'(3)', u'(3) - (2)']])
H_GROUPE_FILIALES = 12.3
CARTE_AUTRES_IMMO = dict(largeurs=[40.0, 20.0, 22.0, 18.0], libre=True, pas=12.8, entete=[[
    u'libellé', u'quantité/ nominal', u'valeur actuelle', u'% de l’actif net']])

# Chasse des caracteres d'Arial, maigre et gras, en 1/2048 d'em : la largeur d'une colonne de montants se regle sur
# son contenu, sans police installee (la fonction Fabric n'en a pas).
_ARIAL_CARACTERES = u''.join(chr(c) for c in range(32, 127)) + u'àâäçéèêëîïôöùûüÿœÀÂÄÇÉÈÊËÎÏÔÖÙÛÜŒ’«»°²€…•  '
_ARIAL_MAIGRE = (u'569 569 727 1139 1139 1821 1366 391 682 682 797 1196 569 682 569 569 1139 1139 1139 1139 1139 1139 1139 '
                 u'1139 1139 1139 569 569 1196 1196 1196 1139 2079 1366 1366 1479 1479 1366 1251 1593 1479 569 1024 1366 '
                 u'1139 1706 1479 1593 1366 1593 1479 1366 1251 1479 1366 1933 1366 1366 1251 569 569 569 961 1139 682 1139 '
                 u'1139 1024 1139 1139 569 1139 1139 455 455 1024 455 1706 1139 1139 1139 1139 682 1024 569 1139 1024 1479 '
                 u'1024 1024 1024 684 532 684 1196 1139 1139 1139 1024 1139 1139 1139 1139 569 569 1139 1139 1139 1139 1139 '
                 u'1024 1933 1366 1366 1366 1479 1366 1366 1366 1366 569 569 1593 1593 1479 1479 1479 2048 455 1139 1139 '
                 u'819 682 1139 2048 717 569 410')
_ARIAL_GRAS = (u'569 682 971 1139 1139 1821 1479 487 682 682 797 1196 569 682 569 569 1139 1139 1139 1139 1139 1139 1139 '
               u'1139 1139 1139 682 682 1196 1196 1196 1251 1997 1479 1479 1479 1479 1366 1251 1593 1479 569 1139 1479 '
               u'1251 1706 1479 1593 1366 1593 1479 1366 1251 1479 1366 1933 1366 1366 1251 682 569 682 1196 1139 682 '
               u'1139 1251 1139 1251 1139 682 1251 1251 569 569 1139 569 1821 1251 1251 1251 1251 797 1139 682 1251 1139 '
               u'1593 1139 1139 1024 797 573 797 1196 1139 1139 1139 1139 1139 1139 1139 1139 569 569 1251 1251 1251 1251 '
               u'1251 1139 1933 1479 1479 1479 1479 1366 1366 1366 1366 569 569 1593 1593 1479 1479 1479 2048 569 1139 '
               u'1139 819 682 1139 2048 717 569 410')
_CHASSE = (dict(zip(_ARIAL_CARACTERES, (int(x) for x in _ARIAL_MAIGRE.split()))),
           dict(zip(_ARIAL_CARACTERES, (int(x) for x in _ARIAL_GRAS.split()))))


def _mesure(texte, taille, gras=False):
    u"""Largeur du texte en twips, en Arial au corps donne (un caractere hors table compte pour un « n »)."""
    t = _CHASSE[1 if gras else 0]
    return sum(t.get(c, 1139) for c in texte) * taille * 20.0 / 2048


def _segments_de_texte(texte, taille, gras):
    u"""Les mots d'un texte, comme Word les coupe : a l'espace, apres un trait d'union et a chaque retour a la ligne
    impose ; par ligne imposee, la liste des (largeur en twips, precede d'une espace)."""
    segments = []
    for seg in texte.split(u'\n'):
        jetons = []
        for i, mot in enumerate(seg.split(u' ')):
            for j, morceau in enumerate(re.split(u'(?<=-)', mot)):
                if morceau:
                    jetons.append((_mesure(morceau, taille, gras), i > 0 and j == 0))
        segments.append(jetons)
    return segments


def _compter_lignes(segments, espace, largeur):
    n = 0
    for jetons in segments:
        n, courant = n + 1, None
        for w, apres_espace in jetons:
            if courant is None:
                courant = w
            elif courant + (espace if apres_espace else 0) + w <= largeur:
                courant += (espace if apres_espace else 0) + w
            else:
                n, courant = n + 1, w
    return n


def _nombre_de_lignes(texte, taille, gras, largeur):
    u"""Le nombre de lignes d'un texte dans une largeur de texte donnee (twips), coupe comme Word le coupe."""
    return _compter_lignes(_segments_de_texte(texte, taille, gras), _mesure(u' ', taille, gras), largeur)


def _largeur_en_lignes(texte, taille, gras, lignes):
    u"""La plus petite largeur de texte (twips) ou un intitule tient sur « lignes » lignes au plus, coupe comme Word le
    coupe : a l'espace, apres un trait d'union et a chaque retour a la ligne impose."""
    segments = _segments_de_texte(texte, taille, gras)
    espace = _mesure(u' ', taille, gras)
    bas = max([w for jetons in segments for w, _ in jetons] or [0])
    haut = max([sum(w for w, _ in jetons) + espace * len(jetons) for jetons in segments] or [0])
    if _compter_lignes(segments, espace, bas) <= lignes:
        return bas
    while haut - bas > 5:
        milieu = (bas + haut) / 2.0
        if _compter_lignes(segments, espace, milieu) <= lignes:
            haut = milieu
        else:
            bas = milieu
    return haut


# Interligne simple d'Arial, rapport au corps (hauteurs de la police : 2 355 / 2 048 em) : 11,5 pt en 10 pt
INTERLIGNE_ARIAL = 1.15


def _interligne(x, taille):
    u"""L'interligne exact du modele (releve au recueil, en points), au moins l'interligne simple du corps : en 10 pt,
    un interligne du recueil plus serre que 11,5 pt couperait le haut des lettres ; None : interligne simple."""
    if not x:
        return None
    return round(max(x, INTERLIGNE_ARIAL * taille), 2)


# CE QUI DEBORDE PASSE SUR 2 LIGNES (06/10/2026) : un montant plus large que sa colonne se coupe a
# l'espace d'un groupe de chiffres (l'espace insecable devient une espace ordinaire, la ou Word peut couper), jamais
# au caractere ; « à remplir » se coupe a son espace. Les autres espaces insecables (« % », « m² ») restent.
_GROUPE = re.compile(u'(?<=\\d)\u00a0(?=\\d{3}(?!\\d))')


def _coupable(texte):
    return _GROUPE.sub(u' ', (texte or u'').replace(A_REMPLIR_INSECABLE, A_REMPLIR))


# UNE VALEUR DE CELLULE : un montant, un nombre, un pourcentage, une surface (format du module, espaces insecables) ou
# « à remplir » ; un texte (denomination, adresse, secteur) passe a la ligne a ses espaces, comme un libelle
_VALEUR = re.compile(u'^-?[0-9]{1,3}(?: [0-9]{3})*(?:,[0-9]+)?(?: %| m²)?$')


def _est_valeur(texte):
    t = (texte or u'').replace(AR, A_REMPLIR_INSECABLE).strip()
    return t in (A_REMPLIR_INSECABLE, A_REMPLIR) or bool(_VALEUR.match(t))

# ============================================================================ rapport de l'expert-comptable

# LA SOURCE UNIQUE EST LA BASE : le texte des 4 formes (exemples E1 a E4 de la NP 2300, recopies du PDF) est dans
# ref_forme_rapport, celui de la forme arretee et la donnee de chaque emplacement dans v_rapport_expert_comptable. Le
# module ne porte que le reperage des emplacements dans ce texte : les « … » et les descriptions entre parentheses,
# que le document remplit de la donnee de la base, ou de « à remplir » quand elle manque. Aucun mot n'est ajoute.

def _motif(litteral):
    u"""Le motif d'un emplacement : ses mots dans l'ordre, separes par un blanc quelconque (espace insecable comprise)
    ou par aucun, l'apostrophe droite ou courbe."""
    mots = [re.escape(m).replace(u'’', u"['’]") for m in litteral.split()]
    return re.compile(u'[\\s\u00a0\u202f]*'.join(mots))


EMPLACEMENTS = dict(
    comptes=_motif(u'comptes … (préciser « annuels », « intermédiaires »)'),
    entite=_motif(u'de … (préciser l’entité concernée)'),
    periode=_motif(u'relatifs à … (préciser la période ou l’exercice concerné)'),
    lettre=_motif(u'en date du …'),
    alternative=re.compile(u'%s|%s' % (_motif(u'annuels (ou « intermédiaires »)').pattern,
                                       _motif(u'annuels (ou intermédiaires)').pattern)))
# La liste des donnees (paragraphes « · ») : le libelle imprime et la colonne de v_rapport_expert_comptable.
DONNEES_RAPPORT = [(u'Total du bilan', u'total_du_bilan'), (u'Chiffre d’affaires', u'chiffre_d_affaires'),
                   (u'Résultat net comptable', u'resultat_net_comptable')]
SIGNATURE = u'Lieu, date et signature'

# ============================================================================ DIP, page « Le DIP » du 06/09/2026

TITRES_DIP = dict(
    bandeau=u'Le document d’information périodique',
    sous=u'Les 7 rubriques du rapport semestriel ; 6 sortent de la mission, la 7e relève de la société de gestion.',
    cartes=u'1 à 3 et 6. La valeur nette d’inventaire, les parts en circulation, la valeur par part, le dividende par '
           u'part',
    r1=u'1. L’état du patrimoine, par famille, en valeur actuelle',
    r5=u'5. Les mouvements du portefeuille dans la période',
    r5t=u'Acquisitions, cessions et travaux de la période',
    r4=u'4. Le portefeuille : l’inventaire du patrimoine, articles 336-1 à 336-3',
    r6=u'6. Les dividendes à verser : l’obligation minimale, article L. 214-69 du CMF',
    r7=u'7. Le récapitulatif des cas de plafonnement des rachats décidés au cours de la période relève de la société '
       u'de gestion : la solution ne porte pas ces décisions, cette rubrique est renseignée par elle au moment de la '
       u'mise au format du document.',
    fondement=u'Rubriques du rapport semestriel : code monétaire et financier, article L. 214-53 ; règlement général '
              u'de l’AMF, article 422-184 ; instruction AMF DOC-2011-23, article 28. Cette page assemble la matière '
              u'que la mission produit ; le document lui-même est établi par la solution au format prescrit et '
              u'attesté.')
CARTES_DIP = [u'ANR de l’arrêté', u'Parts à l’arrêté', u'VL de l’arrêté', u'Distribution par part']
CARTE_MOUVEMENTS = u'Mouvements de la période'
COLONNES_DIP = dict(
    r1=[u'Famille', u'Total actifs', u'Total valeur actuelle', u'Total différence d’estimation'],
    r5=[u'Actif', u'Nature', u'Date du mouvement', u'Montant'],
    r4=[u'Actif', u'Famille', u'Adresse', u'Valeur comptable', u'Valeur actuelle', u'Source de la valeur',
        u'Article du règlement'],
    r6=[u'Catégorie', u'Base de calcul', u'Taux', u'Obligation'])

# ============================================================================ lecture de la base

SQL = dict(
    interface_depot=u"SELECT (SELECT COUNT(*) FROM sys.parameters WHERE object_id = OBJECT_ID('dbo.pr_deposer_classeur') "
                    u"AND name IN ('@sous_dossier', '@empreinte')) AS parametres_depot, "
                    u"CASE WHEN OBJECT_ID('dbo.pr_enregistrer_document') IS NULL THEN 0 ELSE 1 END AS enregistrement, "
                    u"CASE WHEN OBJECT_ID('dbo.v_balance_etats_financiers') IS NULL THEN 0 ELSE 1 END AS balance",
    entite=u'SELECT code, denomination, forme_vehicule, forme_sociale, siren, adresse_1, adresse_2, code_postal, '
           u'ville, pays FROM dbo.ref_entite WHERE code = ?',
    arrete=u'SELECT r.entite, r.arrete, r.exercice, r.type_arrete, r.date_arrete, a.porte_balance, a.arrete_n_1, '
           u'a.premier_exercice, (SELECT r1.date_arrete FROM dbo.ref_arrete r1 WHERE r1.entite = r.entite '
           u'AND r1.arrete = a.arrete_n_1) AS date_arrete_n_1 FROM dbo.ref_arrete r LEFT JOIN dbo.v_arrete_etat a '
           u'ON a.entite = r.entite AND a.arrete = r.arrete WHERE r.entite = ? AND r.arrete = ?',
    livrable=u'SELECT livrable, libelle, article FROM dbo.v_livrables_dus WHERE entite = ? AND arrete = ? '
             u'AND livrable = ?',
    etat=u'SELECT ordre, code, libelle, type_ligne, romain, exercice_n, exercice_n_1, exercice_n_libelle, '
         u'exercice_n_1_libelle FROM dbo.v_ligne_etat_montant WHERE etat = ? AND entite = ? AND arrete = ? '
         u'ORDER BY ordre',
    tableaux=u'SELECT article, libelle, principe, texte_principe FROM dbo.ref_tableau_annexe',
    # L'ANNEXE : la table de la base, une ligne par cellule, a une version donnee (fn_annexe_cellules) ou a la derniere
    # (v_annexe_cellule) ; la version de l'arrete a la 1re lecture d'une production (v_annexe_version) ; le texte du
    # refus de retard de l'entite, tel que la base le leve (fn_ann_texte_retard), NULL sans marque
    annexe_derniere_version=u'SELECT derniere_version FROM dbo.v_annexe_version WHERE entite = ? AND arrete = ?',
    annexe_cellules_version=u'SELECT article, ligne, colonne, version, libelle, ordre, groupe, signe, etat, '
                            u'texte_imprime FROM dbo.fn_annexe_cellules(?, ?, ?) ORDER BY article, ordre, colonne',
    annexe_cellules_derniere=u'SELECT article, ligne, colonne, version, libelle, ordre, groupe, signe, etat, '
                             u'texte_imprime FROM dbo.v_annexe_cellule WHERE entite = ? AND arrete = ? '
                             u'ORDER BY article, ordre, colonne',
    annexe_texte_retard=u'SELECT dbo.fn_ann_texte_retard(?) AS texte',
    # UNE CELLULE DE LA DERNIERE VERSION DE L'ANNEXE (v_annexe_cellule) : le rapport de l'expert-comptable y lit la
    # devise de comptabilite, cellule 311-1 que les comptes annuels impriment (_devise)
    annexe_cellule=u'SELECT etat, texte_imprime FROM dbo.v_annexe_cellule WHERE entite = ? AND arrete = ? '
                   u'AND article = ? AND ligne = ? AND colonne = ?',
    # un autre format du meme document deja enregistre avec une autre version de l'annexe que celle lue (ordre de la
    # base : le 1er enregistre)
    annexe_autre_format=u'SELECT TOP (1) format, annexe_version FROM dbo.document_produit WHERE entite = ? '
                        u'AND arrete = ? AND livrable = ? AND version = ? AND ISNULL(annexe_version, -1) <> ? ORDER BY id',
    balance=u'SELECT compte_num, libelle, debit, credit, solde_n, solde_n_1 FROM dbo.v_balance_etats_financiers '
            u'WHERE entite = ? AND arrete = ? ORDER BY compte_num',
    # RAPPORT DE L'EXPERT-COMPTABLE : la forme arretee, son texte et ses emplacements ; les 4 formes ; la signature
    rapport=u'SELECT entite, arrete, soumise_commissariat_comptes, forme_arretee, forme_arretee_libelle, arretee_par, '
            u'arretee_le, forme_proposee, forme_proposee_libelle, forme_du_texte, intitule, texte, motif_limitation, '
            u'motif_considerations_particulieres, comptes_annuels_ou_intermediaires, entite_concernee, exercice, '
            u'date_arrete, total_du_bilan, chiffre_d_affaires, chiffre_d_affaires_lecture, resultat_net_comptable, '
            u'lettre_de_mission_du, description_des_points, lieu_date_signature, source '
            u'FROM dbo.v_rapport_expert_comptable WHERE entite = ? AND arrete = ?',
    formes=u'SELECT code, exemple, page_np, libelle, intitule, attestation, fondement, texte, ordre '
           u'FROM dbo.ref_forme_rapport ORDER BY ordre',
    signataire=u"SELECT dbo.fn_peut_viser_nature(?, ?, 'PUBLICATION', CAST(SYSUTCDATETIME() AS DATE)) AS peut",
    # DIP : le grain, puis une vue par rubrique ; la ligne « à remplir » d'une rubrique porte etat non NULL
    dip_arrete=u'SELECT entite, arrete, type_arrete, exercice_distribution FROM dbo.v_dip_arrete '
               u'WHERE entite = ? AND arrete = ?',
    dip_r1=u'SELECT famille_libelle, ordre, nombre_actifs, valeur_actuelle, difference_estimation, '
           u'valeur_nette_inventaire, etat FROM dbo.v_dip_rubrique_1 WHERE entite = ? AND arrete = ? ORDER BY ordre',
    dip_r2=u'SELECT nombre_parts, etat FROM dbo.v_dip_rubrique_2 WHERE entite = ? AND arrete = ?',
    dip_r3=u'SELECT actif_net_reevalue, nombre_parts, valeur_par_part, etat FROM dbo.v_dip_rubrique_3 '
           u'WHERE entite = ? AND arrete = ?',
    dip_r4=u'SELECT code_actif, famille_libelle, adresse, valeur_comptable, valeur_actuelle, source_libelle, article, '
           u'etat FROM dbo.v_dip_rubrique_4 WHERE entite = ? AND arrete = ? ORDER BY code_actif',
    dip_r5=u'SELECT nombre_mouvements, code_actif, nature, date_mouvement, montant, etat FROM dbo.v_dip_rubrique_5 '
           u'WHERE entite = ? AND arrete = ? ORDER BY date_mouvement, code_actif',
    dip_r6=u'SELECT exercice_lu, ordre, categorie, base_calcul, taux, obligation, distribution_par_part, lecture, etat '
           u'FROM dbo.v_dip_rubrique_6 WHERE entite = ? AND arrete = ? ORDER BY ordre',
)


def verifier_interface_depot(lire):
    u"""Le depot au sous-dossier Livrables avec empreinte et l'enregistrement par fichier existent en base."""
    r = (lire(SQL[u'interface_depot'], ()) or [{}])[0]
    if int(r.get(u'parametres_depot') or 0) < 2 or not int(r.get(u'enregistrement') or 0):
        raise Refus(u'Production refusée : la base ne porte pas le dépôt d’un livrable avec son empreinte, ou son '
                    u'enregistrement fichier par fichier.')
    return r


def _une(lire, cle, params, absent):
    l = lire(SQL[cle], params)
    if not l:
        raise Refus(absent)
    return l[0]


def lire_contexte(lire, entite, arrete, livrable):
    e = _une(lire, u'entite', (entite,), u'Production refusée : entité %s inconnue.' % entite)
    a = _une(lire, u'arrete', (entite, arrete), u'Production refusée : arrêté %s inconnu pour l’entité %s.'
             % (arrete, entite))
    l = _une(lire, u'livrable', (entite, arrete, livrable), u'Production refusée : le livrable %s n’est pas dû à '
             u'l’arrêté %s.' % (guillemets(NOMS_LIVRABLES.get(livrable, livrable)), arrete))
    return dict(entite=e, arrete=a, livrable=l)


def lire_etat(lire, entite, arrete, etat):
    u"""Les lignes de l'etat, controlees contre le modele du recueil : memes codes, meme ordre."""
    lignes = lire(SQL[u'etat'], (etat, entite, arrete))
    attendus = [c for c, _, _ in MODELES_ETATS[etat]]
    lus = [l[u'code'] for l in lignes]
    if lus != attendus:
        manquants = [c for c in attendus if c not in lus]
        en_trop = [c for c in lus if c not in attendus]
        s = lambda n: u's' if n > 1 else u''
        ecart = u' et '.join(x for x in (
            (u'%d ligne%s du modèle absente%s' % (len(manquants), s(len(manquants)), s(len(manquants))))
            if manquants else u'', (u'%d ligne%s en trop' % (len(en_trop), s(len(en_trop)))) if en_trop else u'')
            if x) or u'les mêmes lignes, en double ou dans un autre ordre'
        quoi, article = NOMS_ETATS[etat]
        raise Refus(u'Production refusée : les lignes %s en base ne sont pas celles du modèle de l’article %s du '
                    u'recueil, dans son ordre : %s.' % (quoi, article, ecart))
    par_code = dict((l[u'code'], l) for l in lignes)
    return [dict(par_code[c], libelle_recueil=lib, gras=bool(g)) for c, lib, g in MODELES_ETATS[etat]]


def lire_comptes_annuels(lire, entite, arrete, annexe_version=None):
    u"""Le contexte, les 3 etats, les intitules des articles, la balance, et les cellules de l'annexe : celles de la
    version annexe_version de l'arrete (la version que la production a fixee a sa 1re lecture, version_annexe), ou
    celles de la derniere version quand annexe_version est None. d['annexe_version'] dit la version lue."""
    ctx = lire_contexte(lire, entite, arrete, u'COMPTES_ANNUELS')
    d = dict(ctx)
    for etat in MODELES_ETATS:
        d[etat] = lire_etat(lire, entite, arrete, etat)
    tableaux = lire(SQL[u'tableaux'], ())
    d[u'tableaux'] = dict((t[u'article'], t[u'libelle']) for t in tableaux)
    # LES ARTICLES DE PRINCIPE : {article: [alinea]}, le texte de la base coupe a ses sauts de ligne
    d[u'principes'] = dict((t[u'article'], t[u'texte_principe'].split(u'\n')) for t in tableaux
                           if t.get(u'principe') and t.get(u'texte_principe'))
    d[u'balance'] = lire(SQL[u'balance'], (entite, arrete))
    lire_annexe(lire, d, entite, arrete, annexe_version)
    return d


def version_annexe(lire, entite, arrete):
    u"""LA VERSION DE L'ANNEXE QUE LA PRODUCTION FIXE, a sa 1re lecture : la derniere version de l'arrete
    (v_annexe_version). Toutes les cellules du document se lisent ensuite a cette version (lire_comptes_annuels),
    et le document l'enregistre."""
    r = (lire(SQL[u'annexe_derniere_version'], (entite, arrete)) or [{}])[0].get(u'derniere_version')
    if r is None:
        raise Refus(u'Production refusée : l’annexe de l’arrêté %s n’a pas encore été calculée ; recalculez-la avec '
                    u'le bouton %s, puis produisez de nouveau.' % (arrete, guillemets(u'Recalculer')))
    return int(r)


def lire_annexe(lire, d, entite, arrete, annexe_version=None):
    u"""LES CELLULES DE L'ANNEXE, TELLES QUE LA TABLE DE LA BASE LES PORTE : aucune n'est calculee, deduite ni
    completee ici. d['cellules'] : {(article, ligne, colonne): dict(etat, texte)}, texte etant le texte imprime de la
    table ; d['annexe'] : {article: [ligne]}, les lignes de chaque article dans l'ordre de la table (ordre), chacune
    avec son code (la cle de ligne de la table), son libelle, son groupe et son signe ; d['annexe_version'] : la
    version lue, une seule pour tout le document."""
    if annexe_version is None:
        rows = lire(SQL[u'annexe_cellules_derniere'], (entite, arrete))
    else:
        rows = lire(SQL[u'annexe_cellules_version'], (entite, arrete, int(annexe_version)))
    if not rows:
        raise Refus(u'Production refusée : la base ne porte aucune cellule de l’annexe de l’arrêté %s%s ; recalculez-la '
                    u'avec le bouton %s, puis produisez de nouveau.'
                    % (arrete, (u' en version %d' % int(annexe_version)) if annexe_version is not None else u'',
                       guillemets(u'Recalculer')))
    versions = sorted(set(int(r[u'version']) for r in rows))
    if len(versions) != 1:
        raise Refus(u'Production refusée : les cellules de l’annexe de l’arrêté %s lues ensemble portent %d versions ; '
                    u'produisez de nouveau.' % (arrete, len(versions)))
    d[u'annexe_version'] = versions[0]
    d[u'cellules'], d[u'annexe'], d[u'sans_cellule'] = {}, {}, []
    lignes = {}
    for r in rows:
        art, code, k = r[u'article'], r[u'ligne'], int(r[u'colonne'])
        l = lignes.get((art, code))
        if l is None:
            l = lignes[(art, code)] = dict(article=art, code=code, libelle=r[u'libelle'] or u'', ordre=r[u'ordre'],
                                           groupe=r.get(u'groupe'), signe=r.get(u'signe'))
            d[u'annexe'].setdefault(art, []).append(l)
        # une cellule a remplir, en attente ou sans objet s'imprime vide, quel que soit le texte lu
        d[u'cellules'][(art, code, k)] = dict(etat=r[u'etat'], texte=u'' if r[u'etat'] in ETATS_VIDES
                                              else r[u'texte_imprime'] or u'')
    for art in d[u'annexe']:
        d[u'annexe'][art].sort(key=lambda l: int(l[u'ordre'] or 0))
    return d


# LES REFUS DE L'ENREGISTREMENT D'UN DOCUMENT, tels que la base les oppose (pr_enregistrer_document, refus 50648 et
# 50649), dans son ordre et avec ses textes : la fonction des documents les juge avant de deposer le fichier, pour ne
# deposer aucun fichier que la base refuserait.
def garde_annexe(lire, entite, arrete, annexe_version, version=None):
    u"""Les comptes annuels imprimes a la version annexe_version de l'annexe, version (le numero du document) etant
    celle de la demande en cours : refus quand la version lue n'est pas donnee, quand l'annexe de l'arrete est passee a
    une version posterieure depuis la lecture, quand un autre format du meme document est deja enregistre avec une
    autre version de l'annexe, ou quand les cellules de l'entite sont en retard sur une ecriture. Rien sinon."""
    if annexe_version is None:
        raise Refus(u'Document refusé : la version de l’annexe lue par le document n’est pas donnée.')
    v = int(annexe_version)
    w = (lire(SQL[u'annexe_derniere_version'], (entite, arrete)) or [{}])[0].get(u'derniere_version')
    if w is None:
        raise Refus(u'Document refusé : l’annexe de l’arrêté %s n’a aucune version calculée ; recalculez-la avec le '
                    u'bouton %s, puis produisez de nouveau.' % (arrete, guillemets(u'Recalculer')))
    if v < int(w):
        raise Refus(u'Document refusé : l’annexe de l’arrêté %s est passée en version %d après la lecture du document '
                    u'(version %d) ; produisez-le de nouveau.' % (arrete, int(w), v))
    if version is not None:
        autre = (lire(SQL[u'annexe_autre_format'], (entite, arrete, u'COMPTES_ANNUELS', int(version), v)) or [{}])[0]
        if autre.get(u'format'):
            raise Refus(u'Document refusé : le format %s de ce document porte la version %s de l’annexe ; produisez une '
                        u'version neuve du document.' % (autre[u'format'], u'inconnue' if autre.get(u'annexe_version')
                                                         is None else int(autre[u'annexe_version'])))
    garde_retard(lire, entite)


def garde_retard(lire, entite):
    u"""Refus quand les cellules de l'annexe de l'entite sont en retard sur une ecriture d'une source : la base l'oppose
    aux comptes annuels et au rapport de l'expert-comptable. Le texte est celui de la base (fn_ann_texte_retard), que
    l'enregistrement du document leve tel quel : rendu sans retouche, rien quand la base ne rend rien."""
    texte = (lire(SQL[u'annexe_texte_retard'], (entite,)) or [{}])[0].get(u'texte')
    if texte is not None:
        raise Refus(texte)


def lire_rapport(lire, entite, arrete):
    u"""La ligne de v_rapport_expert_comptable de l'arrete : la forme arretee, la proposee, le texte de la forme
    arretee (de la proposee tant qu'aucune ne l'est) et une colonne par emplacement ; puis les 4 formes de
    ref_forme_rapport, chacune avec son exemple de la NP 2300 et son texte."""
    d = lire_contexte(lire, entite, arrete, u'ATTESTATION')
    d[u'rapport'] = _une(lire, u'rapport', (entite, arrete),
                         u'Production refusée : la base ne rend pas les données du rapport de l’expert-comptable de '
                         u'l’entité %s à l’arrêté %s.' % (entite, arrete))
    d[u'formes'] = dict((f[u'code'], f) for f in lire(SQL[u'formes'], ()))
    d[u'forme_arretee'] = d[u'rapport'].get(u'forme_arretee')
    lire_cellule(lire, d, entite, arrete, ARTICLE_DEVISE, LIGNE_DEVISE, 1)
    return d


def lire_cellule(lire, d, entite, arrete, article, code, k):
    u"""UNE CELLULE DE LA DERNIERE VERSION DE L'ANNEXE, versee a d['cellules'] comme lire_annexe la verse : vide pour
    une cellule a remplir, en attente ou sans objet ; absente de d['cellules'] quand la table ne la porte pas."""
    d.setdefault(u'cellules', {})
    for r in lire(SQL[u'annexe_cellule'], (entite, arrete, article, code, int(k))):
        d[u'cellules'][(article, code, int(k))] = dict(etat=r[u'etat'], texte=u'' if r[u'etat'] in ETATS_VIDES
                                                       else r[u'texte_imprime'] or u'')
    return d


def lire_dip(lire, entite, arrete):
    u"""Le grain (v_dip_arrete), puis les rubriques 1 a 6, une vue chacune ; a un arrete non annuel, la rubrique 6
    lit le dernier exercice clos et le dit dans sa colonne lecture."""
    d = lire_contexte(lire, entite, arrete, u'DIP')
    d[u'dip'] = _une(lire, u'dip_arrete', (entite, arrete), u'Production refusée : le DIP n’est pas dû à l’arrêté %s '
                     u'de l’entité %s.' % (arrete, entite))
    for k in range(1, 7):
        d[u'r%d' % k] = lire(SQL[u'dip_r%d' % k], (entite, arrete))
    return d


# ============================================================================ Word

_W = u'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_FILETS = {PLEIN: u'w:val="single" w:sz="8" w:space="0" w:color="000000"',
           FIN: u'w:val="single" w:sz="4" w:space="0" w:color="000000"',
           POINTILLE: u'w:val="dotted" w:sz="4" w:space="0" w:color="000000"', None: u'w:val="nil"'}
EPAISSEUR = {PLEIN: 1.0, FIN: 0.5, POINTILLE: 0.5}      # en points, w:sz / 8
CADRE = PLEIN


def _xml(t):
    return t.replace(u'&', u'&amp;').replace(u'<', u'&lt;').replace(u'>', u'&gt;').replace(u'"', u'&quot;')


def _run_xml(texte, gras=False, italique=False, taille=9, couleur=None, surligne=False, exposant=False):
    u"""Un w:r : le texte, ses retours a la ligne (\\n) et ses tabulations (\\t)."""
    rpr = (u'<w:b/><w:bCs/>' if gras else u'') + (u'<w:i/><w:iCs/>' if italique else u'')
    if couleur:
        rpr += u'<w:color w:val="%s"/>' % couleur
    rpr += u'<w:sz w:val="%d"/><w:szCs w:val="%d"/>' % (int(round(taille * 2)), int(round(taille * 2)))
    if surligne:
        rpr += u'<w:highlight w:val="yellow"/>'
    if exposant:
        rpr += u'<w:vertAlign w:val="superscript"/>'
    contenu = u''
    for i, ligne in enumerate(texte.split(u'\n')):
        if i:
            contenu += u'<w:br/>'
        for j, s in enumerate(ligne.split(u'\t')):
            if j:
                contenu += u'<w:tab/>'
            if s:
                contenu += u'<w:t xml:space="preserve">%s</w:t>' % _xml(s)
    return u'<w:r><w:rPr>%s</w:rPr>%s</w:r>' % (rpr, contenu)


def _puce_xml(taille=9):
    u"""La puce du modele, caractere F0B7 de la police Symbol, comme au recueil (SymbolMT)."""
    return (u'<w:r><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:cs="Symbol"/><w:sz w:val="%d"/>'
            u'<w:szCs w:val="%d"/></w:rPr><w:sym w:font="Symbol" w:char="F0B7"/></w:r>'
            % (int(round(taille * 2)), int(round(taille * 2))))


def _runs_xml(segments, taille=9, couleur=None, insecable=True):
    u"""Les w:r d'une suite de segments (texte, gras, italique) : la marque AR devient « à remplir » en italique
    surligne (sans coupure de ligne, sauf dans une cellule ou il deborde : insecable False), ce qui suit EXPOSANT dans
    son morceau s'ecrit en exposant, ITAL bascule l'italique, PUCE est la puce du modele en police Symbol."""
    marque = A_REMPLIR_INSECABLE if insecable else A_REMPLIR
    sortie = u''
    for texte, gras, italique in segments:
        texte = texte or u''
        if texte.startswith(PUCE):
            sortie += _puce_xml(taille)
            texte = texte[1:]
        morceaux = texte.split(AR)
        for i, m in enumerate(morceaux):
            for j, s in enumerate(m.split(EXPOSANT)):
                it = italique
                for k, seg in enumerate(s.split(ITAL)):
                    if k:
                        it = not it
                    if seg:
                        sortie += _run_xml(seg, gras, it, taille, couleur, exposant=bool(j))
            if i < len(morceaux) - 1:
                sortie += _run_xml(marque, gras, True, taille, couleur, surligne=True)
    return sortie


def _para_xml(runs, jc=None, retrait=0, suspendu=0, avant=0, garder=False, tab=None, taille=9, apres=0,
              interligne=None):
    u"""Un w:p de cellule ; la marque de paragraphe prend le corps du tableau : une cellule vide n'agrandit pas sa
    ligne. avant, apres : espace avant et apres le paragraphe, en points ; interligne : l'interligne exact du modele,
    en points (simple sinon)."""
    ppr = u'<w:keepNext/>' if garder else u''
    if tab:
        ppr += u'<w:tabs><w:tab w:val="left" w:pos="%d"/></w:tabs>' % tab
    ligne = u'w:line="240" w:lineRule="auto"'
    if interligne:
        ligne = u'w:line="%d" w:lineRule="exact"' % int(round(interligne * 20))
    ppr += u'<w:spacing w:before="%d" w:after="%d" %s/>' % (int(round(avant * 20)), int(round(apres * 20)), ligne)
    if retrait or suspendu:
        ppr += u'<w:ind w:left="%d"%s/>' % (retrait, (u' w:hanging="%d"' % suspendu) if suspendu else u'')
    if jc:
        ppr += u'<w:jc w:val="%s"/>' % jc
    ppr += u'<w:rPr><w:sz w:val="%d"/><w:szCs w:val="%d"/></w:rPr>' % (int(round(taille * 2)), int(round(taille * 2)))
    return u'<w:p><w:pPr>%s</w:pPr>%s</w:p>' % (ppr, runs)


def _allouer(proportions, besoins, souples, total=LARGEUR_UTILE):
    u"""Largeurs des colonnes en twips. Une colonne souple (montants) prend la largeur du modele, elargie a son contenu
    qui ne se coupe pas ; les colonnes fixes (libelles) se partagent le reste, au prorata du modele. Une colonne fixe
    plus etroite que le modele ou que son plus long mot reprend la largeur des colonnes souples au-dela de leur
    contenu : le libelle garde la place que le modele lui donne quand les montants le permettent."""
    n = len(proportions)
    s = float(sum(proportions))
    base = [total * p / s for p in proportions]
    w = list(base)
    for j in souples:
        w[j] = max(base[j], besoins[j])
    fixes = [j for j in range(n) if j not in souples]
    if fixes:
        cible = sum(max(base[j], besoins[j]) for j in fixes)
        reste = total - sum(w[j] for j in souples)
        manque = cible - reste
        marges = [(k, w[k] - besoins[k]) for k in souples if w[k] > besoins[k]]
        dispo = sum(m for _, m in marges)
        if manque > 0 and dispo > 0:
            pris = min(manque, dispo)
            for k, m in marges:
                w[k] -= pris * m / dispo
            reste += pris
        sf = float(sum(proportions[j] for j in fixes))
        for j in fixes:
            w[j] = reste * proportions[j] / sf
        # une colonne fixe dont le contenu ne se coupe pas (sans_coupure : sa largeur sur une ligne) garde au moins sa
        # largeur quand le reste le permet ; les autres colonnes fixes se partagent ce qui reste, au prorata du modele
        if any(w[j] < besoins[j] for j in fixes) and reste >= sum(besoins[j] for j in fixes):
            libres, r = list(fixes), reste
            while libres:
                part = r / float(sum(proportions[j] for j in libres))
                serrees = [j for j in libres if besoins[j] > part * proportions[j]]
                if not serrees:
                    for j in libres:
                        w[j] = part * proportions[j]
                    break
                for j in serrees:
                    w[j] = besoins[j]
                    r -= besoins[j]
                    libres.remove(j)
    if sum(w) > total:                      # le contenu ne tient pas : on revient aux proportions
        k = total / sum(w)
        w = [x * k for x in w]
    w = [int(round(x)) for x in w]
    w[0] += total - sum(w)
    return w


def _forme_modele(entetes, lignes, carte, numeriques, total):
    u"""LA FORME D'UN TABLEAU A SON MODELE, sans l'ecrire : ses rangees (cellules, filets, hauteurs), les largeurs de
    ses colonnes dans la justification « total » (twips), sa marge de cellule ; deborde : une cellule de montant ou
    « à remplir » ne tient pas sur une ligne dans sa colonne (elle passe sur 2 lignes, coupee a l'espace d'un groupe
    de chiffres)."""
    n = len(carte[u'largeurs'])
    corps_taille = carte.get(u'taille') or CORPS_TABLEAU
    numeriques = set(numeriques)
    centrees = set(carte.get(u'centrees') or ())
    V = set(carte.get(u'verticales') or range(n + 1)) | {0, n}

    def ep(style):
        u"""Le filet a son epaisseur : une carte « fin » trace en FIN ses filets pleins."""
        return FIN if style == PLEIN and carte.get(u'fin') else style
    cadre = ep(CADRE)
    cadre_bas = ep(carte.get(u'cadre_bas') or CADRE)
    h_entete = carte.get(u'h_entete')
    if h_entete is not None and not isinstance(h_entete, (list, tuple)):
        h_entete = [h_entete]
    # ---- en-tete : lignes de cellules (texte ou dict) ; la 1re cellule porte le titre du modele quand la carte le
    # dit (titre_modele) ou que la cellule le dit (titre : le titre long des 3 etats)
    spec = carte.get(u'entete')
    if spec is None:
        e0 = carte.get(u'titre') if carte.get(u'titre') is not None else (entetes[0] if entetes else u'')
        spec = [[dict(t=e0, jc=carte.get(u'jc0', u'left'))] + list((entetes or [])[1:])]
    grille = []                     # lignes : dict(cellules=[dict(col, span, ...)], filet, entete, h)
    for i, rang in enumerate(spec):
        cellules, col = [], 0
        for j, c in enumerate(rang):
            if c is None:
                c = dict(t=(entetes[0] if entetes else u''), jc=carte.get(u'jc0', u'left')) if j == 0 else dict(t=u'')
            elif not isinstance(c, dict):
                c = dict(t=c, jc=carte.get(u'jc0', u'left') if j == 0 else u'center')
            span = c.get(u'span', 1)
            t = c.get(u't')
            titre = i == 0 and j == 0 and bool(c.get(u'titre') or carte.get(u'titre_modele'))
            segs = t if isinstance(t, list) else [(lettre(t or u''), c.get(u'gras', True), False)]
            taille = TITRE_MODELE if titre else (c.get(u'taille') or corps_taille)
            cellules.append(dict(col=col, span=span, segs=segs, jc=c.get(u'jc', u'center'),
                                 va=c.get(u'va') or u'center', vm=c.get(u'vm'), taille=taille,
                                 couleur=BLEU_TITRE if titre else None,
                                 interligne=None if titre else _interligne(carte.get(u'interligne_entete'), taille)))
            col += span
        filets = carte.get(u'filets_entete') or []
        filet = (filets[i] if i < len(filets) else PLEIN) if i < len(spec) - 1 else carte.get(u'filet_entete', PLEIN)
        grille.append(dict(cellules=cellules, filet=ep(filet), entete=True,
                           h=h_entete[min(i, len(h_entete) - 1)] if h_entete else None))
    # ---- corps : les lignes, chacune precedee des lignes « avant » et suivie des lignes « apres » qu'elle porte
    corps = []
    for l in lignes:
        corps += list(l.get(u'avant') or []) + [l] + list(l.get(u'apres') or [])
    for k, l in enumerate(corps):
        cel = l.get(u'cellules') or []
        if l.get(u'dispo'):             # la disposition de la ligne, meme fusionnee pour la lecture
            dispo = [tuple(x) + (None,) * (3 - len(x)) for x in l[u'dispo']]
        elif l.get(u'fusion'):
            dispo = [(0, n, None)]
        else:
            dispo = [(i, 1, None) for i in range(n)]
        gras, italique = bool(l.get(u'gras')), bool(l.get(u'italique'))
        taille = l.get(u'taille') or corps_taille
        retrait = l[u'retrait'] if l.get(u'retrait') is not None else R9 * (l.get(u'niveau') or 0)
        hauteur = l.get(u'h') or carte.get(u'pas')
        va = l.get(u'va') or carte.get(u'va') or (u'center' if hauteur else None)
        cellules, col = [], 0
        for idx, span, jc in dispo:
            v = cel[idx] if idx < len(cel) else None
            para = dict(retrait=0, suspendu=0, tab=None)
            if idx == 0:
                if not l.get(u'donnees'):
                    v = lettre(l.get(u'affiche') or v)
                jc = jc or l.get(u'jc0') or u'left'
                if jc == u'left':
                    para[u'retrait'] = retrait
                if l.get(u'puce') and v:
                    v = PUCE + u'\t' + v
                    para.update(retrait=retrait + 720, suspendu=360, tab=retrait + 720)
                it = italique or bool(l.get(u'italique_libelle'))
            else:
                jc = jc or (u'center' if col in centrees else u'right' if idx in numeriques else u'left')
                it = italique
            plein_ar = v == AR
            segs = [] if v is None else ([(A_REMPLIR_INSECABLE, gras, True)] if plein_ar else [(v, gras, it)])
            t_cel = (l.get(u'taille_libelle') or taille) if idx == 0 else taille
            cellules.append(dict(col=col, span=span, segs=segs, jc=jc, va=va, taille=t_cel,
                                 fond=u'FFF2CC' if plein_ar else None, couleur=l.get(u'couleur'),
                                 avant=l.get(u'espace_avant') or 0, apres=l.get(u'espace_apres') or 0,
                                 interligne=_interligne(l.get(u'interligne') or carte.get(u'interligne'), t_cel),
                                 **para))
            col += span
        filet = PLEIN if carte.get(u'quadrillage') else l.get(u'filet')
        grille.append(dict(cellules=cellules, filet=ep(filet), garder=l.get(u'garder'), h=hauteur,
                           filet_cols=l.get(u'filet_cols'), pas_2_lignes=l.get(u'pas_2_lignes')))
    # les 3 dernieres lignes ne se separent pas : une ligne seule ne passe pas en haut de la page suivante ;
    # l'en-tete reste avec les 2 premieres lignes du corps : il ne finit pas seul une page
    corps_rangs = [r for r in grille if not r.get(u'entete')]
    for r in corps_rangs[-3:-1] + [r for r in grille if r.get(u'entete')] + corps_rangs[:1]:
        r[u'garder'] = True
    # un tableau court (15 lignes au plus et 260 pt au plus de hauteur minimale) ne se coupe pas : il passe entier a
    # la page suivante
    if len(grille) <= GARDE_ENTIER and sum(r.get(u'h') or 12.0 for r in grille) <= 260:
        for r in grille[:-1]:
            r[u'garder'] = True
    marge = carte[u'marge'] if carte.get(u'marge') is not None else (MARGE_FIN if carte.get(u'fin') else
                                                                      MARGE_PLEIN)
    # une cellule de montant cadree a droite, dans un modele a marge etroite, prend la marge des montants (reglage
    # MARGE_MONTANTS ; None : la marge du modele)
    marge_montants = MARGE_MONTANTS if MARGE_MONTANTS is not None and marge < MARGE_MONTANTS else None
    if marge_montants is not None:
        for r in grille:
            for c in r[u'cellules']:
                if not r.get(u'entete') and c[u'jc'] == u'right' and c[u'col'] > 0:
                    c[u'marge'] = marge_montants
    # ---- largeurs : celles du modele, fixes, dans la justification ; un tableau de presentation libre regle ses
    # colonnes de montants sur leur contenu
    if carte.get(u'libre'):
        souples = set(carte.get(u'souples') if carte.get(u'souples') is not None else
                      [j for j in range(1, n) if j in numeriques or j in centrees])
        besoins = [0.0] * n
        for r in grille:
            for c in r[u'cellules']:
                if c[u'span'] != 1:
                    continue
                for texte, g, _ in c[u'segs']:
                    texte = (texte or u'').replace(AR, A_REMPLIR).replace(EXPOSANT, u'').replace(ITAL, u'')
                    if c[u'col'] in souples and not r.get(u'entete'):
                        b = max(_mesure(x, c[u'taille'], g) for x in texte.split(u'\n'))
                    else:            # un intitule se coupe a l'espace et apres le trait d'union
                        b = max([_mesure(x, c[u'taille'], g) for x in re.split(u'(?<=-)|[ \n\t]', texte)] or [0])
                        if r.get(u'entete'):     # un intitule d'en-tete tient sur 3 lignes au plus
                            b = max(b, _largeur_en_lignes(texte, c[u'taille'], g, carte.get(u'lignes_entete') or 3))
                    b += c.get(u'retrait', 0) + marge + c.get(u'marge', marge) + 20
                    besoins[c[u'col']] = max(besoins[c[u'col']], b)
        for groupe in carte.get(u'egales') or ():           # colonnes de meme nature, de meme largeur
            m = max(besoins[j] for j in groupe)
            for j in groupe:
                besoins[j] = m
        w = _allouer(carte[u'largeurs'], besoins, souples, total)
    else:
        somme = float(sum(carte[u'largeurs']))
        w = [int(round(total * p / somme)) for p in carte[u'largeurs']]
        w[0] += total - sum(w)
    # ---- CE QUI DEBORDE PASSE SUR 2 LIGNES : un montant ou un « à remplir » plus large que sa cellule se coupe a
    # l'espace d'un groupe de chiffres (_coupable), jamais au caractere ; une rangee dont le libelle tient sur 2
    # lignes au recueil et sur une seule au Word perd l'interligne de sa 2e ligne (LIBELLES_2_LIGNES)
    deborde = False
    for r in grille:
        if r.get(u'entete'):
            continue
        for c in r[u'cellules']:
            a, b = c[u'col'], c[u'col'] + c[u'span']
            dispo = sum(w[a:b]) - marge - c.get(u'marge', marge) - c.get(u'retrait', 0) - 10
            large = 0.0
            for t, g, _ in c[u'segs']:
                t = (t or u'').replace(AR, A_REMPLIR_INSECABLE).replace(EXPOSANT, u'').replace(ITAL, u'')
                large += max([_mesure(x, c[u'taille'], g) for x in t.split(u'\n')] or [0])
            if c[u'col'] > 0 and large > dispo and _est_valeur(u''.join((t or u'') for t, _, _ in c[u'segs'])):
                c[u'coupable'] = deborde = True
                c[u'segs'] = [(_coupable(t), g, it) for t, g, it in c[u'segs']]
        if r.get(u'pas_2_lignes') and r.get(u'h') and r[u'cellules']:
            c = r[u'cellules'][0]
            texte = u''.join((t or u'') for t, _, _ in c[u'segs']).replace(ITAL, u'').replace(EXPOSANT, u'')
            gras = any(g for _, g, _ in c[u'segs'])
            dispo = sum(w[c[u'col']:c[u'col'] + c[u'span']]) - 2 * marge - c.get(u'retrait', 0) - 20
            if _nombre_de_lignes(texte, c[u'taille'], gras, dispo) < 2:
                r[u'h'] = round(r[u'h'] - r[u'pas_2_lignes'], 2)
    return dict(grille=grille, w=w, marge=marge, cadre=cadre, cadre_bas=cadre_bas, V=V, deborde=deborde)


class _Word(object):
    u"""Le modele Word du cabinet rendu par docxtpl (en-tete, pied de page), puis le corps ecrit a la place du
    paragraphe [CORPS] avec python-docx, que docxtpl embarque."""

    def __init__(self, modele, entete, pied):
        from docxtpl import DocxTemplate
        if not modele:
            raise Refus(u'Production refusée : le modèle Word du cabinet est vide.')
        self.tpl = DocxTemplate(io.BytesIO(modele))
        self.tpl.render({u'entete': entete, u'pied': pied}, autoescape=True)   # « & » d'une denomination
        self.doc = self.tpl.docx
        self._champs_au_format_du_texte()
        self.ancre = None
        for p in self.doc.paragraphs:
            if p.text.strip() == u'[CORPS]':
                self.ancre = p
                break
        self.avant = list(self.doc.element.body)
        self.styles = set(s.name for s in self.doc.styles)
        self.titres = []                # (niveau, texte, signet) des titres de niveau 1 a 3, pour le sommaire
        self.place_sommaire = None
        self.titre_document, self.date_document = entete, None
        self.largeur = LARGEUR_UTILE    # justification de la section courante, en twips (mise_en_page, orientation)
        self.largeur_portrait = LARGEUR_UTILE
        self.marge_page = None
        self.paysage = False
        self.taille_sommaire = 10

    def style(self, nom):
        return nom if nom in self.styles else None

    def _champs_au_format_du_texte(self):
        u"""LE NUMERO DE PAGE A LE FORMAT DU TEXTE QU'IL SUIT : dans l'en-tete et le pied de page, un champ (PAGE) dont
        les runs n'ont aucune mise en forme prend celle du run de texte qui le precede dans son paragraphe (« Version 1
        produite le …, page » en 8 pt gris) ; sans elle, Word l'imprimait au corps de page, en 12 pt noir. Le modele du
        cabinet n'a pas a etre redepose."""
        import copy
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml.ns import qn
        for rel in list(self.doc.part.rels.values()):
            if rel.reltype not in (RT.HEADER, RT.FOOTER) or rel.is_external:
                continue
            for p in rel.target_part.element.iter(qn(u'w:p')):
                modele = None
                for r in p.iter(qn(u'w:r')):
                    rpr = r.find(qn(u'w:rPr'))
                    champ = r.find(qn(u'w:fldChar')) is not None or r.find(qn(u'w:instrText')) is not None
                    if champ and rpr is None and modele is not None:
                        r.insert(0, copy.deepcopy(modele))
                    elif rpr is not None and r.find(qn(u'w:t')) is not None:
                        modele = rpr
                # UN RUN PAR ELEMENT DU CHAMP, CHACUN AVEC CE FORMAT : le modele pose le debut, l'instruction et la fin
                # du champ dans un seul run ; Word l'imprime au format de ce run, LibreOffice au corps du style du
                # paragraphe (12 pt noir), sauf quand chaque element a son run (essai du 06/10/2026)
                for r in list(p.iter(qn(u'w:r'))):
                    rpr = r.find(qn(u'w:rPr'))
                    elements = [e for e in r if e.tag in (qn(u'w:fldChar'), qn(u'w:instrText'))]
                    if len(elements) < 2 or len(elements) != len(r) - (rpr is not None):
                        continue
                    precedent = r
                    for e in elements[1:]:
                        nouveau = r.makeelement(qn(u'w:r'), {})
                        if rpr is not None:
                            nouveau.append(copy.deepcopy(rpr))
                        nouveau.append(e)            # l'element quitte le run d'origine
                        precedent.addnext(nouveau)
                        precedent = nouveau

    def proprietes(self, titre, date=None):
        u"""Le titre du document et sa date de production, ecrits dans ses proprietes (docProps/core.xml)."""
        self.titre_document, self.date_document = titre, date

    def _nettoyer_paquet(self):
        u"""LES PROPRIETES DU DOCUMENT NE PORTENT QUE LE LIVRABLE : son titre et sa date de production ; auteur,
        description, mots-cles et dernier modificateur vides. Les parties que le modele heritait du gabarit de
        l'outil de fabrication sont retirees : la source bibliographique vide (customXml), la vignette blanche et les
        proprietes etendues (docProps/app.xml, application et statistiques perimees), que Word recalcule a
        l'enregistrement."""
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        cp = self.doc.core_properties
        cp.title = self.titre_document or u''
        for champ in (u'author', u'comments', u'keywords', u'last_modified_by', u'subject', u'category'):
            setattr(cp, champ, u'')
        cp.revision = 1
        if self.date_document is not None:
            jour = self.date_document
            if not isinstance(jour, datetime.datetime):
                jour = datetime.datetime(jour.year, jour.month, jour.day)
            cp.created = cp.modified = jour
        for rId, rel in list(self.doc.part.rels.items()):
            if rel.reltype == RT.CUSTOM_XML:
                del self.doc.part.rels[rId]
        paquet = self.doc.part.package
        for rId, rel in list(paquet.rels.items()):
            if rel.reltype in (RT.THUMBNAIL, RT.EXTENDED_PROPERTIES):
                del paquet.rels[rId]

    def octets(self):
        body = self.doc.element.body
        if self.ancre is not None:
            nouveaux = [e for e in body if not any(e is a for a in self.avant)]
            for e in nouveaux:
                self.ancre._p.addprevious(e)
            body.remove(self.ancre._p)
        self._nettoyer_paquet()
        tampon = io.BytesIO()
        self.tpl.save(tampon)
        return tampon.getvalue()

    def _ajouter(self, xml):
        u"""Un element XML (w:p ou w:tbl) a la fin du corps, avant les proprietes de section."""
        from docx.oxml import parse_xml
        from docx.oxml.ns import qn
        el = parse_xml(xml)
        body = self.doc.element.body
        sect = body.find(qn(u'w:sectPr'))
        if sect is not None:
            sect.addprevious(el)
        else:
            body.append(el)
        return el

    def _dernier_est_tableau(self):
        from docx.oxml.ns import qn
        body = self.doc.element.body
        enfants = [e for e in body if e.tag != qn(u'w:sectPr')]
        return bool(enfants) and enfants[-1].tag == qn(u'w:tbl')

    # ---- texte
    def runs(self, p, texte, gras=False, italique=False, taille=None):
        u"""Ecrit le texte ; chaque marque AR devient « à remplir », en italique surligne ; ce qui suit une marque
        EXPOSANT dans son morceau s'ecrit en exposant."""
        from docx.enum.text import WD_COLOR_INDEX
        from docx.shared import Pt
        morceaux = texte.split(AR)
        for i, m in enumerate(morceaux):
            for j, s in enumerate(m.split(EXPOSANT)):
                if not s:
                    continue
                r = p.add_run(s)
                r.bold, r.italic = gras or None, italique or None
                if j:
                    r.font.superscript = True
                if taille:
                    r.font.size = Pt(taille)
            if i < len(morceaux) - 1:
                r = p.add_run(A_REMPLIR_INSECABLE)
                r.italic, r.bold = True, gras or None
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                if taille:
                    r.font.size = Pt(taille)
        return p

    def titre(self, texte, niveau, saut=False):
        u"""Un titre ; ceux de niveau 1 a 3 portent un signet et entrent au sommaire."""
        nom = u'Title' if niveau == 0 else u'Heading %d' % niveau
        p = self.doc.add_paragraph(style=self.style(nom))
        if saut:
            p.paragraph_format.page_break_before = True
        self.runs(p, texte, gras=self.style(nom) is None)
        if 1 <= niveau <= 3:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            k = len(self.titres) + 1
            signet = u'_Toc%08d' % (90000000 + k)
            debut, fin = OxmlElement(u'w:bookmarkStart'), OxmlElement(u'w:bookmarkEnd')
            debut.set(qn(u'w:id'), str(9000 + k))
            debut.set(qn(u'w:name'), signet)
            fin.set(qn(u'w:id'), str(9000 + k))
            ppr = p._p.find(qn(u'w:pPr'))
            if ppr is not None:
                ppr.addnext(debut)
            else:
                p._p.insert(0, debut)
            p._p.append(fin)
            self.titres.append((niveau, texte, signet))
        return p

    def para(self, texte=u'', gras=False, italique=False, taille=None, centre=False, style=None, justifie=False):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = self.doc.add_paragraph(style=self.style(style) if style else None)
        if centre:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif justifie:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return self.runs(p, texte, gras, italique, taille)

    def puce(self, texte):
        st = self.style(u'List Bullet')
        p = self.doc.add_paragraph(style=st)
        return self.runs(p, texte if st else u'· ' + texte)

    def saut_de_page(self):
        from docx.enum.text import WD_BREAK
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- sommaire
    def sommaire(self, intitule=u'Sommaire'):
        u"""La page du sommaire : son intitule, hors sommaire, puis la place du champ TOC, rempli par
        remplir_sommaire une fois tous les titres ecrits."""
        from docx.shared import Pt, RGBColor
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run(intitule)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(14), RGBColor(0x1F, 0x38, 0x64)
        self.place_sommaire = self.doc.add_paragraph()

    def remplir_sommaire(self, pages=None):
        u"""Le champ TOC \\o "1-3" \\h \\z \\u, son resultat en cache : une entree par titre de niveau 1 a 3, lien vers
        son signet et champ PAGEREF. Word recalcule le champ a l'ouverture (w:updateFields) ; pages, s'il est donne,
        porte le numero de page de chaque titre, mis en cache pour une lecture qui ne recalcule pas les champs."""
        from docx.oxml import parse_xml
        from docx.oxml.ns import qn
        if self.place_sommaire is None:
            return
        ids = set(s.get(qn(u'w:styleId')) for s in self.doc.styles.element.findall(qn(u'w:style')))

        def champ(instr, resultat=None):
            x = u'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve">%s' \
                u'</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>' % _xml(instr)
            if resultat is not None:
                x += resultat
            return x
        fin = u'<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        xmls = []
        for i, (niveau, texte, signet) in enumerate(self.titres):
            page = u'' if not pages or i >= len(pages) or pages[i] is None else str(pages[i])
            style = (u'<w:pStyle w:val="TOC%d"/>' % niveau) if (u'TOC%d' % niveau) in ids else u''
            ppr = (u'<w:pPr>%s<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="%d"/></w:tabs>'
                   u'<w:spacing w:before="%d" w:after="0" w:line="240" w:lineRule="auto"/>'
                   u'<w:ind w:left="%d" w:right="425"/>'
                   u'<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/><w:sz w:val="%d"/>'
                   u'<w:szCs w:val="%d"/></w:rPr></w:pPr>'
                   % (style, self.largeur_portrait, 120 if niveau == 1 else 0, (niveau - 1) * 284,
                      self.taille_sommaire * 2, self.taille_sommaire * 2))
            x = u'<w:p %s>%s' % (_W, ppr)
            if i == 0:
                x += champ(u' TOC \\o "1-3" \\h \\z \\u ')
            x += u'<w:hyperlink w:anchor="%s" w:history="1">' % signet
            x += _runs_xml([(texte, niveau == 1, False)], taille=self.taille_sommaire) + u'<w:r><w:tab/></w:r>'
            x += champ(u' PAGEREF %s \\h ' % signet,
                       _run_xml(page, niveau == 1, taille=self.taille_sommaire) if page else u'') + fin
            # le champ se ferme dans la derniere entree : aucun paragraphe de plus sous le sommaire
            x += u'</w:hyperlink>%s</w:p>' % (fin if i == len(self.titres) - 1 else u'')
            xmls.append(x)
        if not xmls:
            xmls.append(u'<w:p %s>%s%s</w:p>' % (_W, champ(u' TOC \\o "1-3" \\h \\z \\u '), fin))
        for x in xmls:
            self.place_sommaire._p.addprevious(parse_xml(x))
        # la place du sommaire disparait, sauf quand elle porte le saut de section qui suit le sommaire (orientation)
        ppr = self.place_sommaire._p.find(qn(u'w:pPr'))
        if ppr is None or ppr.find(qn(u'w:sectPr')) is None:
            self.place_sommaire._p.getparent().remove(self.place_sommaire._p)
        self.place_sommaire = None

    def champs_a_jour_a_l_ouverture(self):
        u"""w:updateFields dans settings.xml : Word met a jour le sommaire et ses numeros de page a l'ouverture.
        L'element se place a son rang du schema, avant w:hdrShapeDefaults, w:footnotePr, w:endnotePr, w:compat..."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        s = self.doc.settings.element
        if s.find(qn(u'w:updateFields')) is not None:
            return
        el = OxmlElement(u'w:updateFields')
        el.set(qn(u'w:val'), u'true')
        suivants = [u'w:hdrShapeDefaults', u'w:footnotePr', u'w:endnotePr', u'w:compat', u'w:docVars', u'w:rsids',
                    u'm:mathPr', u'w:attachedSchema', u'w:themeFontLang', u'w:clrSchemeMapping',
                    u'w:doNotIncludeSubdocsInStats', u'w:doNotAutoCompressPictures', u'w:forceUpgrade', u'w:captions',
                    u'w:readModeInkLockDown', u'w:smartTagType', u'sl:schemaLibrary', u'w:shapeDefaults',
                    u'w:doNotEmbedSmartTags', u'w:decimalSymbol', u'w:listSeparator']
        for e in s:
            if any(e.tag == qn(n) for n in suivants):
                e.addprevious(el)
                return
        s.append(el)

    def police_par_defaut(self, nom=u'Arial'):
        u"""La police par defaut du document (w:docDefaults) : celle du theme du modele, a defaut. Word la donne aux
        marques de paragraphe du sommaire qu'il regenere a l'ouverture ; elle devient Arial, comme tout le texte."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        racine = self.doc.styles.element
        dd = racine.find(qn(u'w:docDefaults'))
        if dd is None:
            return
        rpd = dd.find(qn(u'w:rPrDefault'))
        if rpd is None:
            rpd = OxmlElement(u'w:rPrDefault')
            dd.insert(0, rpd)
        rpr = rpd.find(qn(u'w:rPr'))
        if rpr is None:
            rpr = OxmlElement(u'w:rPr')
            rpd.append(rpr)
        f = rpr.find(qn(u'w:rFonts'))
        if f is None:
            f = OxmlElement(u'w:rFonts')
            rpr.insert(0, f)
        for k in list(f.attrib):
            del f.attrib[k]
        for k in (u'w:ascii', u'w:eastAsia', u'w:hAnsi', u'w:cs'):
            f.set(qn(k), nom)
        # les polices du theme (titres et corps) : Word les ecrit lui-meme sur les marques de paragraphe du sommaire
        # qu'il regenere ; elles deviennent Arial
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        try:
            theme = self.doc.part.part_related_by(RT.THEME)
        except KeyError:
            return
        xml = theme.blob.decode(u'utf-8')
        xml = re.sub(u'(<a:(?:major|minor)Font>\\s*<a:latin typeface=")[^"]*(")', u'\\g<1>%s\\g<2>' % nom, xml)
        theme._blob = xml.encode(u'utf-8')

    # ---- mise en page
    @staticmethod
    def _geometrie(sect, paysage, marge):
        u"""Le format A4 d'une section (w:pgSz, portrait ou paysage) et ses 4 marges de page (w:pgMar, twips)."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        sz = sect.find(qn(u'w:pgSz'))
        if sz is None:
            sz = OxmlElement(u'w:pgSz')
            sect.insert(0, sz)
        largeur, hauteur = (PAGE_A4[1], PAGE_A4[0]) if paysage else PAGE_A4
        sz.set(qn(u'w:w'), str(largeur))
        sz.set(qn(u'w:h'), str(hauteur))
        if paysage:
            sz.set(qn(u'w:orient'), u'landscape')
        elif sz.get(qn(u'w:orient')) is not None:
            del sz.attrib[qn(u'w:orient')]
        mar = sect.find(qn(u'w:pgMar'))
        if mar is None:
            mar = OxmlElement(u'w:pgMar')
            sz.addnext(mar)
        for cote in (u'w:top', u'w:bottom', u'w:left', u'w:right'):
            mar.set(qn(cote), str(marge))

    def mise_en_page(self, marge, corps):
        u"""LA MISE EN PAGE DES COMPTES ANNUELS (06/10/2026) : A4 portrait, les 4 marges de page a
        « marge » twips ; le corps de page (styles Normal et List Bullet) a « corps » points, et les entrees du sommaire
        (styles toc 1 a 3) avec lui, leur numero de page a droite de la justification. Les titres gardent leur taille."""
        from docx.oxml.ns import qn
        from docx.shared import Pt
        sect = self.doc.element.body.find(qn(u'w:sectPr'))
        self._geometrie(sect, False, marge)
        self.marge_page, self.paysage = marge, False
        self.largeur = self.largeur_portrait = PAGE_A4[0] - 2 * marge
        for nom in (u'Normal', u'List Bullet'):
            if self.style(nom):
                self.doc.styles[nom].font.size = Pt(corps)
        self.taille_sommaire = corps
        # les entrees du sommaire a interligne simple, sans espace apres : en 12 pt, ses 39 entrees tiennent sur sa
        # page, meme regenerees par Word a l'ouverture
        for s in self.doc.styles.element.findall(qn(u'w:style')):
            if s.get(qn(u'w:styleId')) in (u'TOC1', u'TOC2', u'TOC3'):
                for t in s.iter(qn(u'w:tab')):
                    t.set(qn(u'w:pos'), str(self.largeur_portrait))
                for z in list(s.iter(qn(u'w:sz'))) + list(s.iter(qn(u'w:szCs'))):
                    z.set(qn(u'w:val'), str(int(corps * 2)))
                for e in s.iter(qn(u'w:spacing')):
                    e.set(qn(u'w:after'), u'0')
                    e.set(qn(u'w:line'), u'240')
                    e.set(qn(u'w:lineRule'), u'auto')
                # un intitule long passe a la ligne avant le numero de page (retrait droit de 0,75 cm)
                for e in s.iter(qn(u'w:ind')):
                    e.set(qn(u'w:right'), u'425')

    def orientation(self, paysage):
        u"""LES PAGES QUI PORTENT UN TABLEAU LARGE SONT EN PAYSAGE (06/10/2026) : un saut de section, page
        suivante, ferme la section courante (ses proprietes, copiees dans un paragraphe d'une ligne de 1 pt) ; la
        section suivante prend l'orientation demandee, les memes marges, l'en-tete et le pied courants (ni page de
        garde, ni en-tete de 1re page)."""
        if paysage == self.paysage or self.marge_page is None:
            return
        import copy
        from docx.oxml import parse_xml
        from docx.oxml.ns import qn
        body = self.doc.element.body
        sect = body.find(qn(u'w:sectPr'))
        ppr = parse_xml(u'<w:pPr %s><w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'
                        u'<w:rPr><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr></w:pPr>' % _W)
        ppr.append(copy.deepcopy(sect))
        # le saut de section se pose sur le dernier paragraphe s'il est vide et sans mise en forme (la place du
        # sommaire) : aucune ligne de plus, aucune page blanche ; sinon sur un paragraphe d'une ligne de 1 pt
        precedent = sect.getprevious()
        if (precedent is not None and precedent.tag == qn(u'w:p') and len(precedent) == 0):
            precedent.append(ppr)
        else:
            p = parse_xml(u'<w:p %s/>' % _W)
            p.append(ppr)
            sect.addprevious(p)
        for e in list(sect):
            if e.tag == qn(u'w:titlePg') or (e.tag in (qn(u'w:headerReference'), qn(u'w:footerReference'))
                                              and e.get(qn(u'w:type')) == u'first'):
                sect.remove(e)
        self._geometrie(sect, paysage, self.marge_page)
        self.paysage = paysage
        self.largeur = (PAGE_A4[1] if paysage else PAGE_A4[0]) - 2 * self.marge_page

    def premiere_page_sans_entete(self):
        u"""La page de garde ne porte ni l'en-tete courant ni le numero de page : w:titlePg, en-tete et pied de 1re page
        vides ; la numerotation continue ensuite."""
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        sec.first_page_header.is_linked_to_previous = False
        sec.first_page_footer.is_linked_to_previous = False

    # ---- tableaux
    def tableau(self, entetes, lignes, largeurs, numeriques=(), taille=None):
        u"""entetes : textes ; lignes : dict(cellules=[texte ou None], gras, niveau, fusion) ; une cellule None est
        vide ; un texte qui porte AR est mis en evidence. numeriques : index des colonnes alignees a droite.
        largeurs : liste de largeurs en cm (grille simple), ou carte du modele du recueil (dict, cf. CARTES_ANNEXE) :
        le tableau prend alors la forme du modele, ou, pour une carte de forme texte, devient des paragraphes."""
        if isinstance(largeurs, dict):
            if largeurs.get(u'forme') == u'texte':
                return self._texte(lignes, largeurs)
            return self._modele(entetes, lignes, largeurs, numeriques)
        return self._grille(entetes, lignes, largeurs, numeriques, taille)

    def _texte(self, lignes, carte):
        u"""Une information redigee : un paragraphe par ligne, son accroche en gras, puis sa valeur ; une rubrique en
        accroche seule (intro_rubrique : la phrase du recueil, en maigre) ; une ligne sans valeur ne s'imprime pas ;
        « avant » porte les phrases d'introduction (puce : une puce maigre) ; sans_accroche : la valeur seule, apres
        les puces (puces) et la phrase qui les suit (suite). Une valeur None ne s'imprime pas (information sans objet) ;
        une valeur vide s'imprime vide, sous son texte du recueil."""
        from docx.shared import Pt
        premier = [True]

        def para(style=None, suite=False):
            u"""suite : le paragraphe annonce ce qui le suit et ne s'en separe pas en fin de page."""
            p = self.doc.add_paragraph(style=self.style(style) if style else None)
            if premier[0]:
                p.paragraph_format.space_before = Pt(6)
                premier[0] = False
            if suite:
                p.paragraph_format.keep_with_next = True
            return p
        for l in lignes:
            for x in l.get(u'avant') or []:
                if x.get(u'intro') and x.get(u'puce'):
                    self.puce(lettre(x[u'intro'])).paragraph_format.keep_with_next = True
                elif x.get(u'intro'):
                    self.runs(para(suite=True), lettre(x[u'intro']))
            cel = l.get(u'cellules') or []
            accroche = lettre(l.get(u'affiche') or (cel[0] if cel else u'') or u'')
            if l.get(u'fusion'):
                self.runs(para(suite=True), accroche, gras=not l.get(u'intro_rubrique'))
                continue
            valeur = cel[1] if len(cel) > 1 else None
            if valeur is None:
                continue
            if l.get(u'puces') or l.get(u'sans_accroche'):
                if not l.get(u'sans_accroche'):
                    self.runs(para(suite=True), accroche + u' :', gras=True)
                for t in l.get(u'puces') or []:
                    self.puce(lettre(t)).paragraph_format.keep_with_next = True
                if l.get(u'suite'):
                    self.runs(para(suite=True), lettre(l[u'suite']))
                self.runs(para(suite=bool(l.get(u'lie_au_suivant'))), valeur)
            else:
                p = para(u'List Bullet' if l.get(u'puce') else None)
                if l.get(u'puce') and not self.style(u'List Bullet'):
                    p.add_run(u'· ')
                self.runs(p, accroche, gras=True)
                self.runs(p, u' : ' + valeur)

    def _modele(self, entetes, lignes, carte, numeriques):
        u"""Un tableau a la forme de son modele au recueil (_forme_modele) : cadre et filets verticaux, filet sous
        l'en-tete, filets horizontaux la ou la ligne en porte (plein, fin ou pointille), aucun fond, intitules en gras
        centres, contenu en 10 pt, titre du modele en 12 pt gras bleu fonce, proportions des colonnes du modele dans la
        justification de la section (portrait ou paysage), hauteur des lignes du modele (h_entete, pas, h de la ligne,
        en points : hauteur minimale, texte centre verticalement sauf va contraire)."""
        f = _forme_modele(entetes, lignes, carte, numeriques, self.largeur)
        grille, w, marge, cadre, cadre_bas, V = (f[u'grille'], f[u'w'], f[u'marge'], f[u'cadre'], f[u'cadre_bas'],
                                                 f[u'V'])
        # ---- XML : chaque cellule porte ses 4 bords ; un filet entre 2 lignes est le bord bas de la 1re et le bord
        # haut de la 2de ; une cellule fusionnee verticalement n'a pas de filet interieur ; un filet borne a des
        # colonnes (filet_cols) ne passe que sous les cellules de ces colonnes
        def bord(cote, style):
            return u'<w:%s %s/>' % (cote, _FILETS[style])

        def filet_sous(r, a, b):
            if r.get(u'filet_cols') is not None and not set(range(a, b)) <= set(r[u'filet_cols']):
                return None
            return r[u'filet']
        nr = len(grille)
        lignes_xml = []
        for i, r in enumerate(grille):
            tcs = u''
            for c in r[u'cellules']:
                a, b = c[u'col'], c[u'col'] + c[u'span']
                h = cadre if i == 0 else filet_sous(grille[i - 1], a, b)
                bb = cadre_bas if i == nr - 1 else filet_sous(r, a, b)
                if c.get(u'vm') == u'restart':
                    bb = None
                if c.get(u'vm') == u'continue':
                    h = None
                    bb = cadre_bas if i == nr - 1 else r[u'filet']
                tcpr = u'<w:tcW w:w="%d" w:type="dxa"/>' % sum(w[a:b])
                if c[u'span'] > 1:
                    tcpr += u'<w:gridSpan w:val="%d"/>' % c[u'span']
                if c.get(u'vm'):
                    tcpr += u'<w:vMerge w:val="restart"/>' if c[u'vm'] == u'restart' else u'<w:vMerge/>'
                tcpr += u'<w:tcBorders>%s%s%s%s</w:tcBorders>' % (
                    bord(u'top', h), bord(u'left', cadre if a in V else None), bord(u'bottom', bb),
                    bord(u'right', cadre if b in V else None))
                if c.get(u'fond'):
                    tcpr += u'<w:shd w:val="clear" w:color="auto" w:fill="%s"/>' % c[u'fond']
                if c.get(u'marge') is not None:
                    tcpr += (u'<w:tcMar><w:left w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tcMar>'
                             % (marge, c[u'marge']))
                tcpr += u'<w:vAlign w:val="%s"/>' % (c.get(u'va') or u'top')
                runs = _runs_xml(c[u'segs'], c[u'taille'], c.get(u'couleur'), insecable=not c.get(u'coupable'))
                tcs += u'<w:tc><w:tcPr>%s</w:tcPr>%s</w:tc>' % (tcpr, _para_xml(
                    runs, c[u'jc'], c.get(u'retrait', 0), c.get(u'suspendu', 0), c.get(u'avant', 0),
                    bool(r.get(u'garder')), c.get(u'tab'), c[u'taille'], c.get(u'apres', 0), c.get(u'interligne')))
            trpr = u'<w:cantSplit/>'
            if r.get(u'h'):
                # la hauteur du recueil se mesure d'un filet au suivant, filet du haut compris ; Word ajoute
                # l'epaisseur du filet a la hauteur de la ligne : elle se deduit
                dessus = cadre if i == 0 else grille[i - 1][u'filet']
                trpr += u'<w:trHeight w:val="%d" w:hRule="atLeast"/>' % int(round(
                    (r[u'h'] - EPAISSEUR.get(dessus, 0)) * 20))
            if r.get(u'entete'):
                trpr += u'<w:tblHeader/>'
            lignes_xml.append(u'<w:tr><w:trPr>%s</w:trPr>%s</w:tr>' % (trpr, tcs))
        nil = u''.join(u'<w:%s w:val="nil"/>' % c for c in (u'top', u'left', u'bottom', u'right', u'insideH',
                                                             u'insideV'))
        xml = (u'<w:tbl %s><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/><w:tblBorders>%s'
               u'</w:tblBorders><w:tblLayout w:type="fixed"/><w:tblCellMar><w:left w:w="%d" w:type="dxa"/>'
               u'<w:right w:w="%d" w:type="dxa"/></w:tblCellMar><w:tblLook w:val="0000"/></w:tblPr><w:tblGrid>%s'
               u'</w:tblGrid>%s</w:tbl>' % (_W, sum(w), nil, marge, marge,
                                            u''.join(u'<w:gridCol w:w="%d"/>' % x for x in w), u''.join(lignes_xml)))
        if self._dernier_est_tableau():           # 2 tableaux qui se suivent ne se fondent pas en un seul
            self._ajouter(u'<w:p %s><w:pPr><w:spacing w:before="0" w:after="80"/><w:rPr><w:sz w:val="%d"/>'
                          u'<w:szCs w:val="%d"/></w:rPr></w:pPr></w:p>' % (_W, CORPS_TABLEAU * 2, CORPS_TABLEAU * 2))
        return self._ajouter(xml)

    def _grille(self, entetes, lignes, largeurs, numeriques=(), taille=None):
        u"""Le tableau quadrille des livrables autres que les comptes annuels (DIP)."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt
        n = len(largeurs)
        taille = taille or (7 if n >= 7 else 7.5 if n == 6 else 8 if n == 5 else 9)
        t = self.doc.add_table(rows=0, cols=n)
        if self.style(u'Table Grid'):
            t.style = self.doc.styles[u'Table Grid']
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, c in enumerate(t.columns):          # la grille : LibreOffice et Word la lisent avant les cellules
            c.width = Cm(largeurs[i])
        # autofit = False pose deja w:tblLayout fixed ; les marges de cellule se placent avant w:tblLook, ordre du
        # schema de Word (un element mal place rend le document illisible pour Word, pas pour LibreOffice)
        tblpr = t._tbl.tblPr
        marges = OxmlElement(u'w:tblCellMar')            # 0,1 cm a gauche et a droite : un montant tient sur 1 ligne
        for cote in (u'left', u'right'):
            m = OxmlElement(u'w:%s' % cote)
            m.set(qn(u'w:w'), u'57')
            m.set(qn(u'w:type'), u'dxa')
            marges.append(m)
        look = tblpr.find(qn(u'w:tblLook'))
        if look is not None:
            look.addprevious(marges)
        else:
            tblpr.append(marges)

        def ombrer(cellule, couleur):
            tcpr = cellule._tc.get_or_add_tcPr()
            shd = OxmlElement(u'w:shd')
            shd.set(qn(u'w:val'), u'clear')
            shd.set(qn(u'w:color'), u'auto')
            shd.set(qn(u'w:fill'), couleur)
            tcpr.append(shd)

        def ecrire(cellule, texte, gras, droite, niveau=0):
            p = cellule.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            if droite:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if niveau:
                p.paragraph_format.left_indent = Cm(0.35 * niveau)
            if texte is None:
                return
            if texte == AR:
                ombrer(cellule, u'FFF2CC')
                r = p.add_run(A_REMPLIR)
                r.italic, r.bold = True, gras or None
                r.font.size = Pt(taille)
                return
            morceaux = texte.split(AR)
            for i, m in enumerate(morceaux):
                for j, s in enumerate(m.split(EXPOSANT)):
                    if not s:
                        continue
                    r = p.add_run(s)
                    r.bold = gras or None
                    r.font.size = Pt(taille)
                    if j:
                        r.font.superscript = True
                if i < len(morceaux) - 1:
                    from docx.enum.text import WD_COLOR_INDEX
                    r = p.add_run(A_REMPLIR)
                    r.italic, r.bold = True, gras or None
                    r.font.size = Pt(taille)
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW

        def largeur(row):
            for i, c in enumerate(row.cells):
                c.width = Cm(largeurs[i])
            trpr = row._tr.get_or_add_trPr()
            cs = OxmlElement(u'w:cantSplit')
            trpr.append(cs)

        if entetes is not None:
            row = t.add_row()
            largeur(row)
            trpr = row._tr.get_or_add_trPr()
            th = OxmlElement(u'w:tblHeader')
            trpr.append(th)
            for i, e in enumerate(entetes):
                ombrer(row.cells[i], u'D9E2F3')
                ecrire(row.cells[i], e, True, i in numeriques)
        for l in lignes:
            row = t.add_row()
            largeur(row)
            cel = l[u'cellules']
            if l.get(u'fusion'):
                c = row.cells[0].merge(row.cells[n - 1])
                ecrire(c, cel[0], l.get(u'gras'), False, l.get(u'niveau') or 0)
                continue
            for i in range(n):
                ecrire(row.cells[i], cel[i] if i < len(cel) else None, l.get(u'gras'), i in numeriques,
                       (l.get(u'niveau') or 0) if i == 0 else 0)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t


def _largeurs(n, total=17.0, premiere=None):
    u"""Largeurs en cm : la 1re colonne prend le reste quand les autres sont des montants."""
    if n == 1:
        return [total]
    autre = {2: 4.0, 3: 3.2, 4: 3.0, 5: 2.6, 6: 2.3, 7: 2.05, 8: 1.85}.get(n, 1.8)
    p = premiere if premiere is not None else total - autre * (n - 1)
    reste = (total - p) / (n - 1)
    return [p] + [reste] * (n - 1)


# ============================================================================ comptes annuels

def _titre_etat(etat, arrete, devise):
    u"""Le titre du modele en texte simple (classeur) : la devise de comptabilite lue comme le Word la lit (_devise),
    « à remplir » quand la table ne la porte pas ou la porte a remplir, regle d'une valeur absente du classeur."""
    texte, cote = TITRES_ETATS[etat]
    return u'%s au %s en %s%s' % (texte, date_fr(arrete[u'date_arrete']), devise or a_remplir(),
                                  (u' - %s' % cote) if cote else u'')


def _montant_etat(l, n1=False):
    v = l[u'exercice_n_1' if n1 else u'exercice_n']
    if v is not None:
        return montant(v)
    if l[u'type_ligne'] == u'RUBRIQUE':
        return None
    return a_remplir()             # NULL hors rubrique : « à remplir » (regle 0.4), libelle de la base ou non


def _entete_pied(d, quoi, version, produit_le, separateur=u' ; '):
    e, a = d[u'entite'], d[u'arrete']
    return (u'%s%s%s au %s' % (e[u'denomination'], separateur, quoi, date_fr(a[u'date_arrete'])),
            u'Version %d produite le %s' % (int(version), date_fr(produit_le)))


def _titre_etat_segments(etat, arrete, devise):
    u"""Le titre du modele dans la 1re cellule de l'en-tete : en gras, la date de cloture et la devise de comptabilite
    en italique maigre, comme l'emplacement qu'elles remplissent au recueil (« Bilan au (date de clôture) en (devise de
    comptabilité) - Actif »)."""
    texte, cote = TITRES_ETATS[etat]
    return [(u'%s au ' % texte, True, False), (date_fr(arrete[u'date_arrete']), False, True), (u' en ', True, False),
            (devise, False, True)] + ([(u' - %s' % cote, True, False)] if cote else [])


def _bloc_etat(w, d, etat):
    u"""Le tableau du modele de l'etat (321-2, 321-6 ou 322-2), son titre dans la 1re cellule de l'en-tete, puis la
    note (1) du modele, collee sous le cadre."""
    c = CARTES_ETATS[etat]
    carte = dict(largeurs=c[u'largeurs'], h_entete=c.get(u'h_entete'), pas=c.get(u'pas'),
                 interligne_entete=c.get(u'interligne_entete'),
                 entete=[[dict(t=_titre_etat_segments(etat, d[u'arrete'], _devise(d)),
                               jc=u'left', titre=True), c[u'exercices'][0], c[u'exercices'][1]]])
    lignes = []
    for l in d[etat]:
        x = dict(cellules=[l[u'libelle_recueil'], _montant_etat(l), _montant_etat(l, True)], gras=l[u'gras'],
                 retrait=0 if l[u'gras'] else c[u'retrait'])
        x.update(c[u'lignes'].get(l[u'code']) or {})
        lignes.append(x)
    w.tableau([u'', u'Exercice N', u'Exercice N-1'], lignes, carte, numeriques=(1, 2))
    if etat in NOTES_ETATS:
        # la note (1) du modele, 9 pt, justifiee, collee sous le cadre (recueil, pages 91 et 92 du fichier)
        w.para(lettre(NOTES_ETATS[etat]), taille=CORPS_NOTE, justifie=True)


def a_remplir():
    u"""LA REGLE D'UNE VALEUR ABSENTE DU BILAN, DU COMPTE DE RESULTAT ET DU CLASSEUR : la marque AR, imprimee
    « à remplir » et mise en evidence. Jamais dans l'annexe, dont une cellule sans valeur s'imprime vide (_imprime)."""
    return AR


# ============================================================================ annexe : la table de la base

# La devise de comptabilite, cellule de l'article 311-1 de la table : les titres des etats et du 335-8, et la page de
# garde, impriment son texte.
ARTICLE_DEVISE, LIGNE_DEVISE = u'311-1', u'DEVISE'
# Les etats d'une cellule de la table qui n'impriment rien au Word : la case reste vide.
ETATS_VIDES = (u'A_REMPLIR', u'EN_ATTENTE', u'SANS_OBJET')


def _cellule(d, article, code, k):
    u"""La cellule de la table, dict(etat, texte), ou None quand la table ne la porte pas."""
    return (d.get(u'cellules') or {}).get((article, code, k))


def _sans_cellule(d, article, code, k):
    u"""Note une cellule que le modele imprime et que la table ne porte pas : d['sans_cellule'], liste de (article,
    ligne, colonne) ; « * » en ligne : une ligne du modele ou un article entier sans ligne dans la table."""
    x = d.setdefault(u'sans_cellule', [])
    if (article, code, k) not in x:
        x.append((article, code, k))


def _imprime(d, article, code, k):
    u"""LE TEXTE QU'UNE CELLULE DE L'ANNEXE IMPRIME : le texte imprime de la table, vide pour une cellule a remplir, en
    attente ou sans objet, sans surlignage ni fond. Une cellule que la table ne porte pas s'imprime vide et se note dans
    d['sans_cellule']."""
    c = _cellule(d, article, code, k)
    if c is None:
        _sans_cellule(d, article, code, k)
        return u''
    return c[u'texte'] or u''


def _valeur_texte(d, article, code):
    u"""LA VALEUR D'UNE INFORMATION REDIGEE (colonne 1) : None quand la cellule est sans objet, et la ligne ne
    s'imprime pas ; le texte imprime de la table sinon, vide pour une cellule a remplir ou en attente, la ligne
    s'imprimant alors avec le texte du recueil qui l'exige et une valeur vide."""
    c = _cellule(d, article, code, 1)
    if c is None:
        _sans_cellule(d, article, code, 1)
        return u''
    if c[u'etat'] == u'SANS_OBJET':
        return None
    return c[u'texte'] or u''


def _devise(d):
    u"""La devise de comptabilite telle que la table l'imprime : vide tant qu'elle est a remplir."""
    return _imprime(d, ARTICLE_DEVISE, LIGNE_DEVISE, 1)


def _colonne(src):
    u"""La colonne de la table que lit une colonne du modele : 'ech', l'echeance du 334-6, est la colonne 6."""
    return 6 if src == u'ech' else src


def _intitule(rangee, n):
    u"""UNE RANGEE D'INTITULE QUE LE MODELE AJOUTE (334-3 : « évaluation des redevances restant à payer », qui annonce
    ses 3 echeances) : son libelle, ses n - 1 autres cellules vides. Une autre rangee passe telle quelle."""
    if u'intitule' not in rangee:
        return rangee
    x = dict(rangee)
    x[u'cellules'] = [x.pop(u'intitule')] + [None for _ in range(n - 1)]
    return x


# Les 2 intitules d'un bloc d'information redigee : le Word l'imprime en paragraphes a accroche en gras (forme texte)
ENTETES_TEXTE = [u'Information', u'Texte de l’annexe']


def _index_bloc(article, b):
    u"""Le rang du bloc dans ANNEXE[article], qui designe sa carte ; None pour le bloc par defaut."""
    return next((i for i, x in enumerate(ANNEXE.get(article) or []) if x is b), None)


def _attributs(carte, code):
    u"""Les attributs de forme d'une ligne : ceux de la carte par defaut, puis ceux de sa ligne."""
    x = dict(carte.get(u'defaut') or {})
    x.update((carte.get(u'lignes') or {}).get(code) or {})
    return x


def _ordre_impression(article, blocs):
    u"""L'ORDRE D'IMPRESSION DES BLOCS, celui du recueil : un bloc dont la carte porte avant_bloc (rang d'un autre
    bloc) s'imprime juste avant ce bloc ; les autres dans l'ordre d'ANNEXE. Le rang d'un bloc, qui designe sa carte,
    ne change pas."""
    def avant(x):
        return (CARTES_ANNEXE.get((article, _index_bloc(article, x[0]))) or {}).get(u'avant_bloc')
    sortie = [x for x in blocs if avant(x) is None]
    for x in blocs:
        if avant(x) is not None:
            j = next((i for i, y in enumerate(sortie) if _index_bloc(article, y[0]) == avant(x)), len(sortie))
            sortie.insert(j, x)
    return sortie


def _libre(b):
    u"""Le bloc prend les lignes qu'aucun autre bloc ne nomme (texte, tableau ou phrase sans codes)."""
    return b[u'genre'] in (u'texte', u'tableau', u'phrase') and b.get(u'codes') is None


def _blocs_article(article, lignes):
    u"""Les blocs de l'article, chacun avec ses lignes de la table, dans l'ordre de la table : les lignes que le bloc
    nomme (codes), les paiements et le total d'un tableau des acomptes (prefixe), les contrats du 334-6 (CC:), toutes
    les lignes pour l'inventaire du 336-2 ; un bloc sans codes prend les lignes qu'aucun autre bloc ne nomme, ni
    comme ligne, ni comme information imprimee par son libelle (avant), ni comme occurrence (334-3, CB:)."""
    blocs = ANNEXE.get(article) or [X()]
    par_code = dict((l[u'code'], l) for l in lignes)

    def propres(b):
        g = b[u'genre']
        if g == u'sans_ligne':
            p = b[u'prefixe']
            return [l for l in lignes if l[u'code'].startswith(p + u':') or l[u'code'] == p + u'_TOTAL']
        if g == u'contrats':
            return [l for l in lignes if l[u'code'].startswith(u'CC:') or l[u'code'] in (b.get(u'codes') or [])]
        if g == u'une_ligne':
            return [par_code[c] for _, c, _ in b[u'colonnes'] if c in par_code]
        if g == u'inventaire':
            return list(lignes)
        return [par_code[c] for c in (b.get(u'codes') or []) if c in par_code]
    nommes = set()
    for b in blocs:
        nommes |= set(CODES_SANS_LIGNE.get(x) for x in b.get(u'avant') or [])
        occ = (CARTES_ANNEXE.get((article, _index_bloc(article, b))) or {}).get(u'occurrences')
        if occ:
            nommes |= set(l[u'code'] for l in lignes if l[u'code'].startswith(occ[u'prefixe']))
        if not _libre(b):
            nommes |= set(l[u'code'] for l in propres(b))
    return [(b, [l for l in lignes if l[u'code'] not in nommes] if _libre(b) else propres(b), par_code)
            for b in blocs]


def _ecrire_article(w, d, article, lignes):
    u"""Les blocs de l'article, chacun a la forme de son modele (CARTES_ANNEXE). Chaque cellule imprimee est le texte
    imprime de la table (_imprime, _valeur_texte) ; les lignes passees a w.tableau suivent les lignes de la table, une
    par ligne, les colonnes dans l'ordre du bloc, et portent la cle de chaque cellule (cles) ; les lignes que le
    modele ajoute (intitule de poste, occurrence) voyagent dans la cle « avant » de la ligne qu'elles precedent."""
    a = d[u'arrete']
    annee = int(date_iso(a[u'exercice'] or a[u'date_arrete'])[:4])
    principe = (d.get(u'principes') or {}).get(article)
    for b, ls, par_code in _ordre_impression(article, _blocs_article(article, lignes)):
        g = b[u'genre']
        k = _index_bloc(article, b)
        carte = dict(CARTES_ANNEXE.get((article, k)) or {})
        if principe and k in (0, None):
            # ARTICLE DE PRINCIPE : le texte de l'article, au mot pres ; aucune grille
            for alinea in principe:
                w.para(alinea)
        if g == u'phrase':
            if not principe:
                for l in ls:
                    w.para(lettre(l[u'libelle']))
        elif g == u'texte':
            _ecrire_texte(w, d, article, b, ls, par_code, carte, principe)
        elif g == u'tableau':
            _ecrire_tableau(w, d, article, b, k, ls, par_code, carte, annee)
        elif g == u'sans_ligne':
            _ecrire_acomptes(w, d, article, b, k, ls, carte)
        elif g == u'une_ligne':
            # 336-3 : le texte de l'article, a la lettre, puis l'inventaire sous l'intitule de sa rubrique du bilan
            _intro(w, TEXTE_336_3)
            cel = [_imprime(d, article, code, 1) for _, code, _ in b[u'colonnes']]
            cles = [(article, code, 1) for _, code, _ in b[u'colonnes']]
            rubrique = dict(cellules=[RUBRIQUE_336_3], fusion=True)
            carte.setdefault(u'largeurs', [40.0] + [60.0 / (len(cel) - 1)] * (len(cel) - 1))
            w.tableau([e for e, _, _ in b[u'colonnes']],
                      [dict(cellules=cel, cles=cles, donnees=True, avant=[rubrique])], carte, numeriques=(1, 2, 4))
        elif g == u'contrats':
            _ecrire_334_6(w, d, b, ls, par_code, carte)
        elif g == u'inventaire':
            _ecrire_336_2(w, d, ls)


def _ecrire_texte(w, d, article, b, ls, par_code, carte, principe):
    u"""UNE INFORMATION REDIGEE : le texte du recueil qui l'exige (accroche, phrase, puces), puis sa valeur, le texte
    imprime de la table. Une cellule sans objet ne s'imprime pas ; une cellule a remplir ou en attente s'imprime
    vide, sous son texte du recueil. Une information que le recueil exige (exiges) et que la table ne porte pas
    s'imprime vide et se note dans d['sans_cellule']."""
    exiges = b.get(u'exiges') or {}
    if not ls and not b.get(u'avant') and not exiges:
        return
    rows = []
    for lib in b.get(u'avant') or []:
        code = CODES_SANS_LIGNE.get(lib, lib)
        rows.append(dict(dict(cellules=[lib, _valeur_texte(d, article, code)], cles=[None, (article, code, 1)],
                              niveau=0), **_attributs(carte, lib)))
    for l in ls:
        x = _attributs(carte, l[u'code'])
        if x.get(u'intro_rubrique'):
            # UNE RUBRIQUE : la phrase du recueil qui l'annonce, sans valeur
            rows.append(dict(cellules=[l[u'libelle']], gras=True, fusion=True))
        elif principe:
            # ARTICLE DE PRINCIPE : titre et texte de l'article, rien d'autre
            rows.append(dict(cellules=[l[u'libelle'], None], niveau=0))
        else:
            rows.append(dict(cellules=[l[u'libelle'], _valeur_texte(d, article, l[u'code'])],
                             cles=[None, (article, l[u'code'], 1)], niveau=0))
        rows[-1].update(x)
    for code in b.get(u'codes') or []:
        if code in exiges and code not in par_code:
            rows.append(dict(dict(cellules=[None, _imprime(d, article, code, 1)], cles=[None, (article, code, 1)],
                                  niveau=0), **_attributs(carte, code)))
    w.tableau(ENTETES_TEXTE, rows, dict(carte, forme=u'texte'))


def _ecrire_tableau(w, d, article, b, k, ls, par_code, carte, annee):
    u"""UN TABLEAU DU MODELE : une ligne par ligne de la table, ses colonnes dans l'ordre du bloc, le texte imprime de
    chaque cellule ; le signe du 333-3 est celui que la table porte."""
    a = d[u'arrete']
    gras = set(GRAS_ANNEXE.get(article, []))
    cols = b[u'colonnes']
    entetes = [e.format(a0=annee, a1=annee - 1, a2=annee - 2, a3=annee - 3, a4=annee - 4) for e, _ in cols]
    if b.get(u'titre') and carte.get(u'titre_au_dessus'):
        # 332-1 : le tableau que le texte exige sans le dessiner, precede de la phrase qui l'exige, a la lettre
        _intro(w, b[u'titre'])
    elif b.get(u'titre'):
        # LE TITRE DU MODELE DANS LA 1RE CELLULE DE L'EN-TETE (335-8), date et devise substituees
        carte[u'titre'] = b[u'titre'].format(date=date_fr(a[u'date_arrete']), devise=_devise(d))
    if (article, k) in BASES_MODELES:
        # LE PAS DU MODELE, rangee par rangee, releve au recueil (BASES_MODELES) ; une rangee dont le libelle tient
        # sur 2 lignes au recueil porte l'interligne de sa 2e ligne (LIBELLES_2_LIGNES)
        bm = BASES_MODELES[(article, k)]
        deux = dict((c, dict(pas_2_lignes=p)) for c, p in (LIBELLES_2_LIGNES.get((article, k)) or {}).items())
        formes = dict((x[0], dict(_attributs(carte, x[0]), **(deux.get(x[0]) or {}))) for x in bm[0])
        carte[u'lignes'] = _maj(carte.get(u'lignes') or {}, _au_pas_du_recueil(*bm, formes=formes), deux)

    def cellules(code):
        cel, cles = [], []
        for _, src in cols[1:]:
            if src == u'sig':
                cel.append((par_code.get(code) or {}).get(u'signe') or None)
                cles.append((article, code, u'signe'))
            else:
                cel.append(_imprime(d, article, code, _colonne(src)))
                cles.append((article, code, _colonne(src)))
        return cel, cles
    exposant = article == u'333-3'
    rows = []
    for l in ls:
        cel, cles = cellules(l[u'code'])
        renvoi = exposant and l[u'code'] in (u'DISTRIB_PREC', u'AUTRES')
        x = dict(cellules=[l[u'libelle'] + (EXPOSANT + u'1' if renvoi else u'')] + cel, cles=[None] + cles,
                 gras=l[u'code'] in gras, niveau=0)
        x.update(_attributs(carte, l[u'code']))
        if x.get(u'avant'):
            x[u'avant'] = [_intitule(r, len(cols)) for r in x[u'avant']]
        if x.get(u'affiche') and renvoi:
            x[u'affiche'] += EXPOSANT + u'1'
        rows.append(x)
    # LES LIGNES DU MODELE QUI SUIVENT UNE LIGNE (suivantes, 334-3 : les 3 echeances du prix d'achat residuel) :
    # chacune se lit sur sa ligne de la table ; une ligne que la table ne porte pas s'imprime vide a sa place dans le
    # modele, apres la derniere ligne lue qui la precede
    par_ligne = dict((l[u'code'], x) for l, x in zip(ls, rows))
    for l, x in list(zip(ls, rows)):
        dernier = x
        for code in x.pop(u'suivantes', None) or []:
            if code in par_ligne:
                dernier = par_ligne[code]
                continue
            s = _attributs(carte, code)
            s.pop(u'suivantes', None)
            cel, cles = cellules(code)
            s[u'cellules'], s[u'cles'] = [None] + cel, [None] + cles
            dernier[u'apres'] = list(dernier.get(u'apres') or []) + [s]
    occ = carte.get(u'occurrences')
    cible = next((x for l, x in zip(ls, rows) if occ and l[u'code'] == occ[u'avant']), None)
    if cible is not None:
        # UNE LIGNE PAR OCCURRENCE (334-3 : une ligne par contrat), avant la ligne qu'elle precede : les lignes de la
        # table dont la cle porte le prefixe, chacune a la forme de la ligne du modele
        lues = [l for l in par_code.values() if l[u'code'].startswith(occ[u'prefixe'])]
        if not lues:
            _sans_cellule(d, article, occ[u'prefixe'] + u'*', None)
        av = []
        for l in lues:
            cel, cles = cellules(l[u'code'])
            av.append(dict(occ[u'forme'], cellules=[l[u'libelle']] + cel, cles=[None] + cles))
        cible[u'avant'] = av + list(cible.get(u'avant') or [])
    n = len(cols)
    if not rows:
        _sans_cellule(d, article, u'*', None)
        rows.append(dict(cellules=[None for _ in cols]))
    carte.setdefault(u'largeurs', [40.0] + [60.0 / (n - 1)] * (n - 1))
    w.tableau(entetes, rows, carte, numeriques=tuple(range(1, n)))
    if b.get(u'note'):
        # la note du modele, justifiee : 9 pt (333-2), 6 pt (333-3) (recueil, pages 98 et 99 du fichier)
        w.para(lettre(b[u'note']), taille=CORPS_NOTES.get(article, CORPS_NOTE), justifie=True)


def _ecrire_acomptes(w, d, article, b, k, ls, carte):
    u"""335-12, UN TABLEAU DES ACOMPTES : une ligne par paiement de la table (lignes prefixe:..., la date de paiement
    dans la 1re colonne, ou le constat « Néant »), puis « Total acomptes », la ligne prefixe_TOTAL de la table. La
    colonne k de la table est la colonne k du modele, 1 la date de paiement."""
    # les tableaux des acomptes sont larges : leurs pages sont en paysage (PAYSAGE)
    if (article, k) in PAYSAGE:
        w.orientation(True)
    cols = b[u'colonnes']
    n = len(cols)
    p = b[u'prefixe']
    dispo = [(1, 1, u'center')] + [(i, 1) for i in range(2, n)]
    rows = [dict(cellules=[None] + [_imprime(d, article, l[u'code'], j) for j in range(1, n)],
                 cles=[None] + [(article, l[u'code'], j) for j in range(1, n)], donnees=True, filet=None,
                 dispo=list(dispo)) for l in ls if l[u'code'].startswith(p + u':')]
    if not rows:
        _sans_cellule(d, article, p + u':*', None)
        rows.append(dict(cellules=[None for _ in cols], donnees=True, dispo=list(dispo)))
    rows[-1][u'filet'] = PLEIN
    total = p + u'_TOTAL'
    rows += [dict(cellules=[lib, None] + [_imprime(d, article, total, j) for j in range(2, n)],
                  cles=[None, None] + [(article, total, j) for j in range(2, n)], gras=gr, h=H_TOTAL_ACOMPTES.get(k),
                  dispo=[(0, 1)] + [(i, 1) for i in range(2, n)])
             for lib, gr in b[u'lignes']]
    w.tableau([e for e, _ in cols], rows, carte, numeriques=tuple(range(1, n)))


def _ecrire_334_6(w, d, b, ls, par_code, carte=None):
    u"""« Ventilation par contrat » (recueil, p. 101 du fichier) : une ligne par contrat de la table (lignes CC:),
    l'echeance de chacun (« Le cas echeant, pour les comptes courants, il est precise leur echeance a la cloture de
    l'exercice »), colonne 6 de la table ; sans contrat, la ligne COMPTES_COURANTS de la table ; puis la ligne TOTAL.
    Forme du modele : filets pointilles entre les contrats, filet plein avant « Total »."""
    art = u'334-6'
    cols = b[u'colonnes']
    ks = [_colonne(src) for _, src in cols[1:]]

    def ligne(l):
        return dict(cellules=[l[u'libelle']] + [_imprime(d, art, l[u'code'], k) for k in ks],
                    cles=[None] + [(art, l[u'code'], k) for k in ks])
    rows = [dict(ligne(l), donnees=True, filet=POINTILLE) for l in ls if l[u'code'].startswith(u'CC:')]
    if not rows and u'COMPTES_COURANTS' in par_code:
        rows.append(ligne(par_code[u'COMPTES_COURANTS']))
    if rows:
        rows[-1][u'filet'] = PLEIN
    if u'TOTAL' in par_code:
        rows.append(dict(_attributs(carte or {}, u'TOTAL'), gras=True, **ligne(par_code[u'TOTAL'])))
    carte = dict(carte or {}, centrees=(1,))
    carte.setdefault(u'largeurs', CARTES_ANNEXE[(u'334-6', 0)][u'largeurs'])
    if not rows:
        _sans_cellule(d, art, u'*', None)
    w.tableau([e for e, _ in cols], rows or [dict(cellules=[None for _ in cols])], carte,
              numeriques=tuple(range(2, len(cols))))


def _sous_titre(w, texte, recueil=True):
    u"""Un intitule de groupe de l'inventaire 336-2, en gras, detache du tableau qui precede ; a la lettre du recueil
    (lettre), ou, pour un texte qui n'en vient pas, a l'apostrophe typographique et a la capitale accentuee (typo)."""
    from docx.shared import Pt
    p = w.para((lettre if recueil else typo)(texte), gras=True)
    forme = getattr(p, u'paragraph_format', None)      # un document de releve (sans mise en page) rend autre chose
    if forme is not None:
        forme.space_before = Pt(8)
        forme.keep_with_next = True
    return p


# LES TIRETS DE L'ARTICLE 336-2, composes comme au recueil (page 106 du fichier) : tiret a 17,9 pt de la marge du texte,
# texte et lignes suivantes a 35,7 pt (liste a tiret suspendu) ; le 3e tiret (« - pour les autres actifs… », page 107)
# est au fer a gauche, comme au recueil
TIRET_SUSPENDU = (35.7, 17.8)


def _intro(w, texte, avant=6, tiret=False):
    u"""Le texte du recueil qui introduit un tableau sans titre dans sa 1re cellule, a la lettre, en maigre, lie au
    tableau qui le suit ; tiret : un alinea a tiret suspendu (TIRET_SUSPENDU), le tiret suivi d'une tabulation."""
    from docx.shared import Pt
    t = lettre(texte)
    if tiret:
        t = t.replace(u'- ', u'-	', 1)
    p = w.para(t)
    forme = getattr(p, u'paragraph_format', None)
    if forme is not None:
        forme.space_before = Pt(avant)
        forme.keep_with_next = True
        if tiret:
            gauche, suspendu = TIRET_SUSPENDU
            forme.left_indent = Pt(gauche)
            forme.first_line_indent = Pt(-suspendu)
            forme.tab_stops.add_tab_stop(Pt(gauche))
    return p


def _ecrire_336_2(w, d, ls):
    u"""L'inventaire detaille des actifs a caractere immobilier, par poste du bilan (336-1), lignes de la table : les
    immeubles et droits reels (lignes IMM:, puis IMM_TOTAL) sous le libelle de leur poste ; le tableau des parts des
    societes non cotees du modele, qui tient lieu de tableau des filiales et participations, chaque filiale (lignes
    FIL:) sous la ligne de groupe de la table qu'elle porte (une 3e ligne de groupe pour les filiales a classer), et sa
    note (1) ; les autres actifs a caractere immobilier (lignes AUT:), sous l'intitule de leur poste de 321-2, dans
    l'ordre du bilan."""
    art = u'336-2'
    par_code = dict((l[u'code'], l) for l in ls)

    def ligne(l, colonnes, **forme):
        return dict(cellules=[_imprime(d, art, l[u'code'], k) for k in colonnes],
                    cles=[(art, l[u'code'], k) for k in colonnes], donnees=True, **forme)
    # LE TEXTE DE L'ARTICLE, a la lettre : son chapeau, puis un tiret avant chaque tableau
    _intro(w, TEXTE_336_2)
    # IMMEUBLES ET DROITS REELS, sous le libelle de leur poste du bilan (336-1)
    _intro(w, POUR_IMMEUBLES, tiret=True)
    _sous_titre(w, TITRE_IMMEUBLES)
    rows = [ligne(l, range(1, 8), retrait=0) for l in ls if l[u'code'].startswith(u'IMM:')]
    if not rows:
        _sans_cellule(d, art, u'IMM:*', None)
        rows = [dict(cellules=[None for _ in ENTETES_IMMEUBLES], donnees=True, retrait=0)]
    rows[-1][u'filet'] = PLEIN
    if u'IMM_TOTAL' in par_code:
        # la ligne Total : son libelle dans la 1re colonne, les colonnes 2 a 7 de la ligne IMM_TOTAL de la table
        t = ligne(par_code[u'IMM_TOTAL'], range(2, 8), gras=True)
        rows.append(dict(t, cellules=[u'Total'] + t[u'cellules'], cles=[None] + t[u'cles'], donnees=False))
    w.tableau(ENTETES_IMMEUBLES, rows, CARTE_IMMEUBLES, numeriques=(1, 4, 5, 6))
    # FILIALES ET PARTICIPATIONS : le modele du recueil ; sous chaque ligne de groupe de la table, son intitule a la
    # lettre du recueil, puis les filiales qui portent ce groupe, sans filet entre les societes d'un groupe ; un groupe
    # dont la cellule n'est pas sans objet (« Néant ») imprime son texte sur une ligne de societe du modele, filets
    # verticaux gardes
    _intro(w, POUR_PARTS, avant=8, tiret=True)
    n = len(ENTETES_FILIALES)
    filiales = [l for l in ls if l[u'code'].startswith(u'FIL:')]

    def filiale(f):
        u"""La ligne d'une filiale : sa denomination (libelle de la table), puis les colonnes 1 a 6 de la table ; la
        colonne 7, son classement, est le groupe sous lequel elle s'imprime."""
        return dict(cellules=[f[u'libelle']] + [_imprime(d, art, f[u'code'], k) for k in range(1, n)],
                    cles=[None] + [(art, f[u'code'], k) for k in range(1, n)], niveau=1, retrait=0, donnees=True)
    rows = []
    for l in ls:
        if l[u'code'] not in GROUPES_FILIALES:
            continue
        if rows:
            rows[-1][u'filet'] = PLEIN             # filet au-dessus de chaque groupe
        rows.append(dict(cellules=[GROUPES_FILIALES[l[u'code']]], gras=False, fusion=True, h=H_GROUPE_FILIALES,
                         garder=True))
        c = _cellule(d, art, l[u'code'], 1)
        if c is None:
            _sans_cellule(d, art, l[u'code'], 1)
        elif c[u'etat'] != u'SANS_OBJET':
            rows.append(dict(cellules=[c[u'texte'] or u''] + [None for _ in range(n - 1)],
                             cles=[(art, l[u'code'], 1)] + [None for _ in range(n - 1)], donnees=True, fusion=True,
                             niveau=1, retrait=0, dispo=[(0, 1)] + [(j, 1) for j in range(1, n)]))
        rows += [filiale(f) for f in filiales if f.get(u'groupe') == l[u'code']]
    for f in filiales:
        if f.get(u'groupe') not in GROUPES_FILIALES or f.get(u'groupe') not in par_code:
            # une filiale dont la table ne porte pas la ligne de groupe : imprimee a la suite, notee
            _sans_cellule(d, art, f.get(u'groupe') or u'*', None)
            rows.append(filiale(f))
    if not rows:
        _sans_cellule(d, art, u'FIL:*', None)
        rows = [dict(cellules=[None for _ in range(n)], donnees=True, niveau=1, retrait=0)]
    rows[-1][u'garder'] = True                     # la note (1) reste sous le tableau
    w.tableau(ENTETES_FILIALES, rows, CARTE_FILIALES, numeriques=(4, 5, 6))
    w.para(lettre(NOTE_FILIALES), italique=True, taille=CORPS_NOTE, justifie=True)   # note (1) : 9 pt, italique, justifiee
    # AUTRES ACTIFS A CARACTERE IMMOBILIER : un intitule par poste restant de l'actif immobilier de 321-2, dans l'ordre
    # du bilan (tous les postes restants recoivent leur inventaire), chacun suivi des lignes AUT: de la table qui portent
    # ce poste (colonne groupe) ; les lignes de la table sans poste (le constat « Néant » de l'ensemble, l'actif a
    # inventorier) sous le dernier poste. Le libelle du poste a la lettre de 321-2, appel de note « (1) » compris.
    _intro(w, POUR_AUTRES, avant=8)
    libelles = dict((c, lib) for c, lib, _ in MODELES_ETATS[u'BILAN_ACTIF'])
    autres = [l for l in ls if l[u'code'].startswith(u'AUT:')]
    sans_poste = [l for l in autres if l.get(u'groupe') not in POSTES_AUTRES_IMMO]
    # L'INVENTAIRE QUE LA TABLE ETABLIT : la table porte une ligne par actif de chaque poste non nul, et aucun poste nul
    # n'a de ligne ; elle l'etablit tant qu'elle ne porte aucune ligne a inventorier (cle <prefixe>:SAISIE). Un poste
    # sans ligne y est alors un poste sans actif, et sa ligne du modele s'imprime « Néant » (ligne gardee sans
    # occurrence) ; sinon cette ligne est une case vide, notee dans d['sans_cellule'].
    etabli = bool(autres) and not any(l[u'code'].endswith(u':SAISIE') for l in autres)
    n = len(COLONNES_AUTRES_IMMO)
    rows = []
    for poste in POSTES_AUTRES_IMMO:
        if rows:
            rows[-1][u'filet'] = PLEIN             # filet sous le bloc de chaque poste
        rows.append(dict(cellules=[libelles[poste]], fusion=True, garder=True))   # l'intitule reste avec sa ligne
        lp = [l for l in autres if l.get(u'groupe') == poste]
        if poste == POSTES_AUTRES_IMMO[-1]:
            lp += sans_poste
        if lp:
            rows += [ligne(l, range(1, n + 1)) for l in lp]
        elif etabli:
            rows.append(dict(cellules=[NEANT] + [None for _ in range(n - 1)], donnees=True,
                             sans_occurrence=(art, poste)))
        else:
            _sans_cellule(d, art, u'AUT:' + poste, None)
            rows.append(dict(cellules=[None for _ in range(n)], donnees=True))
    w.tableau([e for e, _, _ in COLONNES_AUTRES_IMMO], rows, CARTE_AUTRES_IMMO, numeriques=(1, 2, 3))


def _texte_seul(article):
    u"""L'article n'a aucun tableau : ses blocs sont du texte (phrases, informations redigees)."""
    return all(b[u'genre'] in (u'texte', u'phrase') for b in ANNEXE.get(article) or [X()])


def fabriquer_comptes_annuels_docx(d, version, modele, produit_le, pages=None):
    u"""Le Word des comptes annuels : page de garde (sans en-tete ni numero), sommaire automatique (champ TOC, mis a
    jour par Word a l'ouverture), bilan (actif, passif), compte de resultat, annexe par section et par article. pages :
    numeros de page des titres, dans leur ordre, mis en cache dans le sommaire pour une lecture qui ne recalcule pas
    les champs ; sans eux, Word les calcule a l'ouverture."""
    e, a = d[u'entite'], d[u'arrete']
    entete, pied = _entete_pied(d, u'comptes annuels', version, produit_le, separateur=u' - ')
    w = _Word(modele, entete, pied)
    # PROPRIETES DU DOCUMENT : le titre « Comptes annuels de l'OPCI <entite> au <date> », la denomination lue en base
    # (« OPCI OMEGA » : « de l'OPCI OMEGA ») ; aucune trace de l'outil de fabrication (_Word._nettoyer_paquet)
    den = (e[u'denomination'] or u'').strip()
    w.proprietes(u'Comptes annuels de l’%s au %s' % (den if den.upper().startswith(u'OPCI ') else u'OPCI ' + den,
                                                     date_fr(a[u'date_arrete'])), produit_le)
    w.premiere_page_sans_entete()
    w.champs_a_jour_a_l_ouverture()
    w.police_par_defaut(u'Arial')
    # MISE EN PAGE (06/10/2026) : marges de page de 1,5 cm, corps de page en 12 pt
    w.mise_en_page(MARGE_PAGE, CORPS_PAGE)
    # PAGE DE GARDE : chaque element est lu en base ; aucun n'entre au sommaire
    w.titre(e[u'denomination'], 0)
    w.para(ou_ar(e.get(u'forme_vehicule')), gras=True)
    w.para(u'SIREN %s' % ou_ar(e.get(u'siren')))
    adr = u', '.join(x for x in (e.get(u'adresse_1'), e.get(u'adresse_2'),
                                 u' '.join(y for y in (e.get(u'code_postal'), e.get(u'ville')) if y)) if x)
    w.para(ou_ar(adr or None))
    w.para(u'')
    p = w.para(typo(d[u'livrable'][u'libelle']), gras=True, taille=14)   # nom du livrable, lu en base
    for r in p.runs:
        from docx.shared import RGBColor
        r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    w.para(u'Exercice clos le %s' % date_fr(a[u'date_arrete']), gras=True)
    w.para(REGLEMENT[:1].upper() + REGLEMENT[1:])
    w.para(u'Devise de comptabilité : %s' % _devise(d))
    w.para(u'Version %d produite le %s' % (int(version), date_fr(produit_le)))
    # SOMMAIRE
    w.sommaire()
    # BILAN (321-2, 321-6), COMPTE DE RESULTAT (322-2) : titres de navigation courts, le titre long du modele est
    # dans la 1re cellule de son tableau
    # les 3 etats en portrait, comme le modele (PAYSAGE_SUR_DECISION les mettrait en paysage)
    w.orientation(u'BILAN_ACTIF' in PAYSAGE_SUR_DECISION)
    w.titre(u'Bilan', 1, saut=True)
    w.titre(u'Actif du bilan', 2)
    _bloc_etat(w, d, u'BILAN_ACTIF')
    w.orientation(u'BILAN_PASSIF' in PAYSAGE_SUR_DECISION)
    w.titre(u'Passif du bilan', 2, saut=True)
    _bloc_etat(w, d, u'BILAN_PASSIF')
    w.orientation(u'RESULTAT' in PAYSAGE_SUR_DECISION)
    w.titre(u'Compte de résultat', 1, saut=True)
    _bloc_etat(w, d, u'RESULTAT')
    # ANNEXE : les articles 331-1 a 336-3 de ref_tableau_annexe et de la table, dans l'ordre de leur numero, sous leur
    # intitule du recueil ; leurs cellules sont celles de la table, a la version lue (lire_annexe)
    articles = sorted(set(x for x in list(d[u'tableaux']) + list(d.get(u'annexe') or {}) if x[:2] == u'33'),
                      key=lambda x: (int(x[:3]), int(x[4:])))
    # LES PAGES DES TABLEAUX LARGES EN PAYSAGE, LE RESTE EN PORTRAIT : l'article d'un tableau large (PAYSAGE, ou
    # PAYSAGE_SUR_DECISION) ouvre une section en paysage, quels que soient les montants du dossier ; un article sans
    # tableau (texte seul) suivi d'une partie en paysage reste dans l'orientation de ce qui le precede (334-1 apres
    # 333-3, 336-1 apres les acomptes : une ou 2 phrases, qu'une section a part mettrait seules sur une page) ; tout
    # autre article revient au portrait
    large = dict((art, art in PAYSAGE or art in PAYSAGE_SUR_DECISION) for art in articles)

    def orienter(i):
        art = articles[i]
        suivant = articles[i + 1] if i + 1 < len(articles) else None
        if large[art]:
            w.orientation(True)
        elif not (_texte_seul(art) and suivant and large[suivant]):
            w.orientation(False)
    if articles:
        orienter(0)
    w.titre(u'Annexe', 1, saut=True)
    section = None
    for i, art in enumerate(articles):
        orienter(i)
        if art[:3] != section:
            section = art[:3]
            w.titre(lettre(dict(SECTIONS_ANNEXE).get(section, u'Section')), 2)
        w.titre(u'Article %s %s' % (art, lettre(TITRES_ARTICLES.get(art) or d[u'tableaux'].get(art, u''))), 3)
        lignes = (d.get(u'annexe') or {}).get(art) or []
        if lignes or art in (d.get(u'principes') or {}):
            _ecrire_article(w, d, art, lignes)
        else:
            # un article sans ligne dans la table : son titre seul, l'article se note dans d['sans_cellule']
            _sans_cellule(d, art, u'*', None)
    w.remplir_sommaire(pages)
    return w.octets()


def fabriquer_comptes_annuels_xlsx(d, version, produit_le):
    u"""Le classeur a 3 onglets, Bilan, Compte de resultat, Balance ; l'annexe n'est pas un onglet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    FMT = u'#,##0.00;-#,##0.00'
    JAUNE = PatternFill(u'solid', fgColor=u'FFF2CC')
    BLEU = PatternFill(u'solid', fgColor=u'D9E2F3')
    a = d[u'arrete']
    # LE CLASSEUR PORTE LE BILAN, LE COMPTE DE RESULTAT ET LA BALANCE : le libelle du livrable (« Comptes annuels :
    # bilan, compte de résultat et annexe ») decrit le Word ; le classeur garde son nom, avant « : », et dit son contenu
    lib = d[u'livrable'][u'libelle'] or u''
    libelle = (lib.split(u' : ')[0] if u' : ' in lib else lib) + u' : bilan, compte de résultat et balance'
    wb = Workbook()
    wb.properties.title = nom_fichier(u'COMPTES_ANNUELS', d[u'entite'][u'code'], a[u'arrete'], version, u'xlsx')
    wb.properties.subject = u'%s ; version %d produite le %s' % (libelle, int(version), date_fr(produit_le))
    wb.properties.creator = u'fn_documents'

    def titre(ws, r, texte, ncol):
        u"""Le titre du modele en texte simple ; quand il porte « à remplir » (la devise), la cellule est mise en
        evidence comme une cellule a remplir. Le texte enrichi d'openpyxl n'est pas rendu par LibreOffice (essai du
        04/10/2026) : il n'est pas employe."""
        c = ws.cell(r, 1, texte.replace(AR, A_REMPLIR))
        c.font = Font(bold=True, size=11)
        if AR in texte:
            c.fill = JAUNE

    def entete(ws, r, noms, montants_des=2):
        for j, n in enumerate(noms, 1):
            c = ws.cell(r, j, n)
            c.font, c.fill = Font(bold=True), BLEU
            if j >= montants_des:
                c.alignment = Alignment(horizontal=u'right')

    def valeur(c, v, gras):
        if v is None:
            return
        if v == AR:
            c.value, c.font, c.fill = A_REMPLIR, Font(italic=True, bold=gras), JAUNE
            c.alignment = Alignment(horizontal=u'right')
            return
        c.value, c.number_format, c.font = _dec(v), FMT, Font(bold=gras)

    def etat(ws, r, et):
        titre(ws, r, _titre_etat(et, a, _devise(d)), 3)
        entete(ws, r + 1, [u'Ligne du modèle', u'Exercice N', u'Exercice N-1'])
        r += 2
        for l in d[et]:
            ws.cell(r, 1, l[u'libelle_recueil']).font = Font(bold=l[u'gras'])
            for j, k in ((2, u'exercice_n'), (3, u'exercice_n_1')):
                v = l[k]
                if v is None:
                    v = None if l[u'type_ligne'] == u'RUBRIQUE' else AR
                valeur(ws.cell(r, j), v, l[u'gras'])
            r += 1
        if et in NOTES_ETATS:
            ws.cell(r, 1, NOTES_ETATS[et]).font = Font(italic=True, size=9)
            r += 1
        return r

    # L'IDENTIFICATION DU DOSSIER, en ligne 1 de chaque onglet et en en-tete et pied d'impression : l'entite, le
    # livrable, la date de cloture, la version de la demande et la date de production
    e = d[u'entite']
    ident = u'%s ; %s ; exercice clos le %s ; version %d produite le %s' % (
        e[u'denomination'], libelle, date_fr(a[u'date_arrete']), int(version), date_fr(produit_le))

    def identifier(ws):
        ws.cell(1, 1, ident).font = Font(bold=True, size=10)
        ws.oddHeader.center.text = (u'%s ; comptes annuels au %s' % (e[u'denomination'], date_fr(a[u'date_arrete']))
                                    ).replace(u'&', u'&&')
        ws.oddFooter.left.text = u'Version %d produite le %s' % (int(version), date_fr(produit_le))
        ws.oddFooter.right.text = u'Page &P / &N'

    ws = wb.active
    ws.title = u'Bilan'
    identifier(ws)
    r = etat(ws, 2, u'BILAN_ACTIF')
    etat(ws, r + 1, u'BILAN_PASSIF')
    ws.freeze_panes = u'A4'
    ws2 = wb.create_sheet(u'Compte de résultat')
    identifier(ws2)
    etat(ws2, 2, u'RESULTAT')
    ws2.freeze_panes = u'A4'
    for w_ in (ws, ws2):
        w_.column_dimensions[u'A'].width = 78
        w_.column_dimensions[u'B'].width = w_.column_dimensions[u'C'].width = 18
    # BALANCE : v_balance_etats_financiers ; le total est une formule du tableur, la fonction ne somme rien
    ws3 = wb.create_sheet(u'Balance')
    identifier(ws3)
    titre(ws3, 2, u'Balance au %s ; colonne N-1 : clôture du %s ; solde : débit moins crédit'
          % (date_fr(a[u'date_arrete']), date_fr(a.get(u'date_arrete_n_1')) or AR), 6)
    entete(ws3, 3, [u'Compte', u'Libellé', u'Débit', u'Crédit', u'Solde N', u'Solde N-1'], 3)
    r = premiere = 4
    for b in d[u'balance']:
        ws3.cell(r, 1, b[u'compte_num'])
        ws3.cell(r, 2, b[u'libelle'] if b[u'libelle'] else A_REMPLIR)
        for j, k in ((3, u'debit'), (4, u'credit'), (5, u'solde_n'), (6, u'solde_n_1')):
            valeur(ws3.cell(r, j), AR if b[k] is None else b[k], False)
        r += 1
    if r == premiere:
        for j in range(1, 7):
            valeur(ws3.cell(r, j), AR, False)
        r += 1
    else:
        ws3.cell(r, 2, u'Total').font = Font(bold=True)
        for j in range(3, 7):
            col = get_column_letter(j)
            c = ws3.cell(r, j, u'=SUM(%s%d:%s%d)' % (col, premiere, col, r - 1))
            c.number_format, c.font = FMT, Font(bold=True)
    ws3.freeze_panes = u'A4'
    for f_ in (ws, ws2, ws3):                          # impression : 1 page de large, titres de colonnes repetes
        f_.page_setup.orientation = u'portrait' if f_ is not ws3 else u'landscape'
        f_.page_setup.paperSize = f_.PAPERSIZE_A4
        f_.sheet_properties.pageSetUpPr.fitToPage = True
        f_.page_setup.fitToWidth, f_.page_setup.fitToHeight = 1, 0
    ws3.column_dimensions[u'A'].width = 12
    ws3.column_dimensions[u'B'].width = 70
    for col in u'CDEF':
        ws3.column_dimensions[col].width = 17
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


# ============================================================================ rapport de l'expert-comptable

def _norm(t):
    return re.sub(r'\s+', u' ', (t or u'').replace(u'’', u"'").replace(u' ', u' ').replace(u' ', u' ')).strip()


def texte_forme(d, forme):
    u"""L'intitule, l'exemple et les paragraphes de la forme, lus en base : ceux de v_rapport_expert_comptable pour la
    forme de son texte (l'arretee, ou la proposee tant qu'aucune ne l'est), ceux de ref_forme_rapport pour une autre.
    Un paragraphe par ligne du texte ; les paragraphes « · » forment la liste des donnees."""
    f = (d.get(u'formes') or {}).get(forme)
    if f is None:
        raise Refus(u'Production refusée : la forme arrêtée du rapport de l’expert-comptable n’est aucune des formes '
                    u'de la NP 2300 que porte la base.')
    v = d.get(u'rapport') or {}
    src = v if v.get(u'forme_du_texte') == forme else f
    if not (src.get(u'texte') or u'').strip() or not (src.get(u'intitule') or u'').strip():
        raise Refus(u'Production refusée : la base ne porte pas le texte de la forme arrêtée du rapport de '
                    u'l’expert-comptable%s.' % ((u', exemple %s de la NP 2300' % f[u'exemple']) if f.get(u'exemple')
                                                else u''))
    paragraphes = [p.strip() for p in re.split(u'[\\r\\n]+', src[u'texte']) if p.strip()]
    return dict(code=forme, exemple=f.get(u'exemple'), intitule=src[u'intitule'].strip(), paragraphes=paragraphes)


def remplir_rapport(f, d):
    u"""Les paragraphes de la forme, chaque emplacement rempli par sa colonne de v_rapport_expert_comptable, « à
    remplir » quand elle est NULL. Aucun mot ajoute au texte : un emplacement prend la donnee seule. Rend une liste de
    (genre, texte), genre 'p' ou 'puce'."""
    v = d[u'rapport']
    mot = v.get(u'comptes_annuels_ou_intermediaires')
    jour = date_fr(v.get(u'date_arrete'))
    if mot == u'annuels' and jour:
        periode = u'l’exercice clos le %s' % jour
    elif mot == u'intermédiaires' and jour:
        periode = u'la période close le %s' % jour
    else:
        periode = AR
    lettre = date_fr(v.get(u'lettre_de_mission_du'))
    entite = v.get(u'entite_concernee')
    donnees = dict((_norm(lib), col) for lib, col in DONNEES_RAPPORT)
    sortie, vus = [], dict(comptes=0, entite=0, periode=0)
    for t in f[u'paragraphes']:
        t, n = EMPLACEMENTS[u'comptes'].subn(lambda m: u'comptes ' + ou_ar(mot), t)
        vus[u'comptes'] += n
        t, n = EMPLACEMENTS[u'entite'].subn(lambda m: de_elide(entite) if entite else u'de ' + AR, t)
        vus[u'entite'] += n
        t, n = EMPLACEMENTS[u'periode'].subn(lambda m: u'relatifs à ' + periode, t)
        vus[u'periode'] += n
        t = EMPLACEMENTS[u'lettre'].sub(lambda m: u'en date du ' + ou_ar(lettre), t)
        t = EMPLACEMENTS[u'alternative'].sub(lambda m: ou_ar(mot), t)
        puce = re.match(u'^·[\\s  ]*(.+?)[\\s  ]*([;.])$', t)
        if puce:
            col = donnees.get(_norm(puce.group(1)))
            if col is None:
                raise Refus(u'Production refusée : la base ne porte pas la donnée %s de l’exemple %s de la NP 2300.'
                            % (guillemets(puce.group(1)), f[u'exemple']))
            sortie.append((u'puce', u'%s : %s%s' % (puce.group(1), ou_ar(montant(v.get(col))),
                                                    u' ;' if puce.group(2) == u';' else u'.')))
        elif t.startswith(u'(Description') and t.endswith(u')'):
            sortie.append((u'p', ou_ar(v.get(u'description_des_points'))))
        elif _norm(t) == _norm(SIGNATURE):
            sortie.append((u'p', u'%s : %s' % (t, ou_ar(v.get(u'lieu_date_signature')))))
        else:
            sortie.append((u'p', t))
    reste = [t for g, t in sortie if u'…' in t or u'préciser' in t]
    if any(n != 1 for n in vus.values()) or reste or sum(1 for g, _ in sortie if g == u'puce') != len(DONNEES_RAPPORT):
        raise Refus(u'Production refusée : un emplacement à remplir de l’exemple %s de la NP 2300 est introuvable dans '
                    u'le texte que porte la base.' % f[u'exemple'])
    return sortie


def note_donnees_rapport(d):
    u"""Hors du texte de l'exemple, apres la signature : la devise de comptabilite, lue comme les comptes annuels la
    lisent (_devise, cellule 311-1 de la table), « à remplir » quand elle y manque, et la lecture des 2 donnees que le
    reglement ne nomme pas ; celle du chiffre d'affaires est lue en base (03/10/2026)."""
    v = d[u'rapport']
    return (u'Note sur les données : les montants sont exprimés en ' + ou_ar(_devise(d)) + u', devise de '
            u'comptabilité de l’OPCI. '
            u'Chiffre d’affaires : ' + ou_ar(v.get(u'chiffre_d_affaires_lecture')) + u' ; le ' + REGLEMENT
            + u', ne définit pas de chiffre d’affaires. Résultat net comptable : ligne C du modèle de compte de '
            u'résultat de l’article 322-2, « Résultat de l’exercice ».')


def fabriquer_rapport_docx(d, forme, version, modele, produit_le):
    f = texte_forme(d, forme)
    paragraphes = remplir_rapport(f, d)
    entete, pied = _entete_pied(d, u'rapport de l’expert-comptable', version, produit_le)
    w = _Word(modele, entete, pied)
    w.proprietes(f[u'intitule'], produit_le)
    w.para(f[u'intitule'], gras=True, centre=True, taille=12)
    w.para(u'')
    for genre, t in paragraphes:
        if genre == u'puce':
            w.puce(t)
        else:
            w.para(t)
    w.para(u'')
    w.para(note_donnees_rapport(d), italique=True, taille=8)
    return w.octets()


# ============================================================================ DIP

def _lignes(rubrique):
    u"""Les lignes de la rubrique qui portent une donnee : la vue marque « à remplir » (colonne etat) la ligne qu'elle
    rend quand sa source n'a rien a l'arrete."""
    return [x for x in rubrique if not x.get(u'etat')]


def _valeur(rubrique, cle):
    u"""Une valeur de l'arrete que la vue repete sur chaque ligne de la rubrique : celle de la 1re."""
    return rubrique[0].get(cle) if rubrique else None


def fabriquer_dip_docx(d, version, modele, produit_le):
    e, a = d[u'entite'], d[u'arrete']
    jour = date_fr(a[u'date_arrete'])
    entete, pied = _entete_pied(d, u'document d’information périodique', version, produit_le)
    w = _Word(modele, entete, pied)
    w.proprietes(u'Document d’information périodique de l’%s au %s' % (e[u'denomination'], jour), produit_le)
    w.titre(u'Document d’information périodique de l’%s au %s' % (e[u'denomination'], jour), 0)
    w.titre(TITRES_DIP[u'bandeau'], 1)
    w.para(TITRES_DIP[u'sous'])
    # CARTES 1 a 3 et 6 : valeur nette d'inventaire (rubrique 1), parts (2), valeur par part (3), dividende par part (6)
    w.titre(TITRES_DIP[u'cartes'], 2)
    w.tableau(None, [dict(cellules=[CARTES_DIP[0], ou_ar(montant(_valeur(d[u'r1'], u'valeur_nette_inventaire')))]),
                     dict(cellules=[CARTES_DIP[1], ou_ar(nombre(_valeur(d[u'r2'], u'nombre_parts')))]),
                     dict(cellules=[CARTES_DIP[2], ou_ar(montant(_valeur(d[u'r3'], u'valeur_par_part')))]),
                     dict(cellules=[CARTES_DIP[3], ou_ar(montant(_valeur(d[u'r6'], u'distribution_par_part')))])],
              [10.0, 7.0], numeriques=(1,), taille=9)
    # RUBRIQUE 1
    w.titre(TITRES_DIP[u'r1'], 2)
    rows = [dict(cellules=[ou_ar(x[u'famille_libelle']), ou_ar(nombre(x[u'nombre_actifs'])),
                           ou_ar(montant(x[u'valeur_actuelle'])),
                           ou_ar(montant(x[u'difference_estimation']))]) for x in _lignes(d[u'r1'])]
    w.tableau(COLONNES_DIP[u'r1'], rows or [dict(cellules=[AR] * 4)], _largeurs(4, premiere=6.2), numeriques=(1, 2, 3))
    # RUBRIQUE 5 : le nombre de mouvements de la vue ; sa ligne « à remplir » le rend « à remplir », et non 0
    w.titre(TITRES_DIP[u'r5'], 2)
    r5 = _lignes(d[u'r5'])
    w.tableau(None, [dict(cellules=[CARTE_MOUVEMENTS, ou_ar(nombre(_valeur(r5, u'nombre_mouvements')))])],
              [10.0, 7.0], numeriques=(1,), taille=9)
    w.para(TITRES_DIP[u'r5t'], gras=True, taille=9)
    rows = [dict(cellules=[ou_ar(x[u'code_actif']), ou_ar(x[u'nature']), ou_ar(date_fr(x[u'date_mouvement'])),
                           ou_ar(montant(x[u'montant']))]) for x in r5]
    w.tableau(COLONNES_DIP[u'r5'], rows or [dict(cellules=[AR] * 4)], _largeurs(4, premiere=5.0), numeriques=(3,))
    # RUBRIQUE 4
    w.titre(TITRES_DIP[u'r4'], 2)
    rows = [dict(cellules=[ou_ar(x[u'code_actif']), ou_ar(x[u'famille_libelle']), ou_ar(x[u'adresse']),
                           ou_ar(montant(x[u'valeur_comptable'])), ou_ar(montant(x[u'valeur_actuelle'])),
                           ou_ar(x[u'source_libelle']), ou_ar(x[u'article'])]) for x in _lignes(d[u'r4'])]
    w.tableau(COLONNES_DIP[u'r4'], rows or [dict(cellules=[AR] * 7)], [1.8, 2.5, 3.3, 2.4, 2.4, 2.5, 2.1],
              numeriques=(3, 4))
    # RUBRIQUE 6 : l'obligation de l'exercice que la vue lit, et sa lecture, reprise telle quelle (a un arrete non
    # annuel, le dernier exercice clos ; la base date une distribution a la cloture de son exercice, pas a son versement)
    w.titre(TITRES_DIP[u'r6'], 2)
    rows = [dict(cellules=[ou_ar(x[u'categorie']), ou_ar(montant(x[u'base_calcul'])), ou_ar(pourcentage(x[u'taux'])),
                           ou_ar(montant(x[u'obligation']))]) for x in _lignes(d[u'r6'])]
    w.tableau(COLONNES_DIP[u'r6'], rows or [dict(cellules=[AR] * 4)], _largeurs(4, premiere=7.4), numeriques=(1, 2, 3))
    lecture = _valeur(d[u'r6'], u'lecture')
    if lecture:
        w.para(lecture[:1].upper() + lecture[1:] + u'.', italique=True, taille=8)
    # RUBRIQUE 7 ET FONDEMENT
    w.para(TITRES_DIP[u'r7'])
    w.para(TITRES_DIP[u'fondement'], italique=True, taille=8)
    return w.octets()
# <<< fabrication_livrables.py

# ============================================================================ Fabric

import base64 as _base64

import fabric.functions as fn

udf = fn.UserDataFunctions()


def _qui(ctx) -> str:
    u = (ctx.executing_user or {}).get("PreferredUsername") or (ctx.executing_user or {}).get("Oid")
    if not u:
        raise fn.UserThrownError("L'utilisateur connecté n'est pas identifiable ; l'action n'est pas exécutée.", {})
    return u


def _texte_pilote(exc) -> str:
    """Le texte de l'erreur du pilote ODBC : le 2e argument de l'exception pyodbc, (etat SQL, texte), sinon str(exc)."""
    a = getattr(exc, "args", ())
    return a[1] if len(a) > 1 and isinstance(a[1], str) else str(exc)


def _message_sql(exc) -> str:
    s = _texte_pilote(exc)
    m = re.search(r"\[SQL Server\](.*?)(\s\(\d+\)\s\(SQL\w+\)|$)", s, re.S)
    return (m.group(1) if m else s).strip()


# UNE ERREUR NON PREVUE REMONTE EN FRANCAIS (24/10/2026), comme dans fn_ecran_arrete et fn_ecran_valorisation : THROW
# n'admet que des numeros d'au moins 50000, et seul un refus de la base (numero d'au moins 50000) devient le texte du
# bouton et de la zone de message ; toute autre erreur, du moteur SQL (numero inferieur), du pilote, du stockage ou de
# Python, rend TEXTE_ERREUR (_rendre), le detail technique allant au seul journal des refus.
PREMIER_NUMERO_DE_REFUS = 50000


def _numero_sql(exc):
    """Le numero d'erreur de SQL Server que porte l'exception, celui du 1er enregistrement, ou None quand elle n'en
    porte pas (delai ou coupure du pilote, erreur Python)."""
    m = re.search(r"\[SQL Server\].*?\s\((\d+)\)\s\(SQL\w+\)", _texte_pilote(exc), re.S)
    return int(m.group(1)) if m else None


def _refus_de_la_base(exc):
    """Le texte francais d'un refus de la base (numero d'au moins 50000), ou None pour toute autre erreur."""
    n = _numero_sql(exc)
    return _message_sql(exc) if n is not None and n >= PREMIER_NUMERO_DE_REFUS else None


def _detail(exc) -> str:
    """Le detail technique d'une erreur non prevue, pour le seul journal des refus : classe, numero, texte."""
    n = _numero_sql(exc)
    texte = " ".join((_message_sql(exc) or str(exc) or "").split())
    return "%s%s : %s" % (exc.__class__.__name__, "" if n is None else " n° %d" % n, texte)


def _executer(base: fn.FabricSqlConnection, sql: str, params: tuple) -> dict:
    """Execute une procedure et rend sa derniere ligne. Seul un refus de la base (numero d'au moins 50000) devient une
    UserThrownError a son texte ; toute autre erreur est relevee telle quelle apres le rollback, et _rendre la rend en
    francais (24/10/2026)."""
    conn = base.connect()
    try:
        cur = conn.cursor()
        cur.execute(sql, params)
        derniere = None
        while True:
            if cur.description:
                colonnes = [d[0] for d in cur.description]
                for ligne in cur.fetchall():
                    derniere = dict(zip(colonnes, ligne))
            if not cur.nextset():
                break
        conn.commit()
        return derniere or {}
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        message = _refus_de_la_base(e)
        if message is None:
            raise
        raise fn.UserThrownError(message, {"sql": sql.split("@")[0].strip()})
    finally:
        try:
            conn.close()
        except Exception:
            pass


class _Lecteur(object):
    """Une connexion pour toutes les lectures d'une fabrication ; lire(sql, params) rend une liste de dict."""

    def __init__(self, base):
        self.base = base
        self.conn = None

    def __enter__(self):
        self.conn = self.base.connect()

        def lire(sql, params=()):
            cur = self.conn.cursor()
            if params:
                cur.execute(sql, params)
            else:
                cur.execute(sql)
            if not cur.description:
                return []
            colonnes = [d[0] for d in cur.description]
            return [dict(zip(colonnes, l)) for l in cur.fetchall()]
        return lire

    def __exit__(self, *exc):
        self.conn.close()
        return False


def _modele(coffre: fn.FabricLakehouseClient) -> bytes:
    """Le modele Word du cabinet. Illisible, il est refuse en francais ; le texte de l'exception du client de fichiers
    va au seul journal des refus (detail_journal, lu par _rendre), et non plus dans les proprietes de l'UserThrownError,
    que le service rend a l'application qui appelle la fonction (Learn, « Python programming model for Fabric user data
    function », UserThrownError, relu le 24/10/2026)."""
    try:
        return coffre.connectToFiles().get_file_client(MODELE_WORD).download_file().readall()
    except Exception as e:
        texte = insecables(
            "Production refusée : le modèle Word du cabinet est introuvable. Il se dépose sur le site du cabinet, "
            "bibliothèque Documents, dossier Modeles, sous le nom modele_livrables_opci.docx.")
        refus = fn.UserThrownError(texte, {"chemin": "Files/%s" % MODELE_WORD})
        try:
            refus.detail_journal = "%s Détail technique : %s" % (texte, _detail(e))
        except Exception:
            pass
        raise refus from None


def _demander(base, entite: str, arrete: str, livrable: str, par: str) -> int:
    """La garde de la base : une demande REFUSEE rend son motif tel quel, sans rien fabriquer."""
    r = _executer(base, "EXEC dbo.pr_demander_document @entite=?, @arrete=?, @livrable=?, @par=?",
                  (entite, arrete, livrable, par))
    if r.get("etat") != "DEMANDEE":
        raise fn.UserThrownError(r.get("lecture") or "Production refusée.", {"version": r.get("version")})
    return int(r["version"])


# LA VERSION OUVERTE : la derniere demande du livrable, ses fichiers, ceux du format demande, ses fichiers perimes, la
# date de son 1er fichier et celle de la derniere decision de revue de l'arrete ; les versions de l'annexe que ses
# fichiers portent, et la derniere version de l'annexe de l'arrete.
_SQL_VERSION = (
    "SELECT TOP (1) dd.version, dd.etat, "
    "(SELECT COUNT(*) FROM dbo.document_produit d WHERE d.entite = dd.entite AND d.arrete = dd.arrete "
    "AND d.livrable = dd.livrable AND d.version = dd.version) AS fichiers, "
    "(SELECT COUNT(*) FROM dbo.document_produit d WHERE d.entite = dd.entite AND d.arrete = dd.arrete "
    "AND d.livrable = dd.livrable AND d.version = dd.version AND d.format = ?) AS du_format, "
    "(SELECT COUNT(*) FROM dbo.document_produit d WHERE d.entite = dd.entite AND d.arrete = dd.arrete "
    "AND d.livrable = dd.livrable AND d.version = dd.version AND d.perime_le IS NOT NULL) AS perimes, "
    "(SELECT MIN(d.produit_le) FROM dbo.document_produit d WHERE d.entite = dd.entite AND d.arrete = dd.arrete "
    "AND d.livrable = dd.livrable AND d.version = dd.version) AS premier_le, "
    "(SELECT MAX(v.decide_le) FROM dbo.visa v WHERE v.nature = 'REVUE' AND v.objet_ref = dd.entite + '|' + dd.arrete) "
    "AS revue_le, "
    "(SELECT COUNT(d.annexe_version) FROM dbo.document_produit d WHERE d.entite = dd.entite AND d.arrete = dd.arrete "
    "AND d.livrable = dd.livrable AND d.version = dd.version) AS annexe_portees, "
    "(SELECT MIN(d.annexe_version) FROM dbo.document_produit d WHERE d.entite = dd.entite AND d.arrete = dd.arrete "
    "AND d.livrable = dd.livrable AND d.version = dd.version) AS annexe_min, "
    "(SELECT MAX(d.annexe_version) FROM dbo.document_produit d WHERE d.entite = dd.entite AND d.arrete = dd.arrete "
    "AND d.livrable = dd.livrable AND d.version = dd.version) AS annexe_max, "
    "(SELECT a.derniere_version FROM dbo.v_annexe_version a WHERE a.entite = dd.entite AND a.arrete = dd.arrete) "
    "AS annexe_derniere "
    "FROM dbo.demande_document dd WHERE dd.entite = ? AND dd.arrete = ? AND dd.livrable = ? ORDER BY dd.version DESC")


def _reprise(r, livrable: str) -> bool:
    """REPRISE de la derniere demande r (ligne de _SQL_VERSION, None sans demande) si elle est DEMANDEE, porte deja un
    fichier d'un autre format, aucun du format demande, aucun fichier perime, et si son 1er fichier est posterieur a la
    derniere decision de revue ; pour les comptes annuels, si de plus chaque fichier deja enregistre porte la derniere
    version de l'annexe de l'arrete (les formats d'un meme document impriment la meme version)."""
    annexe_a_jour = livrable != "COMPTES_ANNUELS" or (
        r is not None and r.get("annexe_derniere") is not None and r.get("annexe_portees") == r.get("fichiers")
        and r.get("annexe_min") == r.get("annexe_max") == r.get("annexe_derniere"))
    return bool(r and r.get("etat") == "DEMANDEE" and r.get("fichiers") and not r.get("du_format")
                and not r.get("perimes") and r.get("premier_le") is not None and r.get("revue_le") is not None
                and r["premier_le"] > r["revue_le"] and annexe_a_jour)


def _version(base, entite: str, arrete: str, livrable: str, fmt: str, par: str):
    """(version, neuve) : la version ouverte qui attend encore ce format (_reprise), ou une version neuve demandee a la
    base, pr_demander_document, qui recalcule l'annexe. pr_enregistrer_document rejuge les conditions de production au
    moment d'enregistrer."""
    with _Lecteur(base) as lire:
        r = (lire(_SQL_VERSION, (fmt, entite, arrete, livrable)) or [None])[0]
    if _reprise(r, livrable):
        return int(r["version"]), False
    return _demander(base, entite, arrete, livrable, par), True


def _empreinte_stockee(coffre, entite: str, arrete: str, fichier: str):
    """L'empreinte de la COPIE STOCKEE, relue par le raccourci fec juste apres le depot (63_SQL/251, 04/10/2026).

    POURQUOI 2 EMPREINTES, releve du 04/10/2026 : SharePoint reecrit un fichier Office a l'enregistrement. Il ajoute
    des parties de metadonnees de bibliotheque (customXml, docProps/custom.xml) et en modifie 5 ; word/document.xml
    reste identique. L'empreinte des octets produits prouve la fabrication ; celle-ci date l'etat d'une copie a
    l'heure de la relecture, et une modification ulterieure des proprietes du fichier dans SharePoint la change.
    Une relecture qui echoue ne bloque pas l'enregistrement : la base admet son absence."""
    try:
        chemin = "fec/Dossiers de travail/%s/%s/Livrables/%s" % (entite, arrete, fichier)
        return empreinte(coffre.connectToFiles().get_file_client(chemin).download_file().readall())
    except Exception:
        return None


def _deposer(base, entite: str, arrete: str, fichier: str, octets: bytes, par: str) -> dict:
    """Depot au site du cabinet, Dossiers de travail/<entite>/<arrete>/Livrables, avec l'empreinte des octets. La base
    la recalcule sur le contenu recu et refuse l'ecrasement d'un livrable."""
    return _executer(base, "EXEC dbo.pr_deposer_classeur @entite=?, @arrete=?, @fichier=?, @contenu_base64=?, @par=?, "
                           "@sous_dossier=?, @empreinte=?",
                     (entite, arrete, fichier, _base64.b64encode(octets).decode("ascii"), par, "Livrables",
                      empreinte(octets)))


def _enregistrer(base, coffre, entite: str, arrete: str, livrable: str, fichier: str, par: str,
                 annexe_version: int = None) -> dict:
    """Enregistrement du fichier depose, avec l'empreinte de sa copie stockee (63_SQL/251) ; la base reprend du depot
    son adresse, sa taille et l'empreinte des octets produits. Pour les comptes annuels, la version de l'annexe que le
    fichier imprime, passee par son nom (@annexe_version) : la base refuse un document sans elle, sur une version
    depassee, d'une autre version que le format deja enregistre du meme document, ou quand les cellules de l'entite
    sont en retard sur une ecriture."""
    sql = ("EXEC dbo.pr_enregistrer_document @entite=?, @arrete=?, @livrable=?, @fichier=?, @par=?, "
           "@empreinte_stockee=?")
    params = (entite, arrete, livrable, fichier, par, _empreinte_stockee(coffre, entite, arrete, fichier))
    if annexe_version is not None:
        sql += ", @annexe_version=?"
        params += (int(annexe_version),)
    return _executer(base, sql, params)


def _refus(e, version=None):
    """Le refus du module, suivi de la version qu'il laisse ouverte : pr_demander_document l'a creee, aucun fichier ne
    s'y rattache, et la production suivante en demande une autre (la version suivante prend le maximum + 1, 63_SQL/262)."""
    suite = (" La version %d reste ouverte, sans fichier ; la production suivante ouvre la version %d." % (
        version, version + 1) if version else "")
    raise fn.UserThrownError(insecables(str(e) + suite), {})


# ----------------------------------------------------------------------------- la zone de message de l'ecran

def _poser_message(base, par: str, texte: str, genre: str, entite: str, procedure: str, journal: bool = False,
                   texte_journal: str = None) -> None:
    """Le message du geste dans la boite aux lettres de l'ecran Revision (dbo.pr_poser_message_ecran), que le modele lit
    par [Message de la revision] : la phrase rendue a la reussite, le texte du refus sinon, comme fn_ecran_revision.
    Le refus se journalise d'abord (dbo.pr_journaliser_refus). UNE CONNEXION NEUVE, donc une transaction neuve : le
    geste a deja fait commit ou rollback. Ne leve jamais : un message qui echoue ne masque pas le geste qu'il rapporte.
    texte_journal : le texte journalise quand il differe de celui de l'ecran (l'erreur d'execution entiere)."""
    texte = (texte or "").strip()[:2000]
    texte_journal = (texte_journal or "").strip()[:2000] or texte
    if not texte or not par:
        return
    ent = (entite or "").strip()[:20] or None
    try:
        conn = base.connect()
    except Exception:
        return
    try:
        if journal:
            try:
                cur = conn.cursor()
                cur.execute("EXEC dbo.pr_journaliser_refus @procedure_nom=?, @message=?, @par=?, @entite=?",
                            (procedure[:128], texte_journal, par, ent))
                cur.close()
                conn.commit()
            except Exception:
                try:
                    conn.rollback()
                except Exception:
                    pass
        cur = conn.cursor()
        cur.execute("EXEC dbo.pr_poser_message_ecran @pour=?, @message=?, @genre=?, @entite=?, @procedure_nom=?",
                    (par, texte, genre, ent, procedure[:128]))
        cur.close()
        conn.commit()
    except Exception:
        try:
            conn.rollback()
        except Exception:
            pass
    finally:
        try:
            conn.close()
        except Exception:
            pass


def _texte_erreur(e) -> str:
    v = getattr(e, "message", None)
    if isinstance(v, str) and v.strip():
        return v
    return str(e) or e.__class__.__name__


# L'ERREUR D'EXECUTION A L'ECRAN, depuis le 24/10/2026 : le texte francais fixe de fn_ecran_arrete et de
# fn_ecran_valorisation, et non plus la classe et le texte anglais de l'exception, bornes a 160 caracteres (05/10/2026).
# Le detail technique (production, classe, numero, texte entier) va au seul journal des refus, et la fonction leve une
# UserThrownError (reponse 422) au lieu de relancer l'exception (reponse 500).
TEXTE_ERREUR = ("Action interrompue par une erreur d'exécution. Le détail technique est consigné au journal des refus ; "
                "relisez l'écran avant de recommencer.")


def _rendre(base, par: str, entite: str, nom: str, faire):
    """Execute la production, puis pose son message : la phrase rendue a la reussite, le texte du refus sinon. Le refus
    est relance tel quel : la bulle du service le montre aussi ; son detail technique, quand il en porte un
    (detail_journal, modele Word illisible), va au journal a la place du texte. Une erreur non prevue rend TEXTE_ERREUR,
    son detail au journal seul (24/10/2026)."""
    try:
        texte = faire()
    except fn.UserThrownError as e:
        _poser_message(base, par, _texte_erreur(e), "REFUS", entite, nom, journal=True,
                       texte_journal=getattr(e, "detail_journal", None))
        raise
    except Exception as e:
        detail = "Erreur d’exécution de %s : %s" % (nom, _detail(e))
        _poser_message(base, par, TEXTE_ERREUR, "REFUS", entite, nom, journal=True, texte_journal=detail)
        raise fn.UserThrownError(TEXTE_ERREUR, {"fonction": nom}) from e
    _poser_message(base, par, texte, "SUCCES", entite, nom)
    return texte


# ----------------------------------------------------------------------------- les 3 productions

# LA PHRASE DE REUSSITE, recomposee le 05/10/2026 : la 1re version rendait 285 a 317 caracteres, le nom du fichier 2
# fois, et la phrase de la base (pr_enregistrer_document) a la suite ; la 2e ligne de la zone finissait par « … ». La
# fonction compose seule, sur les colonnes que la base rend (demande_produite), et non sur sa phrase : le livrable, sa
# version et son format, le fichier depose et son dossier, ce qu'il porte, puis ce qui reste a produire de la version.
# Le nom du fichier vient une fois, en tete, ou il tient sur la 1re ligne. Chaque phrase tient en 2 lignes de 870 px
# (poser_ecran_revision.lignes_msg_py), pour les noms reels et l'entite au code le plus large de ref_entite :
# 61_DEPLOIEMENT/prouver_messages_fn_documents.py.
_PORTE = {("COMPTES_ANNUELS", "DOCX"): "Il porte le bilan, le compte de résultat et l’annexe.",
          ("COMPTES_ANNUELS", "XLSX"): "Il porte le bilan, le compte de résultat et la balance."}
_LE_FORMAT = {"DOCX": "le Word", "XLSX": "l’Excel"}


def _message_produit(livrable: str, version: int, fmt: str, fichier: str, r: dict, precision: str = "") -> str:
    formats = _FORMATS[livrable]
    if r.get("demande_produite"):
        suite = "La version %d est complète%s." % (
            version, (" : " + " et ".join(_NOM_FORMAT[x] for x in formats)) if len(formats) > 1 else "")
    else:
        reste = [x for x in formats if x != fmt]
        suite = ("Reste à produire : %s de la version %d." % (" et ".join(_LE_FORMAT[x] for x in reste), version)
                 if reste else "La version %d attend encore un fichier." % version)
    tete = "%s, version %d, %s : %s déposé au dossier Livrables de l’arrêté, sur le site du cabinet." % (
        NOMS_LIVRABLES[livrable], version, _NOM_FORMAT[fmt], fichier)
    return insecables(" ".join(x for x in (tete, precision or _PORTE.get((livrable, fmt), ""), suite) if x))


def _faire_comptes_annuels(base, coffre, par: str, entite: str, arrete: str, fmt: str) -> str:
    """Le fichier d'un format des comptes annuels : le Word (bilan, compte de resultat et annexe) ou l'Excel (bilan,
    compte de resultat et balance, 03/10/2026). Il se rattache a la version ouverte qui attend ce format, ou a une
    version neuve ; la version passe a PRODUITE quand les 2 formats y sont enregistres.

    LA VERSION DE L'ANNEXE est fixee a la 1re lecture qui suit la demande (la demande recalcule l'annexe) : les
    cellules de l'annexe se lisent toutes a cette version, le fichier l'imprime et l'enregistre (@annexe_version).
    Avant le depot, la fonction rejuge ce que la base rejugera a l'enregistrement, avec ses textes : l'annexe n'a pas
    change de version depuis la lecture, l'autre format du meme document porte la meme version, et ses cellules ne
    sont pas en retard sur une ecriture ; sinon aucun fichier n'est depose."""
    version, neuve, annexe = None, False, None
    try:
        with _Lecteur(base) as lire:
            verifier_interface_depot(lire)
        modele = _modele(coffre) if fmt == "DOCX" else None
        version, neuve = _version(base, entite, arrete, "COMPTES_ANNUELS", fmt, par)
        with _Lecteur(base) as lire:
            annexe = version_annexe(lire, entite, arrete)
            d = lire_comptes_annuels(lire, entite, arrete, annexe_version=annexe)
        jour = datetime.date.today()
        if fmt == "DOCX":
            octets = fabriquer_comptes_annuels_docx(d, version, modele, jour)
        else:
            octets = fabriquer_comptes_annuels_xlsx(d, version, jour)
        with _Lecteur(base) as lire:
            garde_annexe(lire, entite, arrete, annexe, version)
    except Refus as e:
        _refus(e, version if neuve else None)
    f = nom_fichier("COMPTES_ANNUELS", entite, arrete, version, fmt.lower())
    _deposer(base, entite, arrete, f, octets, par)
    r = _enregistrer(base, coffre, entite, arrete, "COMPTES_ANNUELS", f, par, annexe_version=annexe)
    return _message_produit("COMPTES_ANNUELS", version, fmt, f, r)


def _faire_rapport(base, coffre, par: str, entite: str, arrete: str) -> str:
    """Le texte de la forme arretee, lu en base (exemples E1 a E4 de la NP 2300), ses emplacements remplis par
    v_rapport_expert_comptable, « a remplir » ailleurs ; jamais sur une forme seulement proposee. Le fichier se depose
    au dossier Livrables, puis pr_produire_attestation le rattache aux comptes annuels couverts et l'enregistre, avec
    l'empreinte de sa copie stockee.

    AVANT LA DEMANDE, 2 lectures de la base, pour ne deposer aucun fichier que pr_produire_attestation refuserait : la
    forme arretee, et la signature, par la fonction de la base que la procedure applique (fn_peut_viser_nature,
    PUBLICATION). Ses autres refus (revue non visee, arrete valide pour le client, comptes annuels non produits)
    sont ceux de la demande, qui les oppose avant toute fabrication. AVANT LE DEPOT, la marque de retard de l'entite,
    que l'enregistrement du rapport refuse avec le meme texte (garde_retard)."""
    version = None
    try:
        with _Lecteur(base) as lire:
            verifier_interface_depot(lire)
            pre = lire_rapport(lire, entite, arrete)
            signe = (lire(SQL["signataire"], (entite, par)) or [{}])[0].get("peut")
        if not pre["forme_arretee"]:
            raise Refus("Production refusée : la forme du rapport de l’expert-comptable n’est pas arrêtée pour cet "
                        "arrêté ; le document ne sort jamais sur une forme proposée.")
        if not signe:
            raise Refus("Production refusée : le rapport de l’expert-comptable est signé par l’associé signataire, "
                        "et lui seul le produit.")
        texte_forme(pre, pre["forme_arretee"])
        modele = _modele(coffre)
        version = _demander(base, entite, arrete, "ATTESTATION", par)
        with _Lecteur(base) as lire:
            d = lire_rapport(lire, entite, arrete)
        if not d["forme_arretee"]:
            raise Refus("Production refusée : la forme du rapport de l’expert-comptable n’est plus arrêtée pour cet "
                        "arrêté.")
        word = fabriquer_rapport_docx(d, d["forme_arretee"], version, modele, datetime.date.today())
        with _Lecteur(base) as lire:
            garde_retard(lire, entite)
    except Refus as e:
        _refus(e, version)
    f = nom_fichier("ATTESTATION", entite, arrete, version, "docx")
    _deposer(base, entite, arrete, f, word, par)
    r = _executer(base, "EXEC dbo.pr_produire_attestation @entite=?, @arrete=?, @fichier=?, @par=?, "
                        "@empreinte_stockee=?",
                  (entite, arrete, f, par, _empreinte_stockee(coffre, entite, arrete, f)))
    exemple = (d["formes"].get(d["forme_arretee"]) or {}).get("exemple")
    return _message_produit("ATTESTATION", version, "DOCX", f, r,
                            "Forme : exemple %s de la NP 2300." % exemple if exemple else "")


def _faire_dip(base, coffre, par: str, entite: str, arrete: str) -> str:
    """Les rubriques de la page « Le DIP » du 06/09/2026, lues sur v_dip_rubrique_1 a 6 ; a un arrete non annuel, la
    rubrique 6 lit le dernier exercice clos, et le document reprend la lecture que la base en donne. Word seul."""
    version = None
    try:
        with _Lecteur(base) as lire:
            verifier_interface_depot(lire)
        modele = _modele(coffre)
        version = _demander(base, entite, arrete, "DIP", par)
        with _Lecteur(base) as lire:
            d = lire_dip(lire, entite, arrete)
        word = fabriquer_dip_docx(d, version, modele, datetime.date.today())
    except Refus as e:
        _refus(e, version)
    f = nom_fichier("DIP", entite, arrete, version, "docx")
    _deposer(base, entite, arrete, f, word, par)
    r = _enregistrer(base, coffre, entite, arrete, "DIP", f, par)
    return _message_produit("DIP", version, "DOCX", f, r)


# LE LIVRABLE ET LE FORMAT DESIGNENT LA PRODUCTION : le code de ref_livrable et le format que l'ecran passent
_PRODUCTIONS = {"COMPTES_ANNUELS": ("produire_comptes_annuels", _faire_comptes_annuels),
                "ATTESTATION": ("produire_rapport", _faire_rapport),
                "DIP": ("produire_dip", _faire_dip)}
_FORMATS = {"COMPTES_ANNUELS": ("DOCX", "XLSX"), "ATTESTATION": ("DOCX",), "DIP": ("DOCX",)}
_NOM_FORMAT = {"DOCX": "Word", "XLSX": "Excel"}

# LES CONTROLES D'ENTREE DE LA PRODUCTION (24/10/2026, decision du candidat « 1-A », regle du meme jour : tout message
# a l'ecran en francais). entite, arrete et livrable ont une valeur par defaut vide a produire_livrable ; format,
# constante du bouton, reste obligatoire, en tete. La production lit la base avant son 1er appel de procedure
# (interface de depot, version ouverte, forme arretee, signataire) : le vehicule, le livrable puis l'arrete vides sont
# refuses ici, avant le controle du code du livrable, avec les textes de l'etude du 24/10/2026 (_scratch/fonctions_fr/
# etude_fonctions_20261024.json, champ « plan »), mot pour mot, que la base reprend dans ses procedures.
REFUS_VEHICULE = "Choisissez le véhicule dans le sélecteur Véhicule."
REFUS_LIVRABLE = "Production refusée : choisissez le livrable dans la liste « 1. Livrables de l’arrêté »."
REFUS_ARRETE_RAPPORT = ("Production refusée : le rapport de l’expert-comptable ne se produit pas à cet arrêté en "
                        "l'état ; il exige notamment la forme arrêtée, l'associé signataire, la revue visée et les "
                        "comptes annuels produits.")
# %s : « Word » ou « Excel », selon le format du bouton
REFUS_ARRETE_LIVRABLE = ("Production refusée : le livrable choisi ne se produit pas en %s à cet arrêté en l'état ; "
                         "elle exige notamment la revue visée, un arrêté non validé pour le client et un livrable dû, "
                         "non encore produit dans ce format.")


def _produire(base, coffre, ctx, entite: str, arrete: str, livrable: str, fmt: str = "DOCX") -> str:
    par = _qui(ctx)
    code = (livrable or "").strip().upper()
    f = (fmt or "").strip().upper()
    texte = None
    if (entite or "").strip() == "":
        texte = REFUS_VEHICULE
    elif code == "":
        texte = REFUS_LIVRABLE
    elif (arrete or "").strip() == "":
        texte = REFUS_ARRETE_RAPPORT if code == "ATTESTATION" else REFUS_ARRETE_LIVRABLE % _NOM_FORMAT.get(f, "Word")
    elif code not in _PRODUCTIONS:
        texte = ("Production refusée : %s ; seuls les comptes annuels, le rapport de l’expert-comptable et le DIP se "
                 "produisent à l’étape des livrables." % (
                     "ce livrable n’est pas produit par la solution" if code else "aucun livrable n’est choisi"))
    elif f not in _FORMATS[code]:
        texte = ("Production refusée : le livrable %s ne se produit pas %s ; il se produit en %s." % (
            guillemets(NOMS_LIVRABLES[code]), ("en " + _NOM_FORMAT.get(f, f)) if f else "sans format",
            " et en ".join(_NOM_FORMAT[x] for x in _FORMATS[code])))
    if texte:
        texte = insecables(texte)
        _poser_message(base, par, texte, "REFUS", entite, "produire_livrable", journal=True)
        raise fn.UserThrownError(texte, {})
    nom, faire = _PRODUCTIONS[code]
    if code == "COMPTES_ANNUELS":
        return _rendre(base, par, entite, nom, lambda: faire(base, coffre, par, entite, arrete, f))
    return _rendre(base, par, entite, nom, lambda: faire(base, coffre, par, entite, arrete))


@udf.context(argName="ctx")
@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def produire_livrable(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient,
                      ctx: fn.UserDataFunctionContext, format: str, entite: str = "", arrete: str = "",
                      livrable: str = "") -> str:
    """Boutons « Produire en Word » et « Produire en Excel » de l'etape 6 et des pages des documents : le livrable
    (COMPTES_ANNUELS, ATTESTATION ou DIP) et le format (DOCX ou XLSX) passent en parametre, et le bouton ne produit
    que le fichier de son format. Un bouton n'appelle qu'une fonction, et le livrable choisi a l'ecran change d'un
    clic a l'autre : la fonction aiguille vers la production du livrable. format, constante du bouton, reste
    obligatoire et passe en tete ; les 3 autres ont une valeur par defaut vide (24/10/2026, controles de _produire)."""
    return _produire(base, coffre, ctx, entite, arrete, livrable, format)


@udf.context(argName="ctx")
@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def produire_comptes_annuels(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient,
                             ctx: fn.UserDataFunctionContext, entite: str, arrete: str) -> str:
    """Les comptes annuels en Word puis en Excel : le classeur se rattache a la version que le Word a ouverte. Chaque
    production pose sa phrase ; la fonction rend celle de l'Excel, qui dit la version complete (05/10/2026 : les 2
    phrases mises bout a bout ne tenaient pas en 2 lignes)."""
    _produire(base, coffre, ctx, entite, arrete, "COMPTES_ANNUELS", "DOCX")
    return _produire(base, coffre, ctx, entite, arrete, "COMPTES_ANNUELS", "XLSX")


@udf.context(argName="ctx")
@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def produire_rapport(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient,
                     ctx: fn.UserDataFunctionContext, entite: str, arrete: str) -> str:
    """Bouton « Produire le rapport » : le rapport de l'expert-comptable, sur la forme arretee."""
    return _produire(base, coffre, ctx, entite, arrete, "ATTESTATION")


@udf.context(argName="ctx")
@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def produire_dip(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient,
                 ctx: fn.UserDataFunctionContext, entite: str, arrete: str) -> str:
    """Bouton « Produire le DIP » : le document d'information periodique, Word seul."""
    return _produire(base, coffre, ctx, entite, arrete, "DIP")
