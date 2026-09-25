"""fn_ecran_revision : les boutons de l'ecran Revision, une fonction par bouton.

Ecrit le 16/09/2026 d'apres la maquette 73 validee par le candidat et les scripts 207 et 208. Memes
regles que fn_ecran_client : chaque fonction appelle UNE procedure, rend la phrase du bouton, prend
l'utilisateur dans le contexte d'execution (PreferredUsername), et un refus remonte en francais par
UserThrownError. Connexions : « DossierOPCI » (base) et « Coffre » (lakehouse). Bibliotheque : openpyxl.

Les 6 etapes de la maquette et leurs boutons :
  1 Collecte      reclamer_document (fn_ecran_client la porte deja ; rappelee ici pour la page)
  2 Programme     choisir_programme, ajuster_programme, verifier_couverture, exporter_questions
  3 Questions     repondre_question, ouvrir_feuille, exporter_feuille, importer_classeur (feuille, OD,
                  questions), conclure_feuille, exporter_od, inscrire_piece (fn_ecran_client), valider_od_question
  4 OD libres     saisir_od, exporter_od (sans question), importer_classeur, valider_od_libres
  5 Conclure      conclure_cycle, viser (nature CYCLE ou REVUE, parametre superviseur obligatoire),
                  conclure_revue
  6 Dossier       exporter_dossier, deverrouiller_dossier, importer_classeur (dossier entier)

Le verrou (207) est controle par les procedures : une fonction n'a rien a en savoir, elle rend le refus.
"""

import datetime
import io
import json
import re

import fabric.functions as fn

udf = fn.UserDataFunctions()


# ----------------------------------------------------------------------------- outils, identiques a fn_ecran_client

def _qui(ctx) -> str:
    u = (ctx.executing_user or {}).get("PreferredUsername") or (ctx.executing_user or {}).get("Oid")
    if not u:
        raise fn.UserThrownError("L'utilisateur connecté n'est pas identifiable ; l'action n'est pas exécutée.", {})
    return u


def _message_sql(exc) -> str:
    s = str(exc)
    m = re.search(r"\[SQL Server\](.*?)(\s\(\d+\)\s\(SQLExecDirectW\)|$)", s, re.S)
    return (m.group(1) if m else s).strip()


def _executer(base: fn.FabricSqlConnection, sql: str, params: tuple) -> dict:
    conn = base.connect()
    try:
        cur = conn.cursor()
        cur.execute(sql, params)
        derniere = None
        while True:
            if cur.description:
                colonnes = [d[0] for d in cur.description]
                for ligne in cur.fetchall():
                    derniere = dict(zip(colonnes, ligne))
            if not cur.nextset():
                break
        conn.commit()
        return derniere or {}
    except Exception as e:
        conn.rollback()
        raise fn.UserThrownError(_message_sql(e), {"sql": sql.split("(")[0].strip()})
    finally:
        conn.close()


def _lire(base: fn.FabricSqlConnection, sql: str, params: tuple) -> list:
    conn = base.connect()
    try:
        cur = conn.cursor()
        cur.execute(sql, params)
        colonnes = [d[0] for d in cur.description]
        return [dict(zip(colonnes, l)) for l in cur.fetchall()]
    finally:
        conn.close()


def _vide(s):
    return None if s is None or str(s).strip() == "" else str(s).strip()


def _superviseur(ctx, superviseur: str) -> None:
    if not superviseur or superviseur.strip().lower() != _qui(ctx).strip().lower():
        raise fn.UserThrownError("Le visa est réservé au superviseur connecté ; l'action n'est pas exécutée.", {})


# ----------------------------------------------------------------------------- etape 2 : le programme

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def choisir_programme(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                      entite: str, arrete: str, typeProgramme: str) -> str:
    """Bouton « Choisir » d'un type : ALLEGE, CLASSIQUE ou ETENDU. Rejouable, la selection est vivante."""
    r = _executer(base, "EXEC dbo.pr_choisir_programme @entite=?, @arrete=?, @type=?, @par=?",
                  (entite, arrete, typeProgramme.upper(), _qui(ctx)))
    return r.get("message", "Programme choisi.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def ajuster_programme(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                      entite: str, arrete: str, reference: str, actif: bool) -> str:
    """Boutons « Ajouter une question » et « Retirer » : la selection change, le journal le dit."""
    r = _executer(base, "EXEC dbo.pr_ajuster_programme @entite=?, @arrete=?, @reference=?, @actif=?, @par=?",
                  (entite, arrete, reference, 1 if actif else 0, _qui(ctx)))
    return r.get("message", "Programme ajusté.")


@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def verifier_couverture(base: fn.FabricSqlConnection, entite: str, arrete: str) -> str:
    """Bouton « Vérifier la couverture de la balance » : rend la phrase, ne modifie rien."""
    r = _executer(base, "EXEC dbo.pr_verifier_couverture @entite=?, @arrete=?", (entite, arrete))
    return r.get("message", "Couverture vérifiée.")


# ----------------------------------------------------------------------------- etape 3 : les questions

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def repondre_question(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                      cote: str, reference: str, reponse: str, motif: str = "", commentaire: str = "") -> str:
    """La ligne de la question : OUI, NON, NA avec motif ; la cote est celle de la feuille Q-<cycle>-P<n>."""
    _executer(base, "EXEC dbo.pr_repondre_question @cote=?, @reference=?, @reponse=?, @motif_non_applicable=?, @commentaire=?, @par=?",
              (cote, reference, _vide(reponse), _vide(motif), _vide(commentaire), _qui(ctx)))
    return f"Réponse à {reference} enregistrée."


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def ouvrir_feuille(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                   entite: str, arrete: str, reference: str) -> str:
    """Pop-up Feuille, « Ouvrir » : la feuille de la question, son gabarit, sa cote."""
    r = _executer(base, "EXEC dbo.pr_ouvrir_feuille_question @entite=?, @arrete=?, @reference=?, @par=?",
                  (entite, arrete, reference, _qui(ctx)))
    return r.get("message", "Feuille ouverte.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def conclure_feuille(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                     cote: str, forme: str, conclusion: str = "") -> str:
    """Pop-up Feuille, « Conclure » : forme et texte ; les portes de la feuille (pieces, derogation) parlent."""
    r = _executer(base, "EXEC dbo.pr_importer_feuille @cote=?, @conclusion=?, @forme=?, @par=?",
                  (cote, _vide(conclusion), forme.upper(), _qui(ctx)))
    return r.get("message", "Feuille conclue.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def valider_od_question(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                        entite: str, arrete: str, reference: str) -> str:
    """Pop-up OD, « Valider » : le brouillon de la question devient un lot propose au visa."""
    l = _lire(base, "SELECT fq.cote, q.id FROM dbo.ref_question q JOIN dbo.feuille_question fq ON fq.question_id = q.id "
                    "JOIN dbo.feuille_travail f ON f.cote = fq.cote AND f.entite=? AND f.arrete=? AND f.cote LIKE 'Q-%' WHERE q.reference=?",
              (entite, arrete, reference))
    if not l:
        raise fn.UserThrownError(f"La question {reference} n'est pas au programme de cet arrêté.", {})
    _executer(base, "EXEC dbo.pr_valider_brouillon @entite=?, @arrete=?, @feuille_cote=?, @question_id=?, @valide_par=?",
              (entite, arrete, l[0]["cote"], l[0]["id"], _qui(ctx)))
    return f"OD de la question {reference} validées : le lot est proposé au visa."


# ----------------------------------------------------------------------------- etape 4 : les OD libres

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def saisir_od(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
              entite: str, arrete: str, compte: str, libelle: str, debit: float = 0, credit: float = 0,
              reference: str = "", journal: str = "ODR", piece: str = "") -> str:
    """Une ligne d'OD, liee a une question si reference est donnee, libre sinon."""
    r = _executer(base, "EXEC dbo.pr_saisir_od @entite=?, @arrete=?, @compte=?, @libelle=?, @debit=?, @credit=?, @reference=?, @journal=?, @piece_ref=?, @par=?",
                  (entite, arrete, compte, libelle, float(debit or 0), float(credit or 0), _vide(reference), journal or "ODR", _vide(piece), _qui(ctx)))
    return r.get("message", "Ligne enregistrée.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def valider_od_libres(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext, entite: str, arrete: str) -> str:
    """Bouton « Valider le brouillon » des OD libres : equilibre exige, lot propose au visa."""
    r = _executer(base, "EXEC dbo.pr_valider_od_libres @entite=?, @arrete=?, @valide_par=?", (entite, arrete, _qui(ctx)))
    return r.get("message", "OD libres validées.")


# ----------------------------------------------------------------------------- etape 5 : conclure et viser

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def conclure_cycle(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                   entite: str, arrete: str, cycle: str, conclusion: str, forme: str = "") -> str:
    """Bouton « Enregistrer la conclusion du cycle » : la synthese des feuilles se regenere, le texte se pose."""
    r = _executer(base, "EXEC dbo.pr_conclure_cycle @entite=?, @arrete=?, @cycle=?, @conclusion=?, @forme=?, @par=?",
                  (entite, arrete, cycle, conclusion, _vide(forme), _qui(ctx)))
    return r.get("message", "Cycle conclu.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def conclure_revue(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                   entite: str, arrete: str, conclusion: str) -> str:
    """Bouton « Enregistrer la conclusion générale »."""
    r = _executer(base, "EXEC dbo.pr_conclure_revue @entite=?, @arrete=?, @conclusion=?, @par=?", (entite, arrete, conclusion, _qui(ctx)))
    return r.get("message", "Revue conclue.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def viser(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
          nature: str, objetRef: str, decision: str, superviseur: str, motif: str = "") -> str:
    """Boutons « Viser le cycle », « Viser la revue », « Renvoyer » : pr_ecran_viser, un seul bouton pour
    toutes les natures. objetRef : « entite|arrete|cycle » pour CYCLE, « entite|arrete » pour REVUE.
    Le parametre superviseur est lie a la mesure Superviseur connecte : vide pour un preparateur."""
    _superviseur(ctx, superviseur)
    r = _executer(base, "EXEC dbo.pr_ecran_viser @nature=?, @objet_ref=?, @decision=?, @decide_par=?, @motif=?",
                  (nature.upper(), objetRef, decision.upper(), _qui(ctx), _vide(motif)))
    return r.get("message", f"{nature} : {decision.lower()}.")


# ----------------------------------------------------------------------------- etape 6 : le dossier

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def deverrouiller_dossier(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext, entite: str, arrete: str) -> str:
    """Bouton « Déverrouiller le dossier de travail » : reserve a la personne qui a vise la revue."""
    r = _executer(base, "EXEC dbo.pr_deverrouiller_dossier @entite=?, @arrete=?, @par=?", (entite, arrete, _qui(ctx)))
    return r.get("message", "Dossier déverrouillé.")


# ----------------------------------------------------------------------------- la voie du classeur

# Les exports ecrivent au coffre, Files/exports/, et rendent le chemin ; les imports lisent Files/depots/.
# Meme lecture que 61_DEPLOIEMENT/gabarits_ecran_revision.py et importer_ecran_revision.py.

_COL_Q = [("Cycle", "cycle"), ("Référence", "reference"), ("Question", "enonce"), ("Type", "type"), ("Réponse", "reponse"),
          ("Motif si non applicable", "motif"), ("Commentaire", "commentaire"), ("Feuille", "feuille_cote"), ("OD", "od"), ("Pièces", "pieces")]
_COL_OD = [("Journal", "journal"), ("Compte", "compte"), ("Libellé", "libelle"), ("Débit", "debit"), ("Crédit", "credit"), ("Pièce", "piece"), ("Actif", "actif")]
_SAISIE = {"reponse", "motif", "commentaire", "journal", "compte", "libelle", "debit", "credit", "piece", "actif", "forme", "conclusion", "objectif", "methodologie"}


def _entete(ws, colonnes):
    from openpyxl.styles import Font, PatternFill
    for j, (lib, cle) in enumerate(colonnes, 1):
        c = ws.cell(row=1, column=j, value=lib)
        c.fill = PatternFill("solid", fgColor="1D1A10")
        c.font = Font(bold=True, color="FFFFFF") if cle in _SAISIE else Font(color="D9D2C0")
        ws.cell(row=2, column=j, value=cle).font = Font(italic=True, color="8A8578", size=9)
    ws.freeze_panes = "A3"


def _lignes(ws, colonnes, donnees, vides=0):
    from openpyxl.styles import PatternFill
    creme = PatternFill("solid", fgColor="FFF8E7")
    r = 3
    for d in donnees:
        for j, (lib, cle) in enumerate(colonnes, 1):
            c = ws.cell(row=r, column=j, value=d.get(cle))
            if cle in _SAISIE:
                c.fill = creme
        r += 1
    for _ in range(vides):
        for j, (lib, cle) in enumerate(colonnes, 1):
            if cle in _SAISIE:
                ws.cell(row=r, column=j).fill = creme
        r += 1


def _feuille_questions(wb, titre, donnees):
    ws = wb.create_sheet(titre[:31]); _entete(ws, _COL_Q); _lignes(ws, _COL_Q, donnees)


def _feuille_od(wb, titre, donnees, vides=12):
    ws = wb.create_sheet(titre[:31]); _entete(ws, _COL_OD); _lignes(ws, _COL_OD, donnees, vides)


def _feuille_feuille(wb, titre, d):
    from openpyxl.styles import Font, PatternFill
    creme = PatternFill("solid", fgColor="FFF8E7")
    ws = wb.create_sheet(titre[:31])
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 100
    ws.cell(row=1, column=1, value="Feuille de travail").font = Font(bold=True, size=13)
    ws.cell(row=1, column=2, value=d.get("cote")).font = Font(bold=True, size=13)
    r = 3
    for lib, cle in [("Cote", "cote"), ("Gabarit", "modele"), ("Entité", "entite"), ("Arrêté", "arrete"), ("Cycle", "cycle"),
                     ("Question", "reference"), ("Préparateur", "preparateur"), ("Objectif", "objectif"), ("Méthodologie", "methodologie")]:
        ws.cell(row=r, column=1, value=lib).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=d.get(cle))
        if cle in _SAISIE:
            c.fill = creme
        ws.cell(row=r, column=3, value=cle).font = Font(italic=True, color="8A8578", size=9)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Travaux").font = Font(bold=True)
    r += 1
    for j, lib in enumerate(["Contrôle", "Attendu", "Constaté", "Écart", "Commentaire"], 1):
        c = ws.cell(row=r, column=j, value=lib); c.fill = PatternFill("solid", fgColor="1D1A10"); c.font = Font(bold=True, color="FFFFFF")
    for i in range(r + 1, r + 13):
        for j in range(1, 6):
            ws.cell(row=i, column=j).fill = creme
    r += 14
    for lib, cle in [("Forme de conclusion", "forme"), ("Conclusion", "conclusion")]:
        ws.cell(row=r, column=1, value=lib).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=d.get(cle)); c.fill = creme
        ws.cell(row=r, column=3, value=cle).font = Font(italic=True, color="8A8578", size=9)
        r += 1


def _questions(base, entite, arrete, cycle=None):
    return _lire(base, "SELECT cycle, reference, enonce, type_reponse AS type, reponse_lue AS reponse, motif_non_applicable AS motif, commentaire, "
                       "feuille_cote, od_brouillon + od_lots AS od, pieces FROM dbo.v_questions_programme WHERE entite=? AND arrete=? "
                       + ("AND cycle=? " if cycle else "") + "ORDER BY cycle, ordre",
                 (entite, arrete, cycle) if cycle else (entite, arrete))


def _od(base, entite, arrete, reference=None, libres=False, cycle=None):
    sql = ("SELECT journal_code AS journal, compte_num AS compte, libelle, debit, credit, piece_ref AS piece, code_actif AS actif "
           "FROM dbo.v_od_revision WHERE entite=? AND arrete=? AND etat='BROUILLON' ")
    if reference:
        return _lire(base, sql + "AND reference=? ORDER BY le", (entite, arrete, reference))
    if libres:
        return _lire(base, sql + "AND reference IS NULL ORDER BY le", (entite, arrete))
    if cycle:
        return _lire(base, sql + "AND cycle=? ORDER BY le", (entite, arrete, cycle))
    return _lire(base, sql + "ORDER BY le", (entite, arrete))


def _feuilles(base, entite, arrete, cycle=None):
    return _lire(base, "SELECT f.cote, f.modele_code AS modele, f.entite, f.arrete, f.cycle, q.reference, f.preparateur, f.forme_conclusion AS forme, f.conclusion "
                       "FROM dbo.feuille_travail f LEFT JOIN dbo.ref_question q ON q.id = f.question_id "
                       "WHERE f.entite=? AND f.arrete=? AND f.cycle IS NOT NULL AND f.cote NOT LIKE 'Q-%' " + ("AND f.cycle=? " if cycle else "") + "ORDER BY f.cote",
                 (entite, arrete, cycle) if cycle else (entite, arrete))


def _ecrire_coffre(coffre, chemin: str, wb) -> str:
    tampon = io.BytesIO(); wb.save(tampon)
    coffre.connectToFiles().get_file_client(chemin).upload_data(tampon.getvalue(), overwrite=True)
    return "Files/" + chemin


@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def exporter_questions(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient, entite: str, arrete: str, cycle: str) -> str:
    """Etape 2, « Ouvrir le classeur des questions du cycle » : les questions au programme, reponses comprises."""
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    d = _questions(base, entite, arrete, cycle)
    if not d:
        raise fn.UserThrownError(f"Aucune question du cycle {cycle} au programme de cet arrêté.", {})
    _feuille_questions(wb, "Q-" + cycle, d)
    chemin = _ecrire_coffre(coffre, f"exports/questions_{cycle}_{entite}_{arrete}.xlsx", wb)
    return f"Classeur écrit au coffre : {chemin}, {len(d)} questions."


@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def exporter_od(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient, entite: str, arrete: str, reference: str = "") -> str:
    """Pop-up OD, « Exporter les OD liées » (reference donnee) ou etape 4, « Exporter le gabarit » (OD libres) : preremplies."""
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    ref = _vide(reference)
    d = _od(base, entite, arrete, reference=ref, libres=ref is None)
    _feuille_od(wb, ("OD " + ref) if ref else "OD libres", d)
    chemin = _ecrire_coffre(coffre, f"exports/od_{ref or 'libres'}_{entite}_{arrete}.xlsx", wb)
    return f"Classeur écrit au coffre : {chemin}, {len(d)} ligne(s) préremplie(s)."


@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def exporter_feuille(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient, cote: str) -> str:
    """Pop-up Feuille, « Exporter » : le gabarit generique prerempli de la feuille."""
    from openpyxl import Workbook
    d = _lire(base, "SELECT f.cote, f.modele_code AS modele, f.entite, f.arrete, f.cycle, q.reference, f.preparateur, f.forme_conclusion AS forme, f.conclusion "
                    "FROM dbo.feuille_travail f LEFT JOIN dbo.ref_question q ON q.id = f.question_id WHERE f.cote=?", (cote,))
    if not d:
        raise fn.UserThrownError(f"La feuille {cote} n'existe pas.", {})
    wb = Workbook(); wb.remove(wb.active)
    _feuille_feuille(wb, cote, d[0])
    chemin = _ecrire_coffre(coffre, f"exports/feuille_{cote}.xlsx", wb)
    return f"Classeur écrit au coffre : {chemin}."


@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def exporter_dossier(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient, entite: str, arrete: str) -> str:
    """Etapes 3 et 6, « Exporter le dossier de travail » : Balance, Programme, puis par cycle Q, FT, OD ; OD libres."""
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Balance")
    colb = [("Compte", "compte"), ("Libellé", "libelle"), ("Balance initiale", "b0_initiale"), ("OD de révision", "od_revision"),
            ("OD de valorisation", "od_valorisation"), ("Balance finale", "balance_finale"), ("Cycles", "cycles"), ("Couverture", "etat")]
    _entete(ws, colb)
    _lignes(ws, colb, _lire(base, "SELECT b.compte, b.libelle, b.b0_initiale, b.od_revision, b.od_valorisation, b.balance_finale, c.cycles, c.etat "
                                  "FROM dbo.v_balance b LEFT JOIN dbo.v_couverture_balance c ON c.entite=b.entite AND c.arrete=b.arrete AND c.compte=b.compte "
                                  "WHERE b.entite=? AND b.arrete=? ORDER BY b.compte", (entite, arrete)))
    ws = wb.create_sheet("Programme")
    colp = [("Type", "type_programme"), ("Cycle", "cycle"), ("Référence", "reference"), ("Question", "enonce"), ("Active", "actif"), ("Origine", "origine")]
    prog = _lire(base, "SELECT p.type_programme, q.cycle, q.reference, q.enonce, pq.actif, pq.origine FROM dbo.programme_travail p "
                       "JOIN dbo.programme_question pq ON pq.programme_id=p.id JOIN dbo.ref_question q ON q.id=pq.question_id "
                       "WHERE p.entite=? AND p.arrete=? ORDER BY q.cycle, q.ordre", (entite, arrete))
    _entete(ws, colp); _lignes(ws, colp, prog)
    cycles = []
    for x in prog:
        if x["actif"] and x["cycle"] not in cycles:
            cycles.append(x["cycle"])
    for cy in cycles:
        _feuille_questions(wb, "Q-" + cy, _questions(base, entite, arrete, cy))
        for f in _feuilles(base, entite, arrete, cy):
            _feuille_feuille(wb, f["cote"], f)
        _feuille_od(wb, "OD-" + cy, _od(base, entite, arrete, cycle=cy), vides=6)
    _feuille_od(wb, "OD libres", _od(base, entite, arrete, libres=True))
    chemin = _ecrire_coffre(coffre, f"exports/dossier_{entite}_{arrete}.xlsx", wb)
    return f"Dossier de travail écrit au coffre : {chemin}, {len(wb.sheetnames)} onglets."


def _lignes_de(ws):
    cles = [c.value for c in ws[2]]
    out = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({k: v for k, v in zip(cles, r) if k})
    return out


@udf.context(argName="ctx")
@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def importer_classeur(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient, ctx: fn.UserDataFunctionContext,
                      entite: str, arrete: str, chemin: str) -> str:
    """Bouton « Réimporter » de toutes les etapes : le classeur depose au coffre, reconnu a ses feuilles :
    Q-<cycle> (reponses), OD <reference> ou OD libres (brouillon remplace), <cote> (feuille : forme, conclusion).
    Balance, Programme et OD-<cycle> ne s'importent pas."""
    from openpyxl import load_workbook
    import hashlib
    octets = coffre.connectToFiles().get_file_client(chemin).download_file().readall()
    wb = load_workbook(io.BytesIO(octets), data_only=True)
    par = _qui(ctx)
    faits, refus = [], []
    for ws in wb.worksheets:
        t = ws.title
        try:
            if t in ("Balance", "Programme") or t.startswith("OD-"):
                continue
            if t.startswith("Q-"):
                l = _lire(base, "SELECT 'Q-' + ? + '-P' + CAST(id AS varchar) AS cote FROM dbo.programme_travail WHERE entite=? AND arrete=?", (t[2:], entite, arrete))
                if not l:
                    raise fn.UserThrownError("Aucun programme choisi pour cet arrêté.", {})
                reponses = [{"reference": d.get("reference"), "reponse": None if d.get("reponse") is None else str(d.get("reponse")),
                             "motif": d.get("motif"), "commentaire": d.get("commentaire")} for d in _lignes_de(ws) if d.get("reference")]
                r = _executer(base, "EXEC dbo.pr_importer_reponses @cote=?, @reponses=?, @par=?", (l[0]["cote"], json.dumps(reponses, ensure_ascii=False, default=str), par))
                faits.append(f"{t} : {r.get('message', 'importé')}")
            elif t == "OD libres" or t.startswith("OD "):
                ref = None if t == "OD libres" else t[3:].strip()
                lignes = [{"journal": d.get("journal") or "ODR", "compte": str(d.get("compte")).strip(), "libelle": d.get("libelle"),
                           "debit": float(d.get("debit") or 0), "credit": float(d.get("credit") or 0), "piece": d.get("piece"), "actif": d.get("actif")}
                          for d in _lignes_de(ws) if d.get("compte") and str(d.get("compte")).strip().lower() not in ("total", "écart")]
                r = _executer(base, "EXEC dbo.pr_importer_od @entite=?, @arrete=?, @reference=?, @lignes=?, @par=?",
                              (entite, arrete, ref, json.dumps(lignes, ensure_ascii=False, default=str), par))
                faits.append(f"{t} : {r.get('message', 'importé')}")
            else:
                valeurs = {}
                for row in ws.iter_rows(min_row=1, values_only=True):
                    if len(row) >= 3 and row[2] in ("cote", "forme", "conclusion"):
                        valeurs[row[2]] = row[1]
                cote = valeurs.get("cote") or t
                r = _executer(base, "EXEC dbo.pr_importer_feuille @cote=?, @conclusion=?, @forme=?, @nom_fichier=?, @empreinte=?, @par=?",
                              (cote, _vide(valeurs.get("conclusion")), _vide(valeurs.get("forme")), chemin.rsplit("/", 1)[-1],
                               hashlib.sha256(octets).hexdigest().upper(), par))
                faits.append(f"{t} : {r.get('message', 'importé')}")
        except fn.UserThrownError as e:
            refus.append(f"{t} : {e}")
    texte = f"{len(faits)} onglet(s) importé(s)." + (" " + " ".join(faits) if faits else "")
    if refus:
        texte += " Refusés : " + " ; ".join(refus)
    return texte
