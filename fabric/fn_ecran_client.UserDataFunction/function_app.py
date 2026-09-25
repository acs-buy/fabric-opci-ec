"""fn_ecran_client : les boutons de l'ecran Client et Acceptation, une fonction par bouton.

Ecrit le 16/09/2026 d'apres la maquette 71 validee par le candidat. Chaque fonction appelle UNE
procedure de la base DossierOPCI et rend la phrase que le bouton affiche. Les regles metier ne sont
jamais ici : elles sont dans la procedure, qui refuse en francais, et le refus remonte tel quel.

COPIE DE REFERENCE, VERSIONNEE. La pose se fait par l'API (updateDefinition), Fabric IQ la conduit,
et se relit par getDefinition. La connexion a la base porte l'alias « DossierOPCI », meme alias que
fn_documents et fn_valorisation ; celle du coffre « Coffre ».

CE QUE LA PLATEFORME IMPOSE, verifie sur learn.microsoft.com le 16/09/2026 :
  - parametres en camelCase, sans soulignement, types annotes, retour annote ;
  - l'utilisateur qui clique est connu par UserDataFunctionContext.executing_user, cle
    PreferredUsername, decorateur @udf.context : c'est lui qui devient @par. Le rapport ne le passe
    pas, il ne peut donc pas le falsifier ;
  - une erreur attendue se rend par fn.UserThrownError, dont le message est affiche a l'appelant ;
  - pas d'execution dans Desktop, seulement sur le rapport publie ; les lecteurs recoivent
    « Execute functions » sur l'element.

PATRON D'APPEL. FabricSqlConnection.connect() rend une connexion pyodbc, autocommit desactive :
sans commit explicite, l'EXEC serait annule a la fermeture. Les procedures d'ecran (pr_ecran_*)
n'echouent pas : elles ecrivent le refus sur la ligne concernee, colonne message_ecran, pour que le
rapport le lise. La fonction relit cette colonne apres l'appel et le rend en erreur si elle est
posee, en succes sinon. Les procedures neuves du 204 et du 205 rendent un jeu de resultats
(code, message) et levent le refus : la fonction rend le dernier jeu de resultats.
"""

import datetime
import hashlib
import logging
import re

import fabric.functions as fn

udf = fn.UserDataFunctions()


# ----------------------------------------------------------------------------- outils

def _qui(ctx) -> str:
    """L'utilisateur connecte, tel que la plateforme l'authentifie."""
    u = (ctx.executing_user or {}).get("PreferredUsername") or (ctx.executing_user or {}).get("Oid")
    if not u:
        raise fn.UserThrownError("L'utilisateur connecté n'est pas identifiable ; l'action n'est pas exécutée.", {})
    _UTILISATEUR["upn"] = u                      # retenu pour journaliser un refus apres rollback
    return u


def _message_sql(exc) -> str:
    """Le texte francais d'un THROW, sans le prefixe du pilote ODBC."""
    s = str(exc)
    m = re.search(r"\[SQL Server\](.*?)(\s\(\d+\)\s\(SQLExecDirectW\)|$)", s, re.S)
    return (m.group(1) if m else s).strip()


def _entite_probable(params):
    """Le code d'entite d'un appel, par convention le premier parametre court et alphanumerique."""
    for p in (params or ()):
        if isinstance(p, str) and 2 < len(p) <= 20 and p.replace("-", "").replace("_", "").isalnum():
            return p
    return None


def _nom_procedure(sql):
    m = re.search(r"EXEC dbo\.(\w+)", sql or "")
    return m.group(1)[:128] if m else "inconnue"


def _poser_message_dans(cur, sql, message, params, genre):
    """Ecrit la boite aux lettres dans la transaction en cours, sans committer : l'appelant commit."""
    try:
        if message:
            cur.execute("EXEC dbo.pr_poser_message_ecran @pour=?, @message=?, @genre=?, @entite=?, @procedure_nom=?",
                        (_qui_sans_lever(), str(message)[:2000], genre,
                         _entite_probable(params), _nom_procedure(sql)))
    except Exception:
        pass


def _poser_message(conn, sql, message, params, genre):
    """Ecrit le dernier geste dans la boite aux lettres de l'ecran, reussite comme refus.

    POURQUOI UNE BOITE AUX LETTRES ET NON LE JOURNAL, mesure du 21/09/2026 : un refus lu dans
    journal_refus survivait a la reussite suivante, et le reviseur lisait « deja au perimetre » alors
    que son ajout venait de reussir. journal_refus est une TRACE, elle s'ajoute et ne s'efface pas ;
    le pied de panneau est un ETAT, une seule ligne vivante par personne et par dossier, remplacee a
    chaque geste. Les 2 objets sont desormais distincts (script 221).

    Ne leve jamais."""
    try:
        cur = conn.cursor()
        cur.execute("EXEC dbo.pr_poser_message_ecran @pour=?, @message=?, @genre=?, @entite=?, @procedure_nom=?",
                    (_qui_sans_lever(), (message or "")[:2000], genre,
                     _entite_probable(params), _nom_procedure(sql)))
        conn.commit()
    except Exception:
        try:
            conn.rollback()
        except Exception:
            pass


def _journaliser_refus(conn, sql, message, params):
    """Ecrit le refus dans journal_refus APRES le rollback, dans une transaction neuve.

    POURQUOI ICI ET NON DANS LA PROCEDURE, mesure du 21/09/2026 : une procedure qui leve un THROW voit
    son propre INSERT de journal annule par le rollback que fait cette fonction, la connexion etant sans
    autocommit et SQL Server n'ayant pas de transaction autonome. Un refus journalise par la procedure
    elle-meme (pr_supprimer_filiale) ne survit que si l'appelant ne fait pas rollback.

    SANS CE JOURNAL, LE REVISEUR NE LIT RIEN : au temps 2, ajouter une filiale deja au perimetre laissait
    l'ecran muet, la base refusant sans qu'aucune table lue par l'ecran ne porte le message.

    Ne leve jamais : un journal qui echoue ne doit pas masquer le refus qu'il rapporte."""
    try:
        nom = "inconnue"
        m = re.search(r"EXEC dbo\.(\w+)", sql or "")
        if m:
            nom = m.group(1)
        entite = None
        for p in (params or ()):
            if isinstance(p, str) and 2 < len(p) <= 20 and p.replace("-", "").replace("_", "").isalnum():
                entite = p
                break
        cur = conn.cursor()
        cur.execute("EXEC dbo.pr_journaliser_refus @procedure_nom=?, @message=?, @par=?, @entite=?",
                    (nom[:128], (message or "")[:2000], _qui_sans_lever(), entite))
        cur.execute("EXEC dbo.pr_poser_message_ecran @pour=?, @message=?, @genre=?, @entite=?, @procedure_nom=?",
                    (_qui_sans_lever(), (message or "")[:2000], "REFUS", entite, nom[:128]))
        conn.commit()
    except Exception:                            # le journal ne doit jamais masquer le refus
        try:
            conn.rollback()
        except Exception:
            pass


_UTILISATEUR = {"upn": None}


def _qui_sans_lever():
    """L'utilisateur retenu au dernier appel de _qui, ou « inconnu » : le journal n'a pas de contexte."""
    return _UTILISATEUR["upn"] or "inconnu"


def _executer(base: fn.FabricSqlConnection, sql: str, params: tuple) -> dict:
    """Execute une procedure qui rend (…, message) et leve ses refus. Rend la derniere ligne."""
    conn = base.connect()
    try:
        cur = conn.cursor()
        cur.execute(sql, params)
        derniere = None
        while True:
            if cur.description:
                colonnes = [d[0] for d in cur.description]
                for ligne in cur.fetchall():
                    derniere = dict(zip(colonnes, ligne))
            if not cur.nextset():
                break
        # LE SUCCES S'AFFICHE AUSSI : sans cela le reviseur ne voit rien quand son geste marche, et
        # le pied garde le message du geste precedent.
        _poser_message_dans(cur, sql, (derniere or {}).get("message"), params, "SUCCES")
        conn.commit()
        return derniere or {}
    except Exception as e:                      # pyodbc.Error, sans dependre de son import
        conn.rollback()
        message = _message_sql(e)
        _journaliser_refus(conn, sql, message, params)
        raise fn.UserThrownError(message, {"sql": sql.split("(")[0].strip()})
    finally:
        conn.close()


def _executer_ecran(base: fn.FabricSqlConnection, sql: str, params: tuple, relire: str,
                    relire_params: tuple, succes: str) -> str:
    """Execute une pr_ecran_* qui ecrit son refus dans message_ecran, puis relit cette colonne."""
    conn = base.connect()
    try:
        cur = conn.cursor()
        cur.execute(sql, params)
        while cur.nextset():
            pass
        _poser_message_dans(cur, sql, succes, params, "SUCCES")
        conn.commit()
        cur.execute(relire, relire_params)
        ligne = cur.fetchone()
    except Exception as e:
        conn.rollback()
        message = _message_sql(e)
        _journaliser_refus(conn, sql, message, params)
        raise fn.UserThrownError(message, {"sql": sql.split("(")[0].strip()})
    finally:
        conn.close()
    message = ligne[0] if ligne else None
    if message:
        raise fn.UserThrownError(message, {})
    return succes


def _vide(s) -> str:
    return None if s is None or str(s).strip() == "" else str(s).strip()


# ----------------------------------------------------------------------------- etape 1 : creer le dossier

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def creer_client(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                 code: str, denomination: str, formeVehicule: str, adresse1: str, codePostal: str,
                 ville: str, dirigeantNom: str, dirigeantQualite: str,
                 siren: str = "", formeSociale: str = "", adresse2: str = "", pays: str = "FR",
                 contactNom: str = "", contactCourriel: str = "", contactTelephone: str = "",
                 cloture: str = "31/12", periodiciteVl: str = "", siteUrl: str = "") -> str:
    """Bouton « Créer le dossier » : les 15 champs de la maquette, plus le lien du site SharePoint cree par le
    reviseur via Teams (regle du 19/09/2026), vers pr_creer_client (script 213) : le client, son questionnaire
    d'acceptation et son site dans le meme geste."""
    r = _executer(base,
                  "EXEC dbo.pr_creer_client @code=?, @denomination=?, @siren=?, @forme_vehicule=?, @forme_sociale=?, "
                  "@adresse_1=?, @adresse_2=?, @code_postal=?, @ville=?, @pays=?, @dirigeant_nom=?, @dirigeant_qualite=?, "
                  "@contact_nom=?, @contact_courriel=?, @contact_telephone=?, @cloture=?, @periodicite_vl=?, @site_url=?, @par=?",
                  (code, denomination, _vide(siren), _vide(formeVehicule), _vide(formeSociale), adresse1, _vide(adresse2),
                   codePostal, ville, pays or "FR", dirigeantNom, dirigeantQualite, _vide(contactNom),
                   _vide(contactCourriel), _vide(contactTelephone), cloture or "31/12", _vide(periodiciteVl), _vide(siteUrl), _qui(ctx)))
    return r.get("message", "Client créé.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def modifier_client(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext, code: str,
                    denomination: str = "", siren: str = "", formeVehicule: str = "", formeSociale: str = "",
                    adresse1: str = "", adresse2: str = "", codePostal: str = "", ville: str = "", pays: str = "",
                    dirigeantNom: str = "", dirigeantQualite: str = "", contactNom: str = "", contactCourriel: str = "",
                    contactTelephone: str = "", cloture: str = "", siteUrl: str = "") -> str:
    """Bouton « Enregistrer les modifications » de la fiche du dossier : pr_modifier_client (scripts 212, 213).
    Le code ne change pas ; un champ vide reste inchange ; « - » efface un champ facultatif ; un lien SharePoint
    fourni remplace le precedent."""
    r = _executer(base,
                  "EXEC dbo.pr_modifier_client @code=?, @denomination=?, @siren=?, @forme_vehicule=?, @forme_sociale=?, "
                  "@adresse_1=?, @adresse_2=?, @code_postal=?, @ville=?, @pays=?, @dirigeant_nom=?, @dirigeant_qualite=?, "
                  "@contact_nom=?, @contact_courriel=?, @contact_telephone=?, @cloture=?, @site_url=?, @par=?",
                  (code, _vide(denomination), _vide(siren), _vide(formeVehicule), _vide(formeSociale), _vide(adresse1),
                   _vide(adresse2), _vide(codePostal), _vide(ville), _vide(pays), _vide(dirigeantNom), _vide(dirigeantQualite),
                   _vide(contactNom), _vide(contactCourriel), _vide(contactTelephone), _vide(cloture), _vide(siteUrl), _qui(ctx)))
    return r.get("message", "Informations mises à jour.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def designer_role(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                  entite: str, role: str, personne: str, connexion: str = "") -> str:
    """Bouton « Désigner » de la section Equipe de la mission (fiche du dossier) : pr_designer_role_mission,
    script 219. Designer quelqu'un ferme le mandat en cours du meme role, la base l'imposant. La connexion
    est l'UPN : sans elle, la personne ne sera pas reconnue a l'ecran."""
    r = _executer(base,
                  "EXEC dbo.pr_designer_role_mission @entite=?, @role=?, @personne=?, @connexion=?, @par=?",
                  (entite, role, personne, _vide(connexion), _qui(ctx)))
    return r.get("message", "Rôle désigné.")


# ----------------------------------------------------------------------------- etape 2 : lister les filiales

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def ajouter_filiale(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                    client: str, code: str, droitsDeVote: float, denomination: str = "",
                    formeSociale: str = "SCI") -> str:
    """Bouton « Ajouter cette filiale » : 80 ou 0,80 sont admis pour 80 %."""
    dv = float(droitsDeVote)
    if dv > 1:
        dv = dv / 100.0
    r = _executer(base,
                  "EXEC dbo.pr_ajouter_filiale @client=?, @code=?, @denomination=?, @forme_sociale=?, @droits_de_vote=?, @par=?",
                  (client, code, _vide(denomination), formeSociale or "SCI", round(dv, 6), _qui(ctx)))
    return r.get("message", "Filiale ajoutée.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def modifier_filiale(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                     client: str, code: str, denomination: str = "", formeSociale: str = "",
                     droitsDeVote: str = "") -> str:
    """Bouton « Enregistrer la filiale » (etat 2c) : pr_modifier_filiale, script 214. Le code ne change pas ;
    un champ vide reste inchange ; les droits de vote s'ecrivent 80 ou 0,80."""
    dv = None
    if _vide(droitsDeVote) is not None:
        dv = float(str(droitsDeVote).replace(",", ".").replace("%", "").strip())
        if dv > 1:
            dv = dv / 100.0
        dv = round(dv, 6)
    r = _executer(base,
                  "EXEC dbo.pr_modifier_filiale @client=?, @code=?, @denomination=?, @forme_sociale=?, @droits_de_vote=?, @par=?",
                  (client, code, _vide(denomination), _vide(formeSociale), dv, _qui(ctx)))
    return r.get("message", "Filiale mise à jour.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def supprimer_filiale(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                      client: str, code: str) -> str:
    """Bouton « Supprimer » du pop-up de confirmation (etat 2d) : pr_supprimer_filiale, script 214.
    La base refuse si la filiale porte des donnees de mission, et dit lesquelles."""
    r = _executer(base, "EXEC dbo.pr_supprimer_filiale @client=?, @code=?, @par=?", (client, code, _qui(ctx)))
    return r.get("message", "Filiale retirée du périmètre.")


# ----------------------------------------------------------------------------- etape 3 : le questionnaire

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def ouvrir_questionnaire_acceptation(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                                     entite: str) -> str:
    """Bouton « Ouvrir la saisie en ligne » : la feuille et ses 110 questions, puis l'acceptation."""
    r = _executer(base, "EXEC dbo.pr_ouvrir_questionnaire_acceptation @entite=?, @par=?", (entite, _qui(ctx)))
    return r.get("message", "Questionnaire ouvert.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def repondre_question(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                      cote: str, reference: str, reponse: str, motif: str = "", commentaire: str = "") -> str:
    """Une reponse, a l'ecran : OUI, NON, NA avec motif, ou la valeur d'un choix. Sert aussi a la Revision."""
    _executer(base,
              "EXEC dbo.pr_repondre_question @cote=?, @reference=?, @reponse=?, @motif_non_applicable=?, @commentaire=?, @par=?",
              (cote, reference, _vide(reponse), _vide(motif), _vide(commentaire), _qui(ctx)))
    return f"Réponse à {reference} enregistrée."


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def importer_reponses(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                      cote: str, reponses: list) -> str:
    """Bouton « Réimporter » : la liste des reponses lues du classeur, tout ou rien."""
    import json
    r = _executer(base, "EXEC dbo.pr_importer_reponses @cote=?, @reponses=?, @par=?",
                  (cote, json.dumps(reponses, ensure_ascii=False), _qui(ctx)))
    return r.get("message", "Réponses importées.")


# ----------------------------------------------------------------------------- etape 4 : joindre les pieces

@udf.context(argName="ctx")
@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def inscrire_piece(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient,
                   ctx: fn.UserDataFunctionContext, entite: str, chemin: str, nature: str,
                   arrete: str = "", question: str = "") -> str:
    """Bouton « Déposer une pièce », second temps : le fichier est deja televerse au coffre (chemin
    relatif a Files/), la fonction en calcule l'empreinte et l'inscrit par pr_deposer_piece."""
    fichiers = coffre.connectToFiles()
    client = fichiers.get_file_client(chemin)
    octets = client.download_file().readall()
    empreinte = hashlib.sha256(octets).hexdigest().upper()
    nom = chemin.rsplit("/", 1)[-1]
    r = _executer(base,
                  "EXEC dbo.pr_deposer_piece @entite=?, @nom_fichier=?, @chemin_coffre=?, @empreinte=?, @nature=?, "
                  "@arrete=?, @question=?, @par=?",
                  # meme convention que les pieces du jeu : /Coffre/<entite>/<arrete>/..., relative a Files/
                  (entite, nom, "/Coffre/" + chemin.lstrip("/"), empreinte, nature, _vide(arrete), _vide(question), _qui(ctx)))
    return r.get("message", "Pièce inscrite.") + f" {len(octets) // 1024} Ko."


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def retirer_piece(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                  entite: str, pieceId: str, motif: str) -> str:
    """Bouton « Retirer la pièce » (temps 3, liste des pieces) : pr_retirer_piece, script 215. Le fichier reste
    au coffre ; le rattachement au dossier part, avec sa trace et son motif. La base refuse si la piece fonde
    une donnee de mission, ou si le dossier est verrouille."""
    r = _executer(base, "EXEC dbo.pr_retirer_piece @entite=?, @piece_id=?, @motif=?, @par=?",
                  (entite, int(str(pieceId).strip()), motif, _qui(ctx)))
    return r.get("message", "Pièce retirée du dossier.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def reclamer_document(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                      entite: str, arrete: str, livrable: str) -> str:
    """Bouton « Réclamer » : la demande de document, tracee sur l'arrete."""
    return _executer_ecran(base,
                           "EXEC dbo.pr_ecran_demander_document @entite=?, @arrete=?, @livrable=?, @par=?",
                           (entite, arrete, livrable, _qui(ctx)),
                           "SELECT message_ecran FROM dbo.ref_arrete WHERE entite=? AND arrete=?", (entite, arrete),
                           f"Document {livrable} réclamé pour l'arrêté {arrete}.")


# ----------------------------------------------------------------------------- etape 5 : soumettre au visa

def _superviseur(ctx, superviseur: str) -> None:
    """Le parametre « superviseur » est lie a la mesure Superviseur connecte : vide pour un preparateur,
    donc le bouton reste inactif (parametre obligatoire non fourni). Ici, seconde couche : la valeur
    fournie doit etre l'utilisateur authentifie lui-meme, sinon un parametre force est refuse."""
    if not superviseur or superviseur.strip().lower() != _qui(ctx).strip().lower():
        raise fn.UserThrownError("Le visa est réservé au superviseur connecté ; l'action n'est pas exécutée.", {})


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def approuver_acceptation(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                          entite: str, decision: str, superviseur: str, motif: str = "") -> str:
    """Boutons « Accepter la mission » et « Refuser » : decision ACCEPTEE ou REFUSEE, par l'associe.
    La procedure refuse le proposant et le role insuffisant ; le rapport grise le bouton en amont."""
    _superviseur(ctx, superviseur)
    return _executer_ecran(base,
                           "EXEC dbo.pr_ecran_approuver_acceptation @entite=?, @decision=?, @motif=?, @par=?",
                           (entite, decision, _vide(motif), _qui(ctx)),
                           "SELECT TOP 1 message_ecran FROM dbo.acceptation_mission WHERE entite=? ORDER BY cree_le DESC",
                           (entite,),
                           f"Acceptation de {entite} : décision {decision.lower()} visée.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def reprendre_acceptation(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                          entite: str, motif: str) -> str:
    """Bouton « Reprendre » : une acceptation refusee revient a l'etat ouvert, avec motif."""
    return _executer_ecran(base,
                           "EXEC dbo.pr_ecran_reprendre_acceptation @entite=?, @motif=?, @par=?",
                           (entite, motif, _qui(ctx)),
                           "SELECT TOP 1 message_ecran FROM dbo.acceptation_mission WHERE entite=? ORDER BY cree_le DESC",
                           (entite,),
                           f"Acceptation de {entite} reprise ; le questionnaire se complète à nouveau.")


# ----------------------------------------------------------------------------- etape 6 : ouvrir l'arrete

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def ouvrir_arrete(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                  entite: str, exerciceDebut: str, exerciceFin: str, periodiciteVl: str) -> str:
    """Bouton « Ouvrir l'arrêté » : planifie les arretes de l'exercice selon la periodicite, ouvre le
    premier. Dates au format jj/mm/aaaa ou aaaa-mm-jj."""
    def d(s):
        s = s.strip()
        for f in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.datetime.strptime(s, f).date()
            except ValueError:
                pass
        raise fn.UserThrownError(f"Date illisible : « {s} ». Attendu jj/mm/aaaa.", {})
    r = _executer(base,
                  "EXEC dbo.pr_ecran_ouvrir_arrete @entite=?, @exercice_debut=?, @exercice_fin=?, @periodicite_vl=?, @par=?",
                  (entite, d(exerciceDebut), d(exerciceFin), periodiciteVl.upper(), _qui(ctx)))
    return r.get("message", "Arrêté planifié.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def rouvrir_arrete(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                   entite: str, arrete: str, motif: str) -> str:
    """Bouton « Rouvrir » : un arrete clos se rouvre avec motif, trace sur sa ligne."""
    return _executer_ecran(base,
                           "EXEC dbo.pr_ecran_rouvrir_arrete @entite=?, @arrete=?, @motif=?, @par=?",
                           (entite, arrete, motif, _qui(ctx)),
                           "SELECT message_ecran FROM dbo.ref_arrete WHERE entite=? AND arrete=?", (entite, arrete),
                           f"Arrêté {arrete} de {entite} rouvert.")


# ----------------------------------------------------------------------------- client accepte : le maintien

@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def ouvrir_maintien(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                    entite: str, arreteConclu: str) -> str:
    """Bouton « Ouvrir le maintien » sur la ligne d'un arrete conclu : la feuille et ses questions."""
    r = _executer(base, "EXEC dbo.pr_ouvrir_questionnaire_maintien @entite=?, @arrete_conclu=?, @par=?",
                  (entite, arreteConclu, _qui(ctx)))
    return r.get("message", "Maintien ouvert.")


@udf.context(argName="ctx")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def approuver_maintien(base: fn.FabricSqlConnection, ctx: fn.UserDataFunctionContext,
                       entite: str, arreteConclu: str, decision: str, superviseur: str, motif: str = "") -> str:
    """Bouton « Viser le maintien » : decision MAINTENU ou ROMPU, par le superviseur."""
    _superviseur(ctx, superviseur)
    return _executer_ecran(base,
                           "EXEC dbo.pr_ecran_approuver_maintien @entite=?, @arrete_conclu=?, @decision=?, @motif=?, @par=?",
                           (entite, arreteConclu, decision, _vide(motif), _qui(ctx)),
                           "SELECT TOP 1 message_ecran FROM dbo.maintien_mission WHERE entite=? AND arrete_conclu=? ORDER BY cree_le DESC",
                           (entite, arreteConclu),
                           f"Maintien {arreteConclu[:4]} de {entite} : décision {decision.lower()} visée.")


# ----------------------------------------------------------------------------- la voie du classeur

# Les 2 voies de la maquette, a l'ecran ou par classeur, ecrivent au meme endroit. Un bouton de rapport
# ne lit pas de fichier : le classeur est depose au coffre (Files/depots/<entite>/), puis le bouton
# « Reimporter » donne son chemin. La lecture est celle de 61_DEPLOIEMENT/importer_ecran_client.py :
# la feuille dit la nature, la ligne 2 des en-tetes donne le parametre. Bibliotheque : openpyxl.

_ORDRE_CLIENT = ["code", "denomination", "siren", "forme_vehicule", "forme_sociale", "adresse_1", "adresse_2",
                 "code_postal", "ville", "pays", "dirigeant_nom", "dirigeant_qualite", "contact_nom",
                 "contact_courriel", "contact_telephone"]


def _lire_classeur(octets: bytes):
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(octets), data_only=True).active
    cles = [c.value for c in ws[2]]
    lignes = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        d = {k: (None if v is None or str(v).strip() == "" else v) for k, v in zip(cles, r) if k}
        if d.get("code") is None and d.get("reference") is None:
            continue
        lignes.append(d)
    return ws.title, lignes


def _cote_maintien(base: fn.FabricSqlConnection, client: str) -> str:
    """La cote de maintien porte l'annee : on prend la feuille MAINTIEN la plus recente du client."""
    conn = base.connect()
    try:
        cur = conn.cursor()
        cur.execute("SELECT TOP 1 cote FROM dbo.feuille_travail WHERE entite=? AND phase='MAINTIEN' ORDER BY arrete DESC",
                    (client,))
        ligne = cur.fetchone()
    finally:
        conn.close()
    if not ligne:
        raise fn.UserThrownError("Aucun questionnaire de maintien n'est ouvert pour ce client.", {})
    return ligne[0]


@udf.context(argName="ctx")
@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def importer_classeur(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient,
                      ctx: fn.UserDataFunctionContext, chemin: str, client: str = "") -> str:
    """Bouton « Réimporter » des 3 etapes : le classeur depose au coffre, Client, Filiales ou
    Questionnaire, entre par les memes procedures que la saisie a l'ecran."""
    import json
    octets = coffre.connectToFiles().get_file_client(chemin).download_file().readall()
    feuille, lignes = _lire_classeur(octets)
    par = _qui(ctx)
    if feuille == "Client":
        if len(lignes) != 1:
            raise fn.UserThrownError("Le classeur Client porte %d ligne(s) ; il en porte une et une seule." % len(lignes), {})
        d = lignes[0]
        r = _executer(base,
                      "EXEC dbo.pr_creer_client " + ", ".join("@%s=?" % k for k in _ORDRE_CLIENT) + ", @par=?",
                      tuple(str(d.get(k)) if d.get(k) is not None else None for k in _ORDRE_CLIENT) + (par,))
        return r.get("message", "Client créé.")
    if not client:
        raise fn.UserThrownError("Le classeur %s s'importe sur un client : sélectionnez-le." % feuille, {})
    if feuille == "Filiales":
        faites, refus = [], []
        for d in lignes:
            code = str(d.get("code"))
            try:
                dv = float(str(d.get("droits_de_vote")).replace(",", ".").replace("%", "").strip())
                dv = dv / 100.0 if dv > 1 else dv
                _executer(base, "EXEC dbo.pr_ajouter_filiale @client=?, @code=?, @denomination=?, @droits_de_vote=?, @par=?",
                          (client, code, d.get("denomination"), round(dv, 6), par))
                faites.append(code)
            except fn.UserThrownError as e:
                refus.append("%s : %s" % (code, e))
            except (TypeError, ValueError):
                refus.append("%s : droits de vote illisibles" % code)
        texte = "%d filiale(s) ajoutée(s) au périmètre de %s." % (len(faites), client)
        if refus:
            texte += " Refusées : " + " ; ".join(refus)
        return texte
    if feuille == "Questionnaire":
        reponses = [{"reference": d.get("reference"), "reponse": None if d.get("reponse") is None else str(d.get("reponse")),
                     "motif": d.get("motif"), "commentaire": d.get("commentaire")} for d in lignes if d.get("reference")]
        maintien = any(str(x["reference"]).startswith("MTN") or str(x["reference"]).startswith("MAINTIEN") for x in reponses)
        cote = _cote_maintien(base, client) if maintien else "ACC-" + client
        r = _executer(base, "EXEC dbo.pr_importer_reponses @cote=?, @reponses=?, @par=?",
                      (cote, json.dumps(reponses, ensure_ascii=False, default=str), par))
        return r.get("message", "Réponses importées.")
    raise fn.UserThrownError("Feuille « %s » inconnue : Client, Filiales ou Questionnaire." % feuille, {})


@udf.context(argName="ctx")
@udf.connection(alias="Coffre", argName="coffre")
@udf.connection(alias="DossierOPCI", argName="base")
@udf.function()
def exporter_questionnaire(base: fn.FabricSqlConnection, coffre: fn.FabricLakehouseClient,
                           ctx: fn.UserDataFunctionContext, cote: str) -> str:
    """Bouton « Exporter le questionnaire » : le classeur des questions de la feuille, reponses
    comprises, ecrit au coffre dans Files/exports/<cote>.xlsx ; la fonction rend le chemin."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    conn = base.connect()
    try:
        cur = conn.cursor()
        cur.execute("SELECT q.section, q.reference, q.enonce, q.type_reponse, q.obligatoire, "
                    "COALESCE(fq.reponse, fq.reponse_valeur), fq.motif_non_applicable, fq.commentaire, q.doc_attendu "
                    "FROM dbo.feuille_question fq JOIN dbo.ref_question q ON q.id = fq.question_id "
                    "WHERE fq.cote = ? ORDER BY q.ordre", (cote,))
        lignes = cur.fetchall()
    finally:
        conn.close()
    if not lignes:
        raise fn.UserThrownError("Le questionnaire %s n'existe pas ou ne porte aucune question." % cote, {})
    wb = Workbook()
    ws = wb.active
    ws.title = "Questionnaire"
    entetes = [("Section", "section"), ("Référence", "reference"), ("Question", "enonce"), ("Type", "type"),
               ("Obligatoire", "obligatoire"), ("Réponse", "reponse"), ("Motif si non applicable", "motif"),
               ("Commentaire", "commentaire"), ("Pièce attendue", "doc_attendu")]
    for j, (lib, cle) in enumerate(entetes, 1):
        c = ws.cell(row=1, column=j, value=lib)
        c.fill = PatternFill("solid", fgColor="1D1A10")
        c.font = Font(bold=True, color="FFFFFF") if cle in ("reponse", "motif", "commentaire") else Font(color="D9D2C0")
        ws.cell(row=2, column=j, value=cle).font = Font(italic=True, color="8A8578", size=9)
    for i, l in enumerate(lignes, 3):
        for j, v in enumerate(l, 1):
            ws.cell(row=i, column=j, value=("oui" if v else "non") if j == 5 else v)
    ws.freeze_panes = "A3"
    tampon = io.BytesIO()
    wb.save(tampon)
    chemin = "exports/%s.xlsx" % cote
    coffre.connectToFiles().get_file_client(chemin).upload_data(tampon.getvalue(), overwrite=True)
    return "Classeur écrit au coffre : Files/%s, %d questions, par %s." % (chemin, len(lignes), _qui(ctx))
